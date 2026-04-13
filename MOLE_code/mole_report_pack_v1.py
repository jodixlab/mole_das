"""mole_report_pack_v1

MOLE Control DAS - Report Pack v1

Creates a compact, portable report pack (JSON + CSV + optional PDF) for a
session/run folder produced by the DAQ Runner.

CLI examples:
  python mole_report_pack_v1.py --config ".../runner_config.json"
  python mole_report_pack_v1.py --session-dir ".../sessions/2026-02-02/RUN_123" 

The report pack is written to:
  <session_dir>/exports/report_pack_v1/

Notes:
  - The PDF output is optional (requires reportlab).
  - The JSON/CSV outputs are always produced.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import re
import sqlite3
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
from xml.sax.saxutils import escape as _xml_escape

try:
    import mole_ftir_reference_v1 as mole_ftir_reference
except Exception:
    mole_ftir_reference = None

try:
    import mole_ftir_offset_recommendations_v1 as mole_ftir_offset_recommendations
except Exception:
    mole_ftir_offset_recommendations = None

try:
    import mole_ftir_validation_v1 as mole_ftir_validation
except Exception:
    mole_ftir_validation = None

try:
    import mole_method_spec_engine as mole_spec_engine
except Exception:
    mole_spec_engine = None

try:
    import mole_postcal_policy as mole_postcal_policy
except Exception:
    mole_postcal_policy = None


# -----------------------------
# Small utils
# -----------------------------

def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")


def _safe_float(v: Any) -> Optional[float]:
    try:
        if v is None or v == "":
            return None
        return float(v)
    except Exception:
        return None


def _fmt_num(v: Any, ndp: int = 3, fallback: str = "") -> str:
    try:
        f = _safe_float(v)
        if f is None or math.isnan(f) or math.isinf(f):
            return fallback
        return f"{f:,.{int(ndp)}f}"
    except Exception:
        return fallback


def _parse_iso_dt(text: Any) -> Optional[datetime]:
    try:
        s = str(text or "").strip()
        if not s:
            return None
        if s.endswith("Z"):
            s = s[:-1] + "+00:00"
        dt = datetime.fromisoformat(s)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt
    except Exception:
        return None


def _read_json(path: Path) -> Optional[Dict[str, Any]]:
    try:
        if not path.exists():
            return None
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except Exception:
        return None


def _has_value(v: Any) -> bool:
    if v is None:
        return False
    if isinstance(v, str):
        return bool(v.strip())
    if isinstance(v, (list, tuple, set, dict)):
        return len(v) > 0
    return True


def _first_present(*values: Any) -> Any:
    for value in values:
        if _has_value(value):
            return value
    return None


def _ftir_validation_cfg_from_session(session: Dict[str, Any]) -> Dict[str, Any]:
    block = session.get("ftir_validation") if isinstance(session.get("ftir_validation"), dict) else {}
    analytes_default: List[str] = []
    try:
        pollutants = session.get("pollutants") if isinstance(session.get("pollutants"), dict) else {}
        prescriptions = pollutants.get("prescriptions") if isinstance(pollutants.get("prescriptions"), list) else []
        analytes_default = [
            str((item.get("code") if isinstance(item, dict) else "") or "").strip().upper()
            for item in prescriptions
            if str((item.get("code") if isinstance(item, dict) else "") or "").strip()
        ]
    except Exception:
        analytes_default = []
    if mole_ftir_validation is None:
        return {
            "enabled": bool(block.get("enabled")),
            "validation_mode": str(block.get("validation_mode") or "METHOD_301_INFORMED_COMPARISON"),
            "analytes": analytes_default,
            "ftir_file_path": str(block.get("ftir_file_path") or "").strip(),
        }
    return mole_ftir_validation.normalize_config(block, analytes_default=analytes_default)


def _locked_ftir_validation_snapshot(cfg: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    try:
        if not bool(cfg.get("review_locked")):
            return None
        snapshot = cfg.get("review_snapshot") if isinstance(cfg.get("review_snapshot"), dict) else {}
        json_txt = str(snapshot.get("json_path") or "").strip()
        if not json_txt:
            return None
        json_path = Path(json_txt).expanduser()
        if not json_path.exists():
            return None
        obj = _read_json(json_path)
        return obj if isinstance(obj, dict) else None
    except Exception:
        return None


def _ctx_field(
    value: Any = None,
    source: str = "",
    status: Optional[str] = None,
    note: str = "",
) -> Dict[str, Any]:
    resolved_status = str(status or ("Available" if _has_value(value) else "Gap")).strip() or "Gap"
    return {
        "status": resolved_status,
        "value": value,
        "source": str(source or ""),
        "note": str(note or ""),
    }


def _ctx_counts(node: Any) -> Dict[str, int]:
    counts = {"Available": 0, "Partial": 0, "Gap": 0}

    def _walk(obj: Any) -> None:
        if isinstance(obj, dict):
            if (
                obj.get("status") in counts
                and "value" in obj
                and "source" in obj
                and "note" in obj
            ):
                counts[str(obj.get("status"))] += 1
                return
            for val in obj.values():
                _walk(val)
            return
        if isinstance(obj, list):
            for item in obj:
                _walk(item)

    _walk(node)
    return counts


def _read_jsonl(path: Path) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    try:
        if not path.exists():
            return rows
        with open(path, "r", encoding="utf-8") as f:
            for line in f:
                text = str(line or "").strip()
                if not text:
                    continue
                try:
                    obj = json.loads(text)
                except Exception:
                    continue
                if isinstance(obj, dict):
                    rows.append(obj)
    except Exception:
        return rows
    return rows


def _sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _ensure_dir(p: Path) -> Path:
    p.mkdir(parents=True, exist_ok=True)
    return p


def _reference_cfg_from_session(session: Dict[str, Any]) -> Dict[str, Any]:
    cfg = session.get("reference_audit")
    if not isinstance(cfg, dict):
        runner = session.get("daq_runner") if isinstance(session.get("daq_runner"), dict) else {}
        cfg = runner.get("reference_audit") if isinstance(runner.get("reference_audit"), dict) else {}
    out = None
    if mole_ftir_reference is not None:
        try:
            out = mole_ftir_reference.normalize_config(cfg if isinstance(cfg, dict) else {})
        except Exception:
            out = None
    if not isinstance(out, dict):
        out = dict(cfg if isinstance(cfg, dict) else {})
    out.setdefault("enabled", False)
    out.setdefault("provider", "MG2000_FTIR_PRN")
    out.setdefault("role", "AUDIT")
    out.setdefault("prn_path", "")
    out.setdefault("file_pattern", "*.prn")
    out.setdefault("recursive", False)
    out.setdefault("freshness_s", 120.0)
    out.setdefault("column_overrides", {})
    out.setdefault("units_overrides", {})
    if mole_ftir_offset_recommendations is not None:
        try:
            store = mole_ftir_offset_recommendations.load_store()
            default_use = bool(store.get("auto_apply_future_projects"))
            out["offset_recommendations"] = mole_ftir_offset_recommendations.normalize_project_policy(
                out.get("offset_recommendations"),
                default_use_global=default_use,
            )
        except Exception:
            out["offset_recommendations"] = dict(mole_ftir_offset_recommendations.DEFAULT_PROJECT_POLICY)
    else:
        out.setdefault("offset_recommendations", {
            "use_global": False,
            "promote_latest_for_future_projects": False,
        })
    return out


def _reference_cfg_seeded(cfg: Dict[str, Any]) -> bool:
    if not isinstance(cfg, dict):
        return False
    return bool(
        cfg.get("enabled")
        or cfg.get("prn_path")
        or cfg.get("column_overrides")
        or cfg.get("units_overrides")
    )


def _live_reference_snapshot(cfg: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    if (not _reference_cfg_seeded(cfg)) or mole_ftir_reference is None:
        return None
    try:
        return mole_ftir_reference.MG2000PrnIngestor(cfg).read_snapshot()
    except Exception:
        return None


def _snapshot_to_record(snapshot: Dict[str, Any], cfg: Dict[str, Any]) -> Dict[str, Any]:
    ts_iso = str(snapshot.get("ts_iso") or _now_iso())
    return {
        "ts_utc": ts_iso,
        "frame_ts_iso": ts_iso,
        "provider": str(snapshot.get("provider") or ""),
        "role": str(snapshot.get("role") or ""),
        "ok": bool(snapshot.get("ok")),
        "status": str(snapshot.get("status") or ""),
        "source_path": str(snapshot.get("source_path") or ""),
        "source_mtime_iso": str(snapshot.get("source_mtime_iso") or ""),
        "source_sha256": str(snapshot.get("source_sha256") or ""),
        "source_bytes": snapshot.get("source_bytes"),
        "age_s": snapshot.get("age_s"),
        "header_index": snapshot.get("header_index"),
        "row_index": snapshot.get("row_index"),
        "delimiter": str(snapshot.get("delimiter") or ""),
        "updated_value_count": snapshot.get("updated_value_count"),
        "error": str(snapshot.get("error") or ""),
        "notes": list(snapshot.get("notes") or []),
        "measurements": dict(snapshot.get("measurements") or {}),
        "primary_values": {},
        "deltas": {},
        "config": {
            "enabled": bool(cfg.get("enabled")),
            "provider": str(cfg.get("provider") or ""),
            "role": str(cfg.get("role") or ""),
            "prn_path": str(cfg.get("prn_path") or ""),
            "file_pattern": str(cfg.get("file_pattern") or ""),
            "recursive": bool(cfg.get("recursive")),
            "freshness_s": cfg.get("freshness_s"),
            "column_overrides": dict(cfg.get("column_overrides") or {}),
            "units_overrides": dict(cfg.get("units_overrides") or {}),
        },
    }


def _normalize_unit_text(unit: Any) -> str:
    s = str(unit or "").strip().upper().replace(" ", "")
    if not s:
        return ""
    aliases = {
        "PERCENT": "%",
        "PCT": "%",
        "VOL%": "%",
        "VOLUME%": "%",
        "PPMVD": "PPM",
        "PPMV": "PPM",
        "PPMWET": "PPM",
    }
    return aliases.get(s, s)


def _pollutant_prescriptions(session: Dict[str, Any]) -> Dict[str, Any]:
    poll = session.get("pollutants") if isinstance(session.get("pollutants"), dict) else {}
    pres = poll.get("prescriptions") if isinstance(poll.get("prescriptions"), dict) else {}
    return pres if isinstance(pres, dict) else {}


def _selected_pollutants(session: Dict[str, Any]) -> List[str]:
    poll = session.get("pollutants") if isinstance(session.get("pollutants"), dict) else {}
    selected = poll.get("selected") if isinstance(poll.get("selected"), list) else []
    out: List[str] = []
    for code in selected:
        canon = str(code or "").strip().upper()
        if canon:
            out.append(canon)
    return out


def _limits_by_pollutant(session: Dict[str, Any]) -> Dict[str, List[Dict[str, Any]]]:
    reg = session.get("regulatory") if isinstance(session.get("regulatory"), dict) else {}
    lims = reg.get("limits") if isinstance(reg.get("limits"), list) else []
    out: Dict[str, List[Dict[str, Any]]] = {}
    for row in lims:
        if not isinstance(row, dict):
            continue
        code = str(row.get("pollutant") or "").strip().upper()
        if not code:
            continue
        out.setdefault(code, []).append(row)
    return out


def _default_units_for_pollutant(code: str) -> str:
    return "%" if str(code or "").strip().upper() in ("O2", "CO2", "H2O", "RH") else "ppm"


def _co_pct_from_units(value: Any, units: Any) -> Optional[float]:
    val = _safe_float(value)
    if val is None:
        return None
    u = str(units or "").lower()
    if "ppb" in u:
        return float(val) / 10000000.0
    if "ppm" in u:
        return float(val) / 10000.0
    return (float(val) / 10000.0) if float(val) > 1.0 else float(val)


def _combustion_model_from_session(session: Dict[str, Any]) -> Dict[str, Any]:
    fuel = session.get("fuel") if isinstance(session.get("fuel"), dict) else {}
    analysis = fuel.get("analysis") if isinstance(fuel.get("analysis"), dict) else {}
    comb = analysis.get("combustion") or analysis.get("combustion_model")
    return comb if isinstance(comb, dict) else {}


def _fuel_analysis_snapshot(session: Dict[str, Any]) -> Dict[str, Any]:
    fuel = session.get("fuel") if isinstance(session.get("fuel"), dict) else {}
    analysis = fuel.get("analysis") if isinstance(fuel.get("analysis"), dict) else {}
    combustion = _combustion_model_from_session(session)
    site_conditions = session.get("site_conditions") if isinstance(session.get("site_conditions"), dict) else {}
    methodology_code = str(
        analysis.get("methodology_basis")
        or combustion.get("methodology_basis_code")
        or "ASTM_D3588_2017"
    ).strip() or "ASTM_D3588_2017"
    methodology_label = str(
        analysis.get("methodology_label")
        or combustion.get("methodology_basis_label")
        or ("ASTM D3588-98 (Reapproved 2017)" if methodology_code == "ASTM_D3588_2017" else methodology_code)
    ).strip()
    molecular_weight = combustion.get("MW_mix")
    relative_density = combustion.get("relative_density_air")
    if relative_density is None and _safe_float(molecular_weight) is not None:
        relative_density = float(molecular_weight) / 28.9625
    compressibility_basis = (
        combustion.get("compressibility_basis")
        or site_conditions.get("z_model")
        or "IDEAL"
    )
    compressibility_factor = combustion.get("compressibility_factor")
    if compressibility_factor is None:
        basis_upper = str(compressibility_basis or "").strip().upper()
        if basis_upper == "IDEAL":
            compressibility_factor = 1.0
        elif basis_upper == "FIXED_Z":
            compressibility_factor = _safe_float(site_conditions.get("fixed_z"))
    profile_id = analysis.get("profile_id") or combustion.get("profile_id")
    profile_label = analysis.get("profile_label") or combustion.get("profile_label") or profile_id
    default_source_citation = combustion.get("default_source_citation")
    default_source_url = combustion.get("default_source_url")
    default_source_display = combustion.get("default_source_display")
    if not default_source_display and default_source_citation:
        default_source_display = (
            f"{default_source_citation} | {default_source_url}"
            if default_source_url else str(default_source_citation)
        )
    default_source_note = combustion.get("default_source_note")
    traceability_note = combustion.get("traceability_note")
    if not traceability_note:
        trace_parts: List[str] = [
            f"Methodology: {methodology_label}",
        ]
        if profile_label:
            trace_parts.append(f"Profile: {profile_label}")
        if default_source_citation:
            trace_parts.append(f"Source: {default_source_citation}")
        else:
            trace_parts.append(
                "Source: not recorded in combustion snapshot"
            )
        if default_source_note:
            trace_parts.append(str(default_source_note))
        if default_source_url:
            trace_parts.append(str(default_source_url))
        traceability_note = " | ".join(str(x) for x in trace_parts if str(x).strip())
    return {
        "present": bool(analysis),
        "capture_enabled": bool(fuel.get("capture_fuel_analysis", True)),
        "captured": bool(analysis.get("captured")),
        "source": analysis.get("source"),
        "captured_by": analysis.get("captured_by"),
        "captured_iso": analysis.get("captured_iso"),
        "notes": analysis.get("notes"),
        "profile_id": profile_id,
        "profile_label": profile_label,
        "family": analysis.get("family"),
        "component_count": len(analysis.get("components") or []),
        "methodology_basis_code": methodology_code,
        "methodology_basis_label": methodology_label,
        "default_source_citation": default_source_citation,
        "default_source_url": default_source_url,
        "default_source_display": default_source_display,
        "default_source_note": default_source_note,
        "f_factor_basis": analysis.get("f_factor_basis"),
        "molecular_weight": molecular_weight,
        "relative_density_air": relative_density,
        "compressibility_basis": compressibility_basis,
        "compressibility_factor": compressibility_factor,
        "standard_conditions": combustion.get("std") or (site_conditions.get("standard_conditions") if isinstance(site_conditions.get("standard_conditions"), dict) else {}),
        "site_conditions": combustion.get("site") or (site_conditions.get("manual_entry") if isinstance(site_conditions.get("manual_entry"), dict) else {}),
        "hhv_btu_scf": combustion.get("HHV_Btu_scf"),
        "lhv_btu_scf": combustion.get("LHV_Btu_scf"),
        "hhv_btu_lb": combustion.get("HHV_Btu_lb"),
        "lhv_btu_lb": combustion.get("LHV_Btu_lb"),
        "hhv_btu_gal": combustion.get("HHV_Btu_gal"),
        "lhv_btu_gal": combustion.get("LHV_Btu_gal"),
        "f_factor_selected": combustion.get("F_selected") if combustion.get("F_selected") is not None else combustion.get("Fd_sel_dscf_per_MMBtu"),
        "f_factor_hhv": combustion.get("F_HHV") if combustion.get("F_HHV") is not None else combustion.get("Fd_HHV_dscf_per_MMBtu"),
        "f_factor_lhv": combustion.get("F_LHV") if combustion.get("F_LHV") is not None else combustion.get("Fd_LHV_dscf_per_MMBtu"),
        "traceability_note": traceability_note,
    }


def _pollutant_adjustment_elapsed_hours(session: Dict[str, Any]) -> float:
    daq = session.get("daq_runner") if isinstance(session.get("daq_runner"), dict) else {}
    runs = daq.get("runs") if isinstance(daq.get("runs"), list) else []
    if not runs:
        return 0.0
    idx = int(daq.get("active_run_index") or 0)
    idx = max(0, min(idx, len(runs) - 1))
    candidate_runs: List[Dict[str, Any]] = []
    if 0 <= idx < len(runs) and isinstance(runs[idx], dict):
        candidate_runs.append(runs[idx])
    for run in reversed(runs):
        if isinstance(run, dict) and run not in candidate_runs:
            candidate_runs.append(run)
    for run in candidate_runs:
        start_dt = _parse_iso_dt(run.get("start_iso"))
        if start_dt is None:
            continue
        end_dt = _parse_iso_dt(run.get("end_iso")) or datetime.now(timezone.utc)
        try:
            hours = max(0.0, (end_dt - start_dt).total_seconds() / 3600.0)
        except Exception:
            hours = 0.0
        return hours
    return 0.0


def _pollutant_adjustment_active_run_no(session: Dict[str, Any]) -> Optional[int]:
    daq = session.get("daq_runner") if isinstance(session.get("daq_runner"), dict) else {}
    runs = daq.get("runs") if isinstance(daq.get("runs"), list) else []
    if not runs:
        return None
    idx = int(daq.get("active_run_index") or 0)
    idx = max(0, min(idx, len(runs) - 1))
    candidate_runs: List[Dict[str, Any]] = []
    if 0 <= idx < len(runs) and isinstance(runs[idx], dict):
        candidate_runs.append(runs[idx])
    for run in reversed(runs):
        if isinstance(run, dict) and run not in candidate_runs:
            candidate_runs.append(run)
    for run in candidate_runs:
        try:
            run_no = run.get("run_no")
            if run_no not in (None, ""):
                return int(run_no)
        except Exception:
            continue
    return None


def _pollutant_adjustment_effective_after_run_no(spec: Any) -> Optional[int]:
    try:
        if not isinstance(spec, dict):
            return None
        value = spec.get("effective_after_run_no")
        return int(value) if value not in (None, "") else None
    except Exception:
        return None


def _pollutant_adjustment_expires_after_run_no(spec: Any) -> Optional[int]:
    try:
        if not isinstance(spec, dict):
            return None
        value = spec.get("expires_after_run_no")
        return int(value) if value not in (None, "") else None
    except Exception:
        return None


def _pollutant_adjustment_expires_after_postcal_run_no(spec: Any) -> Optional[int]:
    try:
        if not isinstance(spec, dict):
            return None
        value = spec.get("expires_after_postcal_run_no")
        return int(value) if value not in (None, "") else None
    except Exception:
        return None


def _normalize_pollutant_adjustment_scope(value: Any, source: Any = "") -> str:
    scope = str(value or "").strip().upper().replace("-", "_").replace(" ", "_")
    if scope in ("NEXT_RUN_ONLY", "UNTIL_NEXT_POSTCAL", "PERSISTENT"):
        return scope
    src = str(source or "").strip().upper().replace("-", "_").replace(" ", "_")
    if src == "POSTCAL_AUTO":
        return "UNTIL_NEXT_POSTCAL"
    return "PERSISTENT"


def _pollutant_adjustment_scope(spec: Any) -> str:
    if not isinstance(spec, dict):
        return "PERSISTENT"
    return _normalize_pollutant_adjustment_scope(spec.get("scope"), spec.get("source"))


def _pollutant_adjustment_scope_label(scope: Any) -> str:
    scope_u = _normalize_pollutant_adjustment_scope(scope)
    if scope_u == "NEXT_RUN_ONLY":
        return "next run only"
    if scope_u == "UNTIL_NEXT_POSTCAL":
        return "until next post-cal"
    return "persistent"


def _pollutant_adjustment_completed_postcal_runs(session: Dict[str, Any]) -> set[int]:
    completed: set[int] = set()
    try:
        daq = session.get("daq_runner") if isinstance(session.get("daq_runner"), dict) else {}
        ws = daq.get("worksteps") if isinstance(daq.get("worksteps"), dict) else {}
        postcal = ws.get("postcal") if isinstance(ws.get("postcal"), dict) else {}
        comp = postcal.get("completed_runs") if isinstance(postcal.get("completed_runs"), dict) else {}
        for key, value in comp.items():
            try:
                run_no = int(str(key))
            except Exception:
                continue
            if isinstance(value, dict):
                if bool(str(value.get("completed_iso") or "").strip()) or ("overall_pass" in value):
                    completed.add(run_no)
            else:
                completed.add(run_no)
    except Exception:
        return set()
    return completed


def _pollutant_adjustment_is_expired(session: Dict[str, Any], spec: Any) -> bool:
    if not isinstance(spec, dict):
        return False
    cur_run_no = _pollutant_adjustment_active_run_no(session)
    exp_run_no = _pollutant_adjustment_expires_after_run_no(spec)
    if exp_run_no is not None and cur_run_no is not None and cur_run_no > exp_run_no:
        return True
    exp_postcal_run_no = _pollutant_adjustment_expires_after_postcal_run_no(spec)
    if exp_postcal_run_no is not None and exp_postcal_run_no in _pollutant_adjustment_completed_postcal_runs(session):
        return True
    return False


def _pollutant_adjustment_scope_summary(spec: Any) -> str:
    scope_label = _pollutant_adjustment_scope_label(spec)
    exp_run_no = _pollutant_adjustment_expires_after_run_no(spec)
    exp_postcal_run_no = _pollutant_adjustment_expires_after_postcal_run_no(spec)
    if exp_run_no is not None:
        return f"{scope_label} (expires after Run {exp_run_no})"
    if exp_postcal_run_no is not None:
        return f"{scope_label} (expires after Post-Cal Run {exp_postcal_run_no})"
    return scope_label


def _pollutant_adjustment_spec_is_active(session: Dict[str, Any], spec: Any) -> bool:
    if _pollutant_adjustment_is_expired(session, spec):
        return False
    eff_run_no = _pollutant_adjustment_effective_after_run_no(spec)
    if eff_run_no is None:
        return True
    cur_run_no = _pollutant_adjustment_active_run_no(session)
    if cur_run_no is None:
        return False
    return bool(cur_run_no > eff_run_no)


def _pollutant_adjustment_state_snapshot(session: Dict[str, Any], spec: Any, *, global_enabled: bool = True) -> Dict[str, Any]:
    spec_dict = spec if isinstance(spec, dict) else {}
    spec_enabled = bool(spec_dict.get("enabled", True))
    lifecycle_status = str(spec_dict.get("lifecycle_status") or "").strip().upper()
    status_reason = str(spec_dict.get("status_reason") or "").strip()
    source = str(spec_dict.get("source") or "MANUAL").strip().upper() or "MANUAL"
    source_run_no = spec_dict.get("source_run_no")
    effective_after_run_no = _pollutant_adjustment_effective_after_run_no(spec_dict)
    scope = _pollutant_adjustment_scope(spec_dict)
    expired = _pollutant_adjustment_is_expired(session, spec_dict)
    active_now = bool(global_enabled and spec_enabled and (not expired) and _pollutant_adjustment_spec_is_active(session, spec_dict))
    if not global_enabled:
        state = "GLOBAL_OFF"
        effective_label = "global adjustments disabled"
    elif lifecycle_status in ("POLICY_SUSPENDED", "DISABLED_BY_POLICY", "DISABLED_AFTER_FAIL"):
        state = "SUSPENDED"
        effective_label = "disabled"
    elif expired:
        state = "EXPIRED"
        effective_label = "scope completed"
    elif not spec_enabled:
        state = "DISABLED"
        effective_label = "disabled"
    elif active_now:
        state = "ACTIVE"
        effective_label = "active now"
    elif effective_after_run_no is not None:
        state = "PENDING"
        effective_label = ("at next run start" if int(effective_after_run_no) <= 0 else f"after Run {effective_after_run_no}")
    else:
        state = "READY"
        effective_label = "manual / immediate"
    source_label = source if source_run_no in (None, "") else f"{source} (Run {source_run_no})"
    return {
        "state": state,
        "effective_label": effective_label,
        "active_now": bool(active_now),
        "source": source,
        "source_label": source_label,
        "source_run_no": source_run_no,
        "effective_after_run_no": effective_after_run_no,
        "scope": scope,
        "scope_label": _pollutant_adjustment_scope_label(scope),
        "scope_summary": _pollutant_adjustment_scope_summary(spec_dict),
        "expires_after_run_no": _pollutant_adjustment_expires_after_run_no(spec_dict),
        "expires_after_postcal_run_no": _pollutant_adjustment_expires_after_postcal_run_no(spec_dict),
        "lifecycle_status": lifecycle_status,
        "status_reason": status_reason,
        "enabled": bool(spec_enabled),
    }


def _units_for_pollutant(session: Dict[str, Any], pollutant: str) -> str:
    code = str(pollutant or "").strip().upper()
    pres = _pollutant_prescriptions(session).get(code)
    if isinstance(pres, dict):
        units = str(pres.get("expected_units") or pres.get("units") or "").strip()
        if units:
            return units
    return _default_units_for_pollutant(code)


def _pollutant_adjustment_history(session: Dict[str, Any]) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    if mole_postcal_policy is None:
        return [], []
    try:
        history_rows = list(mole_postcal_policy.collect_postcal_history(
            session,
            units_lookup=lambda sess_local, pollutant: _units_for_pollutant(sess_local, pollutant),
        ))
        review_rows = list(mole_postcal_policy.collect_postcal_reviews(session))
        return history_rows, review_rows
    except Exception:
        return [], []


def _pollutant_adjustment_summary(session: Dict[str, Any]) -> Dict[str, Any]:
    daq = session.get("daq_runner") if isinstance(session.get("daq_runner"), dict) else {}
    blk = daq.get("pollutant_adjustments") if isinstance(daq.get("pollutant_adjustments"), dict) else {}
    channels = blk.get("channels") if isinstance(blk.get("channels"), dict) else {}
    elapsed_hr = _pollutant_adjustment_elapsed_hours(session)
    active_run_no = _pollutant_adjustment_active_run_no(session)
    history_rows, review_rows = _pollutant_adjustment_history(session)
    rows: List[Dict[str, Any]] = []
    for pollutant in sorted(channels.keys()):
        spec = channels.get(pollutant) if isinstance(channels.get(pollutant), dict) else None
        if not isinstance(spec, dict):
            continue
        bias = _safe_float(spec.get("bias")) or 0.0
        drift_per_hr = _safe_float(spec.get("drift_per_hr")) or 0.0
        state = _pollutant_adjustment_state_snapshot(session, spec, global_enabled=bool(blk.get("enabled")))
        source = str(state.get("source") or "MANUAL")
        source_run_no = state.get("source_run_no")
        effective_after_run_no = state.get("effective_after_run_no")
        channel_enabled = bool(state.get("enabled"))
        if (not channel_enabled) and abs(float(bias)) < 1e-12 and abs(float(drift_per_hr)) < 1e-12:
            continue
        units = _units_for_pollutant(session, str(pollutant or ""))
        active_now = bool(state.get("active_now"))
        total_adjustment = (bias + (drift_per_hr * elapsed_hr)) if active_now else 0.0
        note_parts = []
        if str(spec.get("note") or "").strip():
            note_parts.append(str(spec.get("note") or "").strip())
        if str(state.get("status_reason") or "").strip():
            note_parts.append(str(state.get("status_reason") or "").strip())
        source_label = str(state.get("source_label") or source)
        rows.append({
            "pollutant": str(pollutant or "").strip().upper(),
            "enabled": channel_enabled,
            "state": str(state.get("state") or ""),
            "active_now": active_now,
            "source": source,
            "source_label": source_label,
            "source_run_no": source_run_no,
            "effective_after_run_no": effective_after_run_no,
            "effective_label": str(state.get("effective_label") or ""),
            "scope": str(state.get("scope") or ""),
            "scope_label": str(state.get("scope_label") or ""),
            "scope_summary": str(state.get("scope_summary") or ""),
            "expires_after_run_no": state.get("expires_after_run_no"),
            "expires_after_postcal_run_no": state.get("expires_after_postcal_run_no"),
            "lifecycle_status": str(state.get("lifecycle_status") or ""),
            "status_reason": str(state.get("status_reason") or ""),
            "units": units,
            "bias": bias,
            "drift_per_hr": drift_per_hr,
            "elapsed_hours": elapsed_hr,
            "total_adjustment": total_adjustment,
            "note": " | ".join(note_parts),
            "updated_by": str(spec.get("updated_by") or "").strip(),
            "updated_iso": str(spec.get("updated_iso") or "").strip(),
        })
    return {
        "enabled": bool(blk.get("enabled")),
        "drift_basis": str(blk.get("drift_basis") or "ACTIVE_RUN_HR"),
        "formula": str(blk.get("formula") or "adjusted = raw + bias + (drift_per_hr * elapsed_run_hr)"),
        "elapsed_hours": elapsed_hr,
        "active_run_no": active_run_no,
        "row_count": len(rows),
        "rows": rows,
        "history_row_count": len(history_rows),
        "history": history_rows,
        "decision_review_count": len(review_rows),
        "decision_reviews": review_rows,
        "latest_review": (review_rows[0] if review_rows else None),
        "note": (
            "Adjustments apply before dry normalization, O2 correction, and Method 19 calculations. "
            "If O2 is adjusted, the adjusted O2 channel is used as the EPA 3A / 7E correction denominator."
        ),
    }


def _spike_cfg_from_session(session: Dict[str, Any]) -> Dict[str, Any]:
    daq = session.get("daq_runner") if isinstance(session.get("daq_runner"), dict) else {}
    for candidate in (
        daq.get("spike_recovery"),
        session.get("spike_recovery"),
    ):
        if isinstance(candidate, dict):
            return dict(candidate)
    return {}


def _spike_state_from_session(session: Dict[str, Any]) -> Dict[str, Any]:
    daq = session.get("daq_runner") if isinstance(session.get("daq_runner"), dict) else {}
    ws = daq.get("worksteps") if isinstance(daq.get("worksteps"), dict) else {}
    state = ws.get("spike_recovery")
    return dict(state) if isinstance(state, dict) else {}


def _spike_recovery_snapshot(session: Dict[str, Any]) -> Dict[str, Any]:
    cfg = _spike_cfg_from_session(session)
    state = _spike_state_from_session(session)
    channels = state.get("channels") if isinstance(state.get("channels"), dict) else {}
    flow_result = state.get("flow_result") if isinstance(state.get("flow_result"), dict) else {}
    rows: List[Dict[str, Any]] = []
    mole_pass = 0
    mole_fail = 0
    ftir_pass = 0
    ftir_fail = 0
    for pollutant, rec in sorted(channels.items()):
        if not isinstance(rec, dict):
            continue
        mole_result = rec.get("mole_result") if isinstance(rec.get("mole_result"), dict) else {}
        ftir_result = rec.get("ftir_result") if isinstance(rec.get("ftir_result"), dict) else {}
        mole_status = str(mole_result.get("status") or "").strip().upper()
        ftir_status = str(ftir_result.get("status") or "").strip().upper()
        if mole_status == "PASS":
            mole_pass += 1
        elif mole_status == "FAIL":
            mole_fail += 1
        if ftir_status == "PASS":
            ftir_pass += 1
        elif ftir_status == "FAIL":
            ftir_fail += 1
        rows.append({
            "pollutant": pollutant,
            "units": str(rec.get("units") or ""),
            "spike_amount": rec.get("spike_amount"),
            "native_mole": rec.get("native_mole"),
            "spike_mole": rec.get("spike_mole"),
            "mole_recovery_pct": mole_result.get("recovery_pct"),
            "mole_status": mole_status,
            "native_ftir": rec.get("native_ftir"),
            "spike_ftir": rec.get("spike_ftir"),
            "ftir_recovery_pct": ftir_result.get("recovery_pct"),
            "ftir_status": ftir_status,
            "native_timestamp_iso": rec.get("native_timestamp_iso"),
            "spike_timestamp_iso": rec.get("spike_timestamp_iso"),
        })

    flow_status = str(flow_result.get("status") or "").strip().upper()
    return {
        "enabled": bool(cfg.get("enabled")),
        "mandatory": bool(cfg.get("mandatory")),
        "applies_to_mole": bool(cfg.get("applies_to_mole", True)),
        "applies_to_ftir": bool(cfg.get("applies_to_ftir")),
        "requires_flow_restriction": bool(cfg.get("requires_flow_restriction")),
        "flow_restriction_pct": cfg.get("flow_restriction_pct"),
        "criterion_summary": cfg.get("criterion_summary"),
        "flow_restriction_summary": cfg.get("flow_restriction_summary"),
        "method_standard": cfg.get("method_standard"),
        "method_label": cfg.get("method_label"),
        "method_basis": cfg.get("method_basis"),
        "active_phase": state.get("active_phase"),
        "sample_flow_value": state.get("sample_flow_value"),
        "sample_flow_units": state.get("sample_flow_units"),
        "sample_flow_source": state.get("sample_flow_source"),
        "sample_flow_override": state.get("sample_flow_override"),
        "spike_flow_value": state.get("spike_flow_value"),
        "target_spike_flow_value": state.get("target_spike_flow_value"),
        "flow_status": flow_status,
        "flow_reason": flow_result.get("reason"),
        "channel_count": len(rows),
        "mole_pass_count": mole_pass,
        "mole_fail_count": mole_fail,
        "ftir_pass_count": ftir_pass,
        "ftir_fail_count": ftir_fail,
        "channels": rows,
    }


def _side_by_side_snapshot(session: Dict[str, Any]) -> Dict[str, Any]:
    daq = session.get("daq_runner") if isinstance(session.get("daq_runner"), dict) else {}
    blk = daq.get("side_by_side") if isinstance(daq.get("side_by_side"), dict) else {}
    if not blk and isinstance(session.get("side_by_side"), dict):
        blk = session.get("side_by_side") or {}
    runs = daq.get("runs") if isinstance(daq.get("runs"), list) else []
    try:
        idx = int(daq.get("active_run_index") or 0)
    except Exception:
        idx = 0
    idx = max(0, min(idx, len(runs) - 1)) if runs else 0
    active_run = runs[idx] if runs and 0 <= idx < len(runs) and isinstance(runs[idx], dict) else {}
    try:
        active_run_no = int(active_run.get("run_no") or (idx + 1)) if active_run else None
    except Exception:
        active_run_no = None
    tm = session.get("test_matrix") if isinstance(session.get("test_matrix"), dict) else {}
    plan = tm.get("plan") if isinstance(tm.get("plan"), dict) else {}
    planned_min = None
    schedule_source = "SIDE_BY_SIDE"
    run_durations = plan.get("run_durations_min") if isinstance(plan.get("run_durations_min"), list) else []
    if active_run_no is not None and active_run_no > 0 and len(run_durations) >= active_run_no:
        planned_min = _safe_float(run_durations[active_run_no - 1])
        if planned_min is not None and planned_min > 0:
            schedule_source = "TEST_MATRIX_RUN"
    if planned_min is None:
        planned_min = _safe_float(plan.get("minutes_per_run"))
        if planned_min is not None and planned_min > 0:
            schedule_source = "TEST_MATRIX"
    if planned_min is None or not (planned_min > 0):
        planned_min = _safe_float(blk.get("bias_valve_interval_min")) or 20.0
        schedule_source = "SIDE_BY_SIDE"
    if schedule_source in ("TEST_MATRIX_RUN", "TEST_MATRIX"):
        schedule_label = f"Bias after each run ({_fmt_num(planned_min, 3, '')} min planned from Test Matrix)"
    else:
        schedule_label = f"Bias every {_fmt_num(planned_min, 3, '')} min"
    events = blk.get("events") if isinstance(blk.get("events"), list) else []
    clean_events = [dict(ev) for ev in events if isinstance(ev, dict)]
    latest_sync = blk.get("latest_sync_marker") if isinstance(blk.get("latest_sync_marker"), dict) else {}
    return {
        "enabled": bool(blk.get("enabled")),
        "bias_valve_interval_min": blk.get("bias_valve_interval_min"),
        "bias_schedule_minutes": planned_min,
        "bias_schedule_source": schedule_source,
        "bias_schedule_label": schedule_label,
        "bias_due": bool(blk.get("bias_due_after_run_no") not in (None, "")),
        "bias_due_after_run_no": blk.get("bias_due_after_run_no"),
        "bias_due_after_iso": blk.get("bias_due_after_iso"),
        "ambient_purge_required": bool(blk.get("ambient_purge_required", True)),
        "project_drift_required": bool(blk.get("project_drift_required", True)),
        "notes": blk.get("notes"),
        "event_count": len(clean_events),
        "latest_sync_marker": latest_sync,
        "events": clean_events,
    }


def _session_method_profile(session: Dict[str, Any]) -> Dict[str, Any]:
    if mole_spec_engine is None or not hasattr(mole_spec_engine, "get_method_profile"):
        return {}
    src = session.get("source") if isinstance(session.get("source"), dict) else {}
    fuel = session.get("fuel") if isinstance(session.get("fuel"), dict) else {}
    reg = session.get("regulatory") if isinstance(session.get("regulatory"), dict) else {}
    try:
        return mole_spec_engine.get_method_profile(
            src.get("source_type") or src.get("type") or src.get("category"),
            fuel.get("fuel_type") or fuel.get("fuel_button_code") or fuel.get("fuel_code") or fuel.get("code"),
            src.get("duty_type") or src.get("service_type") or src.get("prime_mover"),
            reg.get("regulation") or reg.get("subpart") or reg.get("rule"),
        )
    except Exception:
        return {}


def _latest_raw_sample_snapshot(raw_samples_path: Path) -> Dict[str, Any]:
    rows = _read_jsonl(raw_samples_path)
    latest: Dict[str, Any] = {}
    ts_latest = ""
    for row in rows:
        try:
            ch = str(row.get("channel_id") or "").strip()
            if not ch:
                continue
            val = _safe_float(row.get("value_eng"))
            if val is None:
                continue
            latest[ch] = float(val)
            ts = str(row.get("ts_utc") or "").strip()
            if ts:
                ts_latest = ts
        except Exception:
            continue
    return {"frame": latest, "ts_utc": ts_latest, "count": len(rows)}


def _fallback_bhp_from_session(session: Dict[str, Any]) -> Optional[float]:
    src = session.get("source") if isinstance(session.get("source"), dict) else {}
    ex_cfg = src.get("exhaust_flow") if isinstance(src.get("exhaust_flow"), dict) else {}
    mf_cfg = ex_cfg.get("manufacturer") if isinstance(ex_cfg.get("manufacturer"), dict) else {}
    for key in ("site_rated_power", "max_power", "bhp", "hp", "rated_hp", "site_rated_hp", "nameplate_hp", "rated_power_hp"):
        val = _safe_float(src.get(key))
        if val is not None:
            return float(val)
    for key in ("rated_power_hp", "rated_hp"):
        val = _safe_float(mf_cfg.get(key))
        if val is not None:
            return float(val)
    return None


def _build_regulatory_snapshot(
    session: Dict[str, Any],
    raw_samples_path: Path,
    analyzer_validity: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    if mole_spec_engine is None or not hasattr(mole_spec_engine, "evaluate_regulatory_output"):
        return {"status": "ENGINE_UNAVAILABLE", "rows": [], "per_pollutant": {}}

    snapshot = _latest_raw_sample_snapshot(raw_samples_path)
    frame = snapshot.get("frame") if isinstance(snapshot.get("frame"), dict) else {}
    if not frame:
        return {
            "status": "NO_DATA",
            "rows": [],
            "per_pollutant": {},
            "frame_ts_utc": snapshot.get("ts_utc"),
            "sample_count": snapshot.get("count", 0),
        }

    pres = _pollutant_prescriptions(session)
    limits_by = _limits_by_pollutant(session)
    selected = _selected_pollutants(session)
    if not selected:
        selected = sorted([str(k).strip().upper() for k in pres.keys() if str(k).strip()])

    profile = _session_method_profile(session)
    default_o2_ref = _safe_float(profile.get("o2_reference_pct"))
    if default_o2_ref is None:
        default_o2_ref = 15.0

    o2_meas_pct = _safe_float(frame.get("O2"))
    co2_meas_pct = _safe_float(frame.get("CO2"))
    co_units = str(((pres.get("CO") or {}).get("units") or (pres.get("CO") or {}).get("expected_units") or _default_units_for_pollutant("CO"))).strip()
    co_raw = _safe_float(frame.get("CO"))
    co_meas_pct = _co_pct_from_units(co_raw, co_units)

    qd_dscfh = _safe_float(frame.get("qd_dscfh_calc")) or _safe_float(frame.get("qd_dscfh")) or _safe_float(frame.get("exh_flow_dscfh"))
    heat_mmbtu_hr = _safe_float(frame.get("heat_input_R_mmbtu_hr")) or _safe_float(frame.get("heat_input_mmbtu_hr")) or _safe_float(frame.get("R_mmbtu_hr"))
    bhp = _safe_float(frame.get("hp_used")) or _safe_float(frame.get("bhp_used")) or _safe_float(frame.get("actual_hp")) or _fallback_bhp_from_session(session)

    rows: List[Dict[str, Any]] = []
    per_pollutant: Dict[str, Any] = {}
    for code in selected:
        spec = pres.get(code) if isinstance(pres.get(code), dict) else {}
        units = str(spec.get("units") or spec.get("expected_units") or _default_units_for_pollutant(code)).strip()
        eval_out = mole_spec_engine.evaluate_regulatory_output(
            code,
            {"raw": frame.get(code), "units": units},
            limits_by.get(code) or [],
            {
                "raw_is_wet": True,
                "apply_o2_correction": True,
                "default_o2_ref": default_o2_ref,
                "o2_meas_pct": o2_meas_pct,
                "co2_meas_pct": co2_meas_pct,
                "co_meas_pct": co_meas_pct,
                "combustion_model": _combustion_model_from_session(session),
                "qd_dscfh": qd_dscfh,
                "heat_mmbtu_hr": heat_mmbtu_hr,
                "bhp": bhp,
                "hours_per_year": ((session.get("regulatory") or {}).get("hours_per_year") if isinstance(session.get("regulatory"), dict) else None),
            },
        )
        conc = eval_out.get("concentration") if isinstance(eval_out.get("concentration"), dict) else {}
        mass = eval_out.get("mass") if isinstance(eval_out.get("mass"), dict) else {}
        views = eval_out.get("views") if isinstance(eval_out.get("views"), dict) else {}
        row = {
            "pollutant": code,
            "units": units,
            "raw": views.get("raw"),
            "dry": views.get("dry"),
            "corr": views.get("corrected"),
            "o2_ref": views.get("o2_ref"),
            "limit_value": conc.get("limit_value"),
            "limit_units": conc.get("limit_units"),
            "compare_value": conc.get("compare_value"),
            "compare_source": conc.get("compare_source"),
            "pct_limit": conc.get("pct_limit"),
            "status": conc.get("status"),
            "mass_value": mass.get("display_value"),
            "mass_units": mass.get("display_units"),
            "mass_limit_value": mass.get("limit_value"),
            "mass_limit_units": mass.get("limit_units"),
            "mass_pct_limit": mass.get("pct_limit"),
            "mass_status": mass.get("status"),
        }
        rows.append(row)
        per_pollutant[code] = row

    pass_count = sum(1 for row in rows if str(row.get("status") or "").upper() == "PASS")
    fail_count = sum(1 for row in rows if str(row.get("status") or "").upper() == "FAIL")
    blocked_count = 0
    validity_by = (analyzer_validity or {}).get("per_pollutant") if isinstance((analyzer_validity or {}).get("per_pollutant"), dict) else {}
    if validity_by:
        pass_count = 0
        fail_count = 0
        blocked_count = 0
        for row in rows:
            code = str(row.get("pollutant") or "").strip().upper()
            validity = validity_by.get(code) if isinstance(validity_by.get(code), dict) else {}
            gates = validity.get("gates") if isinstance(validity.get("gates"), dict) else {}
            state = str(validity.get("state") or "VALID").strip().upper() or "VALID"
            allow_regulatory = bool(gates.get("allow_regulatory", state == "VALID"))
            row["validity_state"] = state
            row["validity_reason"] = str(validity.get("reason") or "")
            row["regulatory_allowed"] = allow_regulatory
            if not allow_regulatory:
                row["status"] = state
                row["mass_status"] = state
                blocked_count += 1
            else:
                if str(row.get("status") or "").upper() == "PASS":
                    pass_count += 1
                elif str(row.get("status") or "").upper() == "FAIL":
                    fail_count += 1
    return {
        "status": "OK",
        "frame_ts_utc": snapshot.get("ts_utc"),
        "sample_count": snapshot.get("count", 0),
        "default_o2_ref": default_o2_ref,
        "profile_id": profile.get("profile_id"),
        "context": {
            "o2_meas_pct": o2_meas_pct,
            "co2_meas_pct": co2_meas_pct,
            "qd_dscfh": qd_dscfh,
            "heat_mmbtu_hr": heat_mmbtu_hr,
            "bhp": bhp,
        },
        "counts": {
            "pollutants": len(rows),
            "pass": pass_count,
            "fail": fail_count,
            "blocked": blocked_count,
        },
        "rows": rows,
        "per_pollutant": per_pollutant,
    }


def _worksteps_path_from_session(session: Dict[str, Any], cfg_path: Path, session_dir: Path) -> Path:
    daq = session.get("daq_runner") if isinstance(session.get("daq_runner"), dict) else {}
    ws = daq.get("worksteps") if isinstance(daq.get("worksteps"), dict) else {}
    explicit = str(ws.get("events_path") or "").strip()
    if explicit:
        try:
            path = Path(explicit).expanduser()
            if not path.is_absolute():
                path = (cfg_path.parent / path).resolve()
            else:
                path = path.resolve()
            if path.exists():
                return path
        except Exception:
            pass
    job = str(
        (session.get("project") or {}).get("job_id")
        or (session.get("project_session") or {}).get("job_id")
        or session.get("job_id")
        or "session"
    ).strip() or "session"
    filename = f"mole_worksteps_{job}.jsonl"
    candidates: List[Path] = [
        (session_dir / filename).resolve(),
        (cfg_path.parent / filename).resolve(),
    ]
    try:
        data_root = session_dir.parents[2]
        candidates.append((data_root / "logs" / filename).resolve())
    except Exception:
        pass
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0]


def _session_job_id(session: Dict[str, Any]) -> str:
    return str(
        ((session.get("project") or {}).get("job_id"))
        or ((session.get("project_session") or {}).get("job_id"))
        or session.get("job_id")
        or ""
    ).strip()


def _session_run_id(session: Dict[str, Any], session_dir: Path) -> str:
    return str(session.get("run_id") or session.get("session_id") or session_dir.name or "").strip()


def _parse_run_id_dt(text: Any, tz_hint: Optional[timezone] = None) -> Optional[datetime]:
    try:
        m = re.search(r"__(\d{8}_\d{6})$", str(text or "").strip())
        if not m:
            return None
        dt = datetime.strptime(m.group(1), "%Y%m%d_%H%M%S")
        tzinfo = tz_hint or datetime.now().astimezone().tzinfo or timezone.utc
        return dt.replace(tzinfo=tzinfo)
    except Exception:
        return None


def _discover_job_session_profiles(session_dir: Path, job_id: str) -> List[Dict[str, Any]]:
    entries: List[Dict[str, Any]] = []
    if not job_id:
        return entries
    try:
        data_root = session_dir.parents[2]
    except Exception:
        return entries

    seen: set[Tuple[str, str]] = set()
    for base in (data_root / "sessions", data_root / "training" / "sessions"):
        if not base.exists():
            continue
        for profile_path in base.rglob("session_profile.json"):
            profile = _read_json(profile_path)
            if not isinstance(profile, dict):
                continue
            profile_job_id = str(
                profile.get("site_id")
                or profile.get("job_id")
                or str(profile.get("run_id") or "").split("__")[0]
                or ""
            ).strip()
            if profile_job_id != job_id:
                continue
            profile_run_id = str(profile.get("run_id") or profile.get("session_id") or profile_path.parent.name or "").strip()
            created_dt = _parse_iso_dt(profile.get("created_at")) or _parse_run_id_dt(profile_run_id)
            if created_dt is None:
                continue
            key = (str(profile_path.parent.resolve()), profile_run_id)
            if key in seen:
                continue
            seen.add(key)
            entries.append({
                "run_id": profile_run_id,
                "job_id": profile_job_id,
                "created_dt": created_dt,
                "created_ts_iso": created_dt.isoformat(),
                "dir": str(profile_path.parent.resolve()),
                "source": str(profile_path),
            })
    entries.sort(key=lambda item: (item.get("created_dt"), str(item.get("run_id") or ""), str(item.get("dir") or "")))
    return entries


def _build_session_time_scope(session: Dict[str, Any], session_dir: Path) -> Dict[str, Any]:
    run_id = _session_run_id(session, session_dir)
    job_id = _session_job_id(session)
    profile = _read_json(session_dir / "session_profile.json") or {}
    meta = session.get("meta") if isinstance(session.get("meta"), dict) else {}

    profile_dt = _parse_iso_dt(profile.get("created_at"))
    applied_dt = _parse_iso_dt(meta.get("applied_iso"))
    tz_hint = None
    for dt in (profile_dt, applied_dt):
        if dt is not None and dt.tzinfo is not None:
            tz_hint = dt.tzinfo
            break
    run_id_dt = _parse_run_id_dt(run_id, tz_hint=tz_hint)

    start_dt = None
    start_source = ""
    for label, candidate in (
        ("session_profile.created_at", profile_dt),
        ("session.meta.applied_iso", applied_dt),
        ("run_id suffix", run_id_dt),
    ):
        if candidate is not None:
            start_dt = candidate
            start_source = label
            break

    siblings = _discover_job_session_profiles(session_dir, job_id)
    session_dir_resolved = str(session_dir.resolve())
    current_index = None
    for idx, entry in enumerate(siblings):
        if str(entry.get("dir") or "") == session_dir_resolved:
            current_index = idx
            if start_dt is None:
                start_dt = entry.get("created_dt")
                start_source = "session profile scan"
            break
    if current_index is None and start_dt is not None:
        for idx, entry in enumerate(siblings):
            created_dt = entry.get("created_dt")
            if isinstance(created_dt, datetime) and created_dt == start_dt and str(entry.get("run_id") or "") == run_id:
                current_index = idx
                break

    next_entry = None
    if current_index is not None:
        if current_index + 1 < len(siblings):
            next_entry = siblings[current_index + 1]
    elif start_dt is not None:
        for entry in siblings:
            created_dt = entry.get("created_dt")
            if isinstance(created_dt, datetime) and created_dt > start_dt:
                next_entry = entry
                break

    end_dt = next_entry.get("created_dt") if isinstance(next_entry, dict) else None
    status = "BOUNDED" if (start_dt is not None and end_dt is not None) else ("START_ONLY" if start_dt is not None else "UNSCOPED")
    return {
        "status": status,
        "job_id": job_id,
        "run_id": run_id,
        "start_ts_iso": start_dt.isoformat() if isinstance(start_dt, datetime) else None,
        "end_ts_iso": end_dt.isoformat() if isinstance(end_dt, datetime) else None,
        "start_source": start_source or None,
        "end_source": ("next session profile" if end_dt is not None else None),
        "next_run_id": str(next_entry.get("run_id") or "") if isinstance(next_entry, dict) else None,
        "sibling_session_count": len(siblings),
    }


def _load_scoped_worksteps(
    session: Dict[str, Any],
    *,
    cfg_path: Path,
    session_dir: Path,
) -> Dict[str, Any]:
    worksteps_path = _worksteps_path_from_session(session, cfg_path, session_dir)
    all_rows = _read_jsonl(worksteps_path)
    session_scope = _build_session_time_scope(session, session_dir)

    try:
        path_scope = "SESSION_LOCAL" if worksteps_path.resolve().is_relative_to(session_dir.resolve()) else "EXTERNAL"
    except Exception:
        path_scope = "EXTERNAL"

    start_dt = _parse_iso_dt(session_scope.get("start_ts_iso"))
    end_dt = _parse_iso_dt(session_scope.get("end_ts_iso"))
    scoped_rows: List[Dict[str, Any]] = []
    excluded_before = 0
    excluded_after = 0
    excluded_missing_ts = 0

    if path_scope == "SESSION_LOCAL":
        scoped_rows = list(all_rows)
        scope_note = "Session-local worksteps file used directly."
    elif start_dt is None:
        scope_note = "External worksteps file has no reliable session window; report aggregation is suppressed."
    else:
        for row in all_rows:
            evt_dt = _parse_iso_dt((row.get("ts_iso") if isinstance(row, dict) else None) or (row.get("timestamp") if isinstance(row, dict) else None))
            if evt_dt is None:
                excluded_missing_ts += 1
                continue
            if evt_dt < start_dt:
                excluded_before += 1
                continue
            if end_dt is not None and evt_dt >= end_dt:
                excluded_after += 1
                continue
            scoped_rows.append(row)
        if end_dt is not None:
            scope_note = "External worksteps file filtered to the current session window."
        else:
            scope_note = "External worksteps file filtered using the current session start; no upper bound was available."

    return {
        "path": worksteps_path,
        "path_scope": path_scope,
        "session_scope": session_scope,
        "scope_note": scope_note,
        "rows": scoped_rows,
        "total_rows": all_rows,
        "excluded_counts": {
            "before_window": excluded_before,
            "after_window": excluded_after,
            "missing_ts": excluded_missing_ts,
        },
    }


def _path_summary(path: Path, *, count: Optional[int] = None, extra: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    info: Dict[str, Any] = {
        "path": str(path),
        "exists": bool(path.exists()),
    }
    if count is not None:
        info["count"] = int(count)
    if path.exists():
        try:
            info["bytes"] = int(path.stat().st_size)
        except Exception:
            pass
        try:
            info["sha256"] = _sha256_file(path)
        except Exception:
            pass
    if isinstance(extra, dict):
        info.update(extra)
    return info


def _flatten_spec_engine_shadow_row(row: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    if not isinstance(row, dict):
        return None
    event = str(row.get("event") or "").strip().upper()
    if event not in ("SPEC_ENGINE_SHADOW", "SPEC_ENGINE_SHADOW_MISMATCH"):
        return None
    payload = row.get("payload") if isinstance(row.get("payload"), dict) else {}
    result = payload.get("result") if isinstance(payload.get("result"), dict) else {}
    comparison = payload.get("comparison") if isinstance(payload.get("comparison"), dict) else {}
    return {
        "ts_iso": str(row.get("ts_iso") or payload.get("timestamp") or ""),
        "event": event,
        "step": str(payload.get("step") or result.get("step") or ""),
        "channel": str(payload.get("channel") or result.get("details", {}).get("channel") or ""),
        "formula_version": str(payload.get("formula_version") or ""),
        "status": str(result.get("status") or ""),
        "pass": result.get("pass"),
        "stable": result.get("stable"),
        "within_tolerance": result.get("within_tolerance"),
        "avg": result.get("avg"),
        "std": result.get("std"),
        "n": result.get("n"),
        "target": result.get("target"),
        "tolerance": result.get("tolerance"),
        "basis": result.get("basis"),
        "comparison_value": result.get("comparison_value"),
        "recovery_pct": result.get("recovery_pct"),
        "legacy_pass": comparison.get("legacy_pass"),
        "shadow_pass": comparison.get("shadow_pass"),
        "match": comparison.get("match"),
        "reason": str(result.get("reason") or ""),
    }


def _static_artifact_state(session: Dict[str, Any]) -> Dict[str, List[Dict[str, Any]]]:
    daq = session.get("daq_runner") if isinstance(session.get("daq_runner"), dict) else {}
    for blk in (
        daq.get("static_artifacts"),
        session.get("static_artifacts"),
    ):
        if isinstance(blk, dict):
            out: Dict[str, List[Dict[str, Any]]] = {}
            for key, value in blk.items():
                if isinstance(value, list):
                    out[str(key)] = [dict(item) for item in value if isinstance(item, dict)]
            return out
    return {}


def _build_evidence_bundle(
    session: Dict[str, Any],
    *,
    cfg_path: Path,
    session_dir: Path,
    raw_samples_path: Path,
    raw_events_path: Path,
    health_path: Path,
    reference_path: Path,
    raw_count: int,
    first_ts: Optional[str],
    last_ts: Optional[str],
    health_counts: Dict[str, int],
) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    scoped_worksteps = _load_scoped_worksteps(session, cfg_path=cfg_path, session_dir=session_dir)
    worksteps_path = scoped_worksteps.get("path")
    workstep_rows = scoped_worksteps.get("rows") if isinstance(scoped_worksteps.get("rows"), list) else []
    all_workstep_rows = scoped_worksteps.get("total_rows") if isinstance(scoped_worksteps.get("total_rows"), list) else []
    raw_event_rows = _read_jsonl(raw_events_path)
    reference_rows = _read_jsonl(reference_path)

    workstep_event_counts: Dict[str, int] = {}
    all_workstep_event_counts: Dict[str, int] = {}
    raw_event_counts: Dict[str, int] = {}
    step_eval_rows: List[Dict[str, Any]] = []
    latest_shadow: Optional[Dict[str, Any]] = None
    latest_mismatch: Optional[Dict[str, Any]] = None

    for row in all_workstep_rows:
        event = str(row.get("event") or "").strip().upper() or "(UNKNOWN)"
        all_workstep_event_counts[event] = all_workstep_event_counts.get(event, 0) + 1

    for row in workstep_rows:
        event = str(row.get("event") or "").strip().upper() or "(UNKNOWN)"
        workstep_event_counts[event] = workstep_event_counts.get(event, 0) + 1
        flat = _flatten_spec_engine_shadow_row(row)
        if not isinstance(flat, dict):
            continue
        step_eval_rows.append(flat)
        latest_shadow = flat
        if flat.get("match") is False:
            latest_mismatch = flat

    for row in raw_event_rows:
        event = str(row.get("event") or "").strip().upper() or "(UNKNOWN)"
        raw_event_counts[event] = raw_event_counts.get(event, 0) + 1

    static_state = _static_artifact_state(session)
    active_artifacts: List[Dict[str, Any]] = []
    for category, entries in static_state.items():
        for entry in entries:
            active_artifacts.append({
                "category": str(category),
                "label": str(entry.get("label") or ""),
                "stored_rel_path": str(entry.get("stored_rel_path") or ""),
                "stored_path": str(entry.get("stored_path") or ""),
                "bytes": entry.get("bytes"),
                "sha256": str(entry.get("sha256") or ""),
                "uploaded_iso": str(entry.get("uploaded_iso") or ""),
            })

    channels = sorted({str(row.get("channel") or "").strip().upper() for row in step_eval_rows if str(row.get("channel") or "").strip()})
    steps = sorted({str(row.get("step") or "").strip().upper() for row in step_eval_rows if str(row.get("step") or "").strip()})
    mismatch_count = sum(1 for row in step_eval_rows if row.get("match") is False)
    reference_ok_count = 0
    for row in reference_rows:
        try:
            if bool(row.get("ok")):
                reference_ok_count += 1
        except Exception:
            continue

    analyzer_validity = _build_analyzer_validity_summary(session, workstep_rows)

    bundle = {
        "generated_iso": _now_iso(),
        "formula_version": getattr(mole_spec_engine, "SPEC_ENGINE_FORMULA_VERSION", "spec_engine_v1"),
        "session": {
            "job_id": str(((session.get("project") or {}).get("job_id")) or ((session.get("project_session") or {}).get("job_id")) or session.get("job_id") or ""),
            "run_id": str(session.get("run_id") or ""),
            "session_dir": str(session_dir),
        },
        "session_scope": dict(scoped_worksteps.get("session_scope") or {}),
        "sources": {
            "raw_samples": _path_summary(raw_samples_path, count=raw_count, extra={"first_ts": first_ts, "last_ts": last_ts}),
            "raw_events": _path_summary(raw_events_path, count=len(raw_event_rows)),
            "health_states": _path_summary(health_path, count=sum(int(v or 0) for v in (health_counts or {}).values()), extra={"counts": dict(health_counts or {})}),
            "reference_audit": _path_summary(reference_path, count=len(reference_rows), extra={"ok_count": reference_ok_count}),
            "worksteps": _path_summary(
                Path(str(worksteps_path)),
                count=len(workstep_rows),
                extra={
                    "event_counts": workstep_event_counts,
                    "total_count": len(all_workstep_rows),
                    "total_event_counts": all_workstep_event_counts,
                    "filtered_out_count": max(0, len(all_workstep_rows) - len(workstep_rows)),
                    "excluded_counts": dict(scoped_worksteps.get("excluded_counts") or {}),
                    "path_scope": scoped_worksteps.get("path_scope"),
                    "scope_note": scoped_worksteps.get("scope_note"),
                    "session_scope": dict(scoped_worksteps.get("session_scope") or {}),
                },
            ),
        },
        "counts": {
            "raw_samples": raw_count,
            "raw_events": len(raw_event_rows),
            "worksteps": len(workstep_rows),
            "step_eval_shadow": len(step_eval_rows),
            "step_eval_mismatch": mismatch_count,
            "reference_audit_records": len(reference_rows),
            "static_artifact_active": len(active_artifacts),
        },
        "raw_events": {
            "event_counts": raw_event_counts,
            "latest_event": raw_event_rows[-1] if raw_event_rows else None,
        },
        "spec_engine_shadow": {
            "total": len(step_eval_rows),
            "mismatches": mismatch_count,
            "parity_ok": mismatch_count == 0,
            "channels": channels,
            "steps": steps,
            "latest": latest_shadow,
            "latest_mismatch": latest_mismatch,
        },
        "static_artifacts": {
            "active_categories": sorted(static_state.keys()),
            "active_count": len(active_artifacts),
            "add_events": raw_event_counts.get("STATIC_ARTIFACT_ADDED", 0),
            "remove_events": raw_event_counts.get("STATIC_ARTIFACT_REMOVED", 0),
            "files": active_artifacts,
        },
        "reference_audit": {
            "records": len(reference_rows),
            "ok_records": reference_ok_count,
        },
        "analyzer_validity": analyzer_validity,
    }
    return bundle, step_eval_rows


def _workstep_payload_pass(payload: Any) -> Optional[bool]:
    if not isinstance(payload, dict):
        return None
    if "pass" in payload:
        value = payload.get("pass")
        if value is None:
            return None
        try:
            return bool(value)
        except Exception:
            return None
    status = str(payload.get("status") or "").strip().upper()
    if status in ("OK", "PASS", "GOOD", "VALID"):
        return True
    if status in ("FAIL", "INVALID"):
        return False
    return None


def _build_analyzer_validity_summary(session: Dict[str, Any], workstep_rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    if mole_spec_engine is None or not hasattr(mole_spec_engine, "derive_analyzer_validity"):
        return {"status": "ENGINE_UNAVAILABLE", "per_pollutant": {}, "counts": {}}

    selected = _selected_pollutants(session)
    zero_by: Dict[str, Dict[str, Any]] = {}
    span_by: Dict[str, Dict[str, Any]] = {}
    post_zero_by: Dict[str, Dict[str, Any]] = {}
    post_span_by: Dict[str, Dict[str, Any]] = {}
    latest_postcal_valid: Dict[str, Any] = {}
    latest_postcal_run = ""

    for row in workstep_rows:
        event = str(row.get("event") or "").strip().upper()
        payload = row.get("payload") if isinstance(row.get("payload"), dict) else {}
        code = str(payload.get("pollutant") or payload.get("channel") or "").strip().upper()
        if event == "CAPTURE_ZERO" and code:
            zero_by[code] = payload
        elif event == "CAPTURE_SPAN" and code:
            span_by[code] = payload
        elif event == "CAPTURE_POST_ZERO" and code:
            post_zero_by[code] = payload
        elif event == "CAPTURE_POST_SPAN" and code:
            post_span_by[code] = payload
        elif event == "POSTCAL_COMPLETE":
            latest_postcal_run = str(payload.get("run_no") or latest_postcal_run or "")
            channel_valid = payload.get("channel_valid")
            if isinstance(channel_valid, dict):
                latest_postcal_valid = {
                    str(k or "").strip().upper(): v
                    for k, v in channel_valid.items()
                    if str(k or "").strip()
                }

    codes = set(selected)
    codes.update(zero_by.keys())
    codes.update(span_by.keys())
    codes.update(post_zero_by.keys())
    codes.update(post_span_by.keys())
    codes.update(latest_postcal_valid.keys())

    per_pollutant: Dict[str, Any] = {}
    counts = {
        "total": 0,
        "valid": 0,
        "degraded": 0,
        "invalid": 0,
        "regulatory_blocked": 0,
    }
    for code in sorted(codes):
        zero_rec = zero_by.get(code)
        span_rec = span_by.get(code)
        drift_rec = None
        latched_invalid = bool(latest_postcal_valid.get(code) is False)
        if code in latest_postcal_valid:
            drift_rec = {
                "pass": bool(latest_postcal_valid.get(code)),
                "status": "PASS" if bool(latest_postcal_valid.get(code)) else "FAIL",
            }
        else:
            post_zero_pass = _workstep_payload_pass(post_zero_by.get(code))
            post_span_pass = _workstep_payload_pass(post_span_by.get(code))
            if post_zero_pass is not None or post_span_pass is not None:
                if post_zero_pass is None or post_span_pass is None:
                    drift_rec = {"pass": None, "status": "NO_DATA"}
                else:
                    drift_ok = bool(post_zero_pass and post_span_pass)
                    drift_rec = {"pass": drift_ok, "status": "PASS" if drift_ok else "FAIL"}
        validity = mole_spec_engine.derive_analyzer_validity(
            zero_rec,
            span_rec,
            drift_rec,
            latched_invalid=latched_invalid,
            invalidate_on_drift=True,
        )
        state = str((validity or {}).get("state") or "VALID").strip().upper() or "VALID"
        gates = validity.get("gates") if isinstance(validity.get("gates"), dict) else {}
        per_pollutant[code] = {
            "state": state,
            "reason": str(validity.get("reason") or ""),
            "scores": dict(validity.get("scores") or {}),
            "gates": dict(gates),
            "latched_invalid": bool(validity.get("latched_invalid")),
            "latest_postcal_run": latest_postcal_run or None,
            "sources": {
                "zero_capture": bool(code in zero_by),
                "span_capture": bool(code in span_by),
                "post_zero_capture": bool(code in post_zero_by),
                "post_span_capture": bool(code in post_span_by),
                "postcal_channel_valid": latest_postcal_valid.get(code),
            },
        }
        counts["total"] += 1
        if state == "VALID":
            counts["valid"] += 1
        elif state == "DEGRADED":
            counts["degraded"] += 1
        else:
            counts["invalid"] += 1
        if not bool(gates.get("allow_regulatory", state == "VALID")):
            counts["regulatory_blocked"] += 1

    status = "NO_DATA"
    if per_pollutant:
        status = "OK"
    return {
        "status": status,
        "latest_postcal_run": latest_postcal_run or None,
        "counts": counts,
        "per_pollutant": per_pollutant,
    }


def _reference_acceptance_cfg(reference_cfg: Dict[str, Any]) -> Dict[str, Any]:
    blk = reference_cfg.get("acceptance") if isinstance(reference_cfg.get("acceptance"), dict) else {}
    by_pollutant = blk.get("by_pollutant")
    if not isinstance(by_pollutant, dict):
        by_pollutant = blk.get("pollutants")
    if not isinstance(by_pollutant, dict):
        by_pollutant = {}

    per_pollutant: Dict[str, str] = {}
    for code, spec in by_pollutant.items():
        canon = str(code or "").strip().upper()
        if not canon:
            continue
        if isinstance(spec, dict):
            expr = str(spec.get("tolerance_expr") or spec.get("expr") or "").strip()
        else:
            expr = str(spec or "").strip()
        if expr:
            per_pollutant[canon] = expr

    return {
        "enabled": bool(blk.get("enabled", True)),
        "require_reference_status_ok": bool(blk.get("require_reference_status_ok", True)),
        "default_tolerance_expr": str(blk.get("default_tolerance_expr") or blk.get("tolerance_expr") or "").strip(),
        "per_pollutant": per_pollutant,
    }


def _span_basis_for_pollutant(session: Dict[str, Any], pollutant: str) -> Dict[str, Any]:
    code = str(pollutant or "").strip().upper()
    pres = _pollutant_prescriptions(session).get(code)
    if not isinstance(pres, dict):
        pres = {}

    candidates: List[Tuple[str, Any]] = [
        (f"pollutants.prescriptions[{code}].span_target", pres.get("span_target")),
        (f"pollutants.prescriptions[{code}].span_cylinder_id", pres.get("span_cylinder_id")),
        (f"pollutants.prescriptions[{code}].expected_max", pres.get("expected_max")),
    ]

    resolved = (((session.get("qaqc") or {}).get("resolved") or {}).get("resolved_by_analyte") or {})
    if isinstance(resolved, dict):
        r = resolved.get(code)
        if isinstance(r, dict):
            candidates.append((f"qaqc.resolved.resolved_by_analyte[{code}].span_suggestion", r.get("span_suggestion")))

    for source, raw in candidates:
        val = _safe_float(raw)
        if val is not None and val > 0:
            return {
                "value": float(val),
                "source": source,
                "units": str(pres.get("expected_units") or ""),
            }

    return {
        "value": None,
        "source": "",
        "units": str(pres.get("expected_units") or ""),
    }


def _parse_tolerance_expr(expr: str, units_hint: str = "") -> Optional[Dict[str, Any]]:
    raw = str(expr or "").strip()
    if not raw:
        return None

    up = raw.upper()
    pct = re.search(r"([-+]?\d*\.?\d+)\s*%\s*SPAN\b", up)
    if pct:
        return {
            "type": "PERCENT_SPAN",
            "expr": raw,
            "value": float(pct.group(1)),
            "units": "%SPAN",
        }

    abs_match = re.search(r"([-+]?\d*\.?\d+)", raw)
    if not abs_match:
        return None

    num = float(abs_match.group(1))
    unit_text = raw[abs_match.end():].strip() or str(units_hint or "").strip()
    return {
        "type": "ABSOLUTE",
        "expr": raw,
        "value": num,
        "units": unit_text,
    }


def _resolve_acceptance_basis(session: Dict[str, Any], reference_cfg: Dict[str, Any], pollutant: str) -> Dict[str, Any]:
    code = str(pollutant or "").strip().upper()
    acceptance_cfg = _reference_acceptance_cfg(reference_cfg)
    pres = _pollutant_prescriptions(session).get(code)
    if not isinstance(pres, dict):
        pres = {}

    expected_units = str(pres.get("expected_units") or "")
    span_basis = _span_basis_for_pollutant(session, code)

    explicit_expr = acceptance_cfg["per_pollutant"].get(code)
    if explicit_expr:
        parsed = _parse_tolerance_expr(explicit_expr, units_hint=expected_units)
        out = {
            "enabled": bool(acceptance_cfg.get("enabled")),
            "expected_units": expected_units,
            "span_basis_value": span_basis.get("value"),
            "span_basis_source": span_basis.get("source"),
            "selection_strategy": "reference_audit.acceptance.by_pollutant",
            "tolerance_expr": explicit_expr,
            "tolerance_source": f"reference_audit.acceptance.by_pollutant[{code}]",
            "tolerance_type": parsed.get("type") if isinstance(parsed, dict) else None,
            "tolerance_units": parsed.get("units") if isinstance(parsed, dict) else expected_units,
            "tolerance_abs": None,
        }
        if isinstance(parsed, dict):
            if parsed.get("type") == "PERCENT_SPAN":
                span_val = _safe_float(span_basis.get("value"))
                if span_val is not None:
                    out["tolerance_abs"] = (float(parsed["value"]) / 100.0) * float(span_val)
            elif parsed.get("type") == "ABSOLUTE":
                out["tolerance_abs"] = float(parsed["value"])
        return out

    default_expr = acceptance_cfg.get("default_tolerance_expr")
    if default_expr:
        parsed = _parse_tolerance_expr(str(default_expr), units_hint=expected_units)
        out = {
            "enabled": bool(acceptance_cfg.get("enabled")),
            "expected_units": expected_units,
            "span_basis_value": span_basis.get("value"),
            "span_basis_source": span_basis.get("source"),
            "selection_strategy": "reference_audit.acceptance.default_tolerance_expr",
            "tolerance_expr": str(default_expr),
            "tolerance_source": "reference_audit.acceptance.default_tolerance_expr",
            "tolerance_type": parsed.get("type") if isinstance(parsed, dict) else None,
            "tolerance_units": parsed.get("units") if isinstance(parsed, dict) else expected_units,
            "tolerance_abs": None,
        }
        if isinstance(parsed, dict):
            if parsed.get("type") == "PERCENT_SPAN":
                span_val = _safe_float(span_basis.get("value"))
                if span_val is not None:
                    out["tolerance_abs"] = (float(parsed["value"]) / 100.0) * float(span_val)
            elif parsed.get("type") == "ABSOLUTE":
                out["tolerance_abs"] = float(parsed["value"])
        return out

    tm = session.get("test_matrix") if isinstance(session.get("test_matrix"), dict) else {}
    steps = tm.get("steps") if isinstance(tm.get("steps"), list) else []
    step_type_rank = {"CAL_ERROR": 0, "SPAN": 1, "ZERO": 2, "DRIFT": 3}
    candidates: List[Dict[str, Any]] = []
    for step in steps:
        if not isinstance(step, dict):
            continue
        analyte = str(step.get("analyte") or "").strip().upper()
        if analyte != code:
            continue
        expr = str(step.get("tol_override") or step.get("tol_default") or "").strip()
        if not expr:
            continue
        parsed = _parse_tolerance_expr(expr, units_hint=expected_units)
        if not isinstance(parsed, dict):
            continue
        tol_abs = None
        if parsed.get("type") == "PERCENT_SPAN":
            span_val = _safe_float(span_basis.get("value"))
            if span_val is not None:
                tol_abs = (float(parsed["value"]) / 100.0) * float(span_val)
        elif parsed.get("type") == "ABSOLUTE":
            tol_abs = float(parsed["value"])
        candidates.append({
            "expr": expr,
            "source": f"test_matrix.steps[{step.get('seq')}].{'tol_override' if str(step.get('tol_override') or '').strip() else 'tol_default'}",
            "type": parsed.get("type"),
            "units": parsed.get("units") or expected_units,
            "tolerance_abs": tol_abs,
            "step_type": str(step.get("type") or ""),
            "step_label": str(step.get("label") or ""),
            "step_rank": step_type_rank.get(str(step.get("type") or "").strip().upper(), 99),
        })

    if candidates:
        candidates.sort(
            key=lambda c: (
                1 if c.get("tolerance_abs") is None else 0,
                float(c.get("tolerance_abs")) if c.get("tolerance_abs") is not None else float("inf"),
                int(c.get("step_rank") or 99),
                str(c.get("expr") or ""),
            )
        )
        chosen = candidates[0]
        return {
            "enabled": bool(acceptance_cfg.get("enabled")),
            "expected_units": expected_units,
            "span_basis_value": span_basis.get("value"),
            "span_basis_source": span_basis.get("source"),
            "selection_strategy": "strictest_test_matrix_tolerance",
            "tolerance_expr": chosen.get("expr"),
            "tolerance_source": chosen.get("source"),
            "tolerance_type": chosen.get("type"),
            "tolerance_units": chosen.get("units") or expected_units,
            "tolerance_abs": chosen.get("tolerance_abs"),
            "tolerance_step_type": chosen.get("step_type"),
            "tolerance_step_label": chosen.get("step_label"),
        }

    return {
        "enabled": bool(acceptance_cfg.get("enabled")),
        "expected_units": expected_units,
        "span_basis_value": span_basis.get("value"),
        "span_basis_source": span_basis.get("source"),
        "selection_strategy": "none",
        "tolerance_expr": "",
        "tolerance_source": "",
        "tolerance_type": None,
        "tolerance_units": expected_units,
        "tolerance_abs": None,
    }


def _evaluate_reference_acceptance(row: Dict[str, Any], basis: Dict[str, Any], require_reference_status_ok: bool) -> Dict[str, Any]:
    status = str(row.get("status") or "").strip().upper()
    primary_value = _safe_float(row.get("primary_value"))
    reference_value = _safe_float(row.get("reference_value"))
    delta_value = _safe_float(row.get("delta"))
    delta_abs = abs(float(delta_value)) if delta_value is not None else None

    expected_units = str(basis.get("expected_units") or "")
    ref_units = str(row.get("units") or "")
    expected_norm = _normalize_unit_text(expected_units)
    ref_norm = _normalize_unit_text(ref_units)
    units_match = None
    if expected_norm and ref_norm:
        units_match = (expected_norm == ref_norm)

    tolerance_abs = _safe_float(basis.get("tolerance_abs"))
    span_basis_value = _safe_float(basis.get("span_basis_value"))
    delta_pct_span = None
    if delta_abs is not None and span_basis_value not in (None, 0):
        try:
            delta_pct_span = (float(delta_abs) / float(span_basis_value)) * 100.0
        except Exception:
            delta_pct_span = None

    acceptance_status = "NOT_EVALUATED"
    acceptance_ok: Optional[bool] = None

    if not bool(basis.get("enabled", True)):
        acceptance_status = "DISABLED"
    elif primary_value is None:
        acceptance_status = "NO_PRIMARY"
    elif reference_value is None:
        acceptance_status = "NO_REFERENCE"
    elif delta_abs is None:
        acceptance_status = "NO_DELTA"
    elif require_reference_status_ok and status != "OK":
        acceptance_status = "REFERENCE_NOT_OK"
    elif units_match is False:
        acceptance_status = "UNITS_MISMATCH"
    elif tolerance_abs is None:
        acceptance_status = "NO_TOLERANCE"
    elif delta_abs <= float(tolerance_abs):
        acceptance_status = "PASS"
        acceptance_ok = True
    else:
        acceptance_status = "FAIL"
        acceptance_ok = False

    return {
        "acceptance_status": acceptance_status,
        "acceptance_ok": acceptance_ok,
        "acceptance_tolerance_expr": basis.get("tolerance_expr"),
        "acceptance_tolerance_source": basis.get("tolerance_source"),
        "acceptance_tolerance_type": basis.get("tolerance_type"),
        "acceptance_tolerance_abs": tolerance_abs,
        "acceptance_tolerance_units": basis.get("tolerance_units"),
        "acceptance_expected_units": expected_units,
        "acceptance_reference_units": ref_units,
        "acceptance_units_match": units_match,
        "acceptance_span_basis_value": span_basis_value,
        "acceptance_span_basis_source": basis.get("span_basis_source"),
        "acceptance_delta_abs": delta_abs,
        "acceptance_delta_pct_span": delta_pct_span,
        "acceptance_selection_strategy": basis.get("selection_strategy"),
        "acceptance_tolerance_step_type": basis.get("tolerance_step_type"),
        "acceptance_tolerance_step_label": basis.get("tolerance_step_label"),
    }


def _summarize_reference_audit(
    session: Dict[str, Any],
    reference_path: Path,
) -> Tuple[Dict[str, Any], List[Dict[str, Any]], List[Dict[str, Any]], Optional[Dict[str, Any]]]:
    cfg = _reference_cfg_from_session(session)
    configured = _reference_cfg_seeded(cfg)
    captured_records = _read_jsonl(reference_path)
    live_snapshot = _live_reference_snapshot(cfg) if (configured and not captured_records) else None

    working_records = list(captured_records)
    origin = "captured"
    if not working_records and isinstance(live_snapshot, dict):
        working_records = [_snapshot_to_record(live_snapshot, cfg)]
        origin = "live_fallback"

    first_ts = None
    last_ts = None
    status_counts: Dict[str, int] = {}
    per_poll: Dict[str, Dict[str, Any]] = {}
    source_rows: Dict[Tuple[str, str], Dict[str, Any]] = {}
    trace_rows: List[Dict[str, Any]] = []
    latest_record: Optional[Dict[str, Any]] = None
    latest_trace_by_poll: Dict[str, Dict[str, Any]] = {}
    acceptance_cfg = _reference_acceptance_cfg(cfg)
    acceptance_status_counts: Dict[str, int] = {}
    acceptance_pass_count = 0
    acceptance_fail_count = 0
    acceptance_unknown_count = 0
    basis_cache: Dict[str, Dict[str, Any]] = {}

    for rec in working_records:
        if not isinstance(rec, dict):
            continue
        latest_record = rec
        ts_iso = str(rec.get("frame_ts_iso") or rec.get("ts_utc") or "")
        if ts_iso:
            if first_ts is None:
                first_ts = ts_iso
            last_ts = ts_iso

        status = str(rec.get("status") or "(UNKNOWN)").strip().upper() or "(UNKNOWN)"
        status_counts[status] = status_counts.get(status, 0) + 1

        src_key = (
            str(rec.get("source_path") or ""),
            str(rec.get("source_sha256") or ""),
        )
        src = source_rows.setdefault(src_key, {
            "record_origin": origin,
            "source_path": src_key[0],
            "source_sha256": src_key[1],
            "source_bytes": rec.get("source_bytes"),
            "provider": str(rec.get("provider") or ""),
            "role": str(rec.get("role") or ""),
            "first_ts": ts_iso or None,
            "last_ts": ts_iso or None,
            "record_count": 0,
            "latest_status": status,
            "source_mtime_iso": str(rec.get("source_mtime_iso") or ""),
        })
        src["record_count"] = int(src.get("record_count") or 0) + 1
        if ts_iso and not src.get("first_ts"):
            src["first_ts"] = ts_iso
        if ts_iso:
            src["last_ts"] = ts_iso
        src["latest_status"] = status
        if rec.get("source_bytes") not in (None, ""):
            src["source_bytes"] = rec.get("source_bytes")
        if rec.get("source_mtime_iso"):
            src["source_mtime_iso"] = rec.get("source_mtime_iso")

        measurements = rec.get("measurements") if isinstance(rec.get("measurements"), dict) else {}
        primary_values = rec.get("primary_values") if isinstance(rec.get("primary_values"), dict) else {}
        deltas = rec.get("deltas") if isinstance(rec.get("deltas"), dict) else {}
        for code, meas in measurements.items():
            if not isinstance(meas, dict):
                continue
            ref_val = _safe_float(meas.get("value"))
            pri_val = _safe_float(primary_values.get(code))
            delta_val = _safe_float(deltas.get(code))
            units = str(meas.get("units") or "").strip()
            source_column = str(meas.get("source_column") or "").strip()
            mapping_source = str(meas.get("mapping_source") or "").strip()

            row = {
                "record_origin": origin,
                "ts_utc": rec.get("ts_utc"),
                "frame_ts_iso": rec.get("frame_ts_iso"),
                "status": status,
                "ok": bool(rec.get("ok")),
                "provider": rec.get("provider"),
                "role": rec.get("role"),
                "pollutant": str(code),
                "primary_value": pri_val,
                "reference_value": ref_val,
                "delta": delta_val,
                "units": units,
                "source_column": source_column,
                "mapping_source": mapping_source,
                "source_path": rec.get("source_path"),
                "source_sha256": rec.get("source_sha256"),
                "source_mtime_iso": rec.get("source_mtime_iso"),
                "age_s": rec.get("age_s"),
                "header_index": rec.get("header_index"),
                "row_index": rec.get("row_index"),
                "delimiter": rec.get("delimiter"),
                "error": rec.get("error"),
            }
            basis = basis_cache.get(str(code))
            if basis is None:
                basis = _resolve_acceptance_basis(session, cfg, str(code))
                basis_cache[str(code)] = basis
            row.update(_evaluate_reference_acceptance(row, basis, bool(acceptance_cfg.get("require_reference_status_ok", True))))
            trace_rows.append(row)
            latest_trace_by_poll[str(code)] = row

            agg = per_poll.setdefault(str(code), {
                "units": units,
                "reference_count": 0,
                "reference_sum": 0.0,
                "reference_min": None,
                "reference_max": None,
                "primary_count": 0,
                "primary_sum": 0.0,
                "delta_count": 0,
                "delta_sum": 0.0,
                "abs_delta_sum": 0.0,
                "source_columns": set(),
                "mapping_sources": set(),
                "acceptance_status_counts": {},
                "acceptance_pass_count": 0,
                "acceptance_fail_count": 0,
                "acceptance_unknown_count": 0,
            })
            if units and not agg.get("units"):
                agg["units"] = units
            if source_column:
                agg["source_columns"].add(source_column)
            if mapping_source:
                agg["mapping_sources"].add(mapping_source)
            if ref_val is not None:
                agg["reference_count"] += 1
                agg["reference_sum"] += float(ref_val)
                agg["reference_min"] = float(ref_val) if agg["reference_min"] is None else min(float(agg["reference_min"]), float(ref_val))
                agg["reference_max"] = float(ref_val) if agg["reference_max"] is None else max(float(agg["reference_max"]), float(ref_val))
            if pri_val is not None:
                agg["primary_count"] += 1
                agg["primary_sum"] += float(pri_val)
            if delta_val is not None:
                agg["delta_count"] += 1
                agg["delta_sum"] += float(delta_val)
                agg["abs_delta_sum"] += abs(float(delta_val))
            acc_status = str(row.get("acceptance_status") or "NOT_EVALUATED")
            agg["acceptance_status_counts"][acc_status] = int(agg["acceptance_status_counts"].get(acc_status) or 0) + 1
            acceptance_status_counts[acc_status] = acceptance_status_counts.get(acc_status, 0) + 1
            if row.get("acceptance_ok") is True:
                agg["acceptance_pass_count"] += 1
                acceptance_pass_count += 1
            elif row.get("acceptance_ok") is False:
                agg["acceptance_fail_count"] += 1
                acceptance_fail_count += 1
            else:
                agg["acceptance_unknown_count"] += 1
                acceptance_unknown_count += 1

    per_poll_out: Dict[str, Any] = {}
    for code, agg in per_poll.items():
        ref_count = int(agg.get("reference_count") or 0)
        pri_count = int(agg.get("primary_count") or 0)
        delta_count = int(agg.get("delta_count") or 0)
        basis = basis_cache.get(code) or _resolve_acceptance_basis(session, cfg, code)
        latest_row = latest_trace_by_poll.get(code) or {}
        per_poll_out[code] = {
            "units": agg.get("units") or "",
            "reference_count": ref_count,
            "reference_avg": (agg["reference_sum"] / ref_count) if ref_count else None,
            "reference_min": agg.get("reference_min"),
            "reference_max": agg.get("reference_max"),
            "primary_count": pri_count,
            "primary_avg": (agg["primary_sum"] / pri_count) if pri_count else None,
            "delta_count": delta_count,
            "delta_avg": (agg["delta_sum"] / delta_count) if delta_count else None,
            "abs_delta_avg": (agg["abs_delta_sum"] / delta_count) if delta_count else None,
            "source_columns": sorted(list(agg.get("source_columns") or [])),
            "mapping_sources": sorted(list(agg.get("mapping_sources") or [])),
            "acceptance": {
                "status_counts": dict(agg.get("acceptance_status_counts") or {}),
                "pass_count": int(agg.get("acceptance_pass_count") or 0),
                "fail_count": int(agg.get("acceptance_fail_count") or 0),
                "unknown_count": int(agg.get("acceptance_unknown_count") or 0),
                "latest_status": latest_row.get("acceptance_status"),
                "latest_ok": latest_row.get("acceptance_ok"),
                "latest_primary_value": latest_row.get("primary_value"),
                "latest_reference_value": latest_row.get("reference_value"),
                "latest_delta": latest_row.get("delta"),
                "latest_delta_abs": latest_row.get("acceptance_delta_abs"),
                "latest_delta_pct_span": latest_row.get("acceptance_delta_pct_span"),
                "basis": {
                    "selection_strategy": basis.get("selection_strategy"),
                    "expected_units": basis.get("expected_units"),
                    "tolerance_expr": basis.get("tolerance_expr"),
                    "tolerance_source": basis.get("tolerance_source"),
                    "tolerance_type": basis.get("tolerance_type"),
                    "tolerance_abs": basis.get("tolerance_abs"),
                    "tolerance_units": basis.get("tolerance_units"),
                    "span_basis_value": basis.get("span_basis_value"),
                    "span_basis_source": basis.get("span_basis_source"),
                    "tolerance_step_type": basis.get("tolerance_step_type"),
                    "tolerance_step_label": basis.get("tolerance_step_label"),
                },
            },
        }

    latest_summary = None
    if isinstance(latest_record, dict):
        latest_summary = {
            "record_origin": origin,
            "ts_utc": latest_record.get("ts_utc"),
            "frame_ts_iso": latest_record.get("frame_ts_iso"),
            "status": latest_record.get("status"),
            "ok": latest_record.get("ok"),
            "provider": latest_record.get("provider"),
            "role": latest_record.get("role"),
            "source_path": latest_record.get("source_path"),
            "source_sha256": latest_record.get("source_sha256"),
            "source_bytes": latest_record.get("source_bytes"),
            "source_mtime_iso": latest_record.get("source_mtime_iso"),
            "age_s": latest_record.get("age_s"),
            "header_index": latest_record.get("header_index"),
            "row_index": latest_record.get("row_index"),
            "delimiter": latest_record.get("delimiter"),
            "updated_value_count": latest_record.get("updated_value_count"),
            "error": latest_record.get("error"),
            "measurements": latest_record.get("measurements") or {},
            "primary_values": latest_record.get("primary_values") or {},
            "deltas": latest_record.get("deltas") or {},
            "acceptance_by_pollutant": {
                code: {
                    "status": row.get("acceptance_status"),
                    "ok": row.get("acceptance_ok"),
                    "delta": row.get("delta"),
                    "delta_abs": row.get("acceptance_delta_abs"),
                    "delta_pct_span": row.get("acceptance_delta_pct_span"),
                    "tolerance_abs": row.get("acceptance_tolerance_abs"),
                    "tolerance_expr": row.get("acceptance_tolerance_expr"),
                    "tolerance_source": row.get("acceptance_tolerance_source"),
                }
                for code, row in latest_trace_by_poll.items()
            },
        }

    summary = {
        "configured": configured,
        "config": cfg if configured else {},
        "evidence": {
            "path": str(reference_path),
            "exists": reference_path.exists(),
            "captured_record_count": len(captured_records),
            "first_ts": first_ts if captured_records else None,
            "last_ts": last_ts if captured_records else None,
            "status_counts": status_counts if captured_records else {},
            "mode": "captured" if captured_records else ("live_fallback" if live_snapshot else "none"),
        },
        "latest": latest_summary,
        "acceptance": {
            "enabled": bool(acceptance_cfg.get("enabled", True)),
            "require_reference_status_ok": bool(acceptance_cfg.get("require_reference_status_ok", True)),
            "status_counts": dict(acceptance_status_counts),
            "pass_count": acceptance_pass_count,
            "fail_count": acceptance_fail_count,
            "unknown_count": acceptance_unknown_count,
        },
        "per_pollutant": per_poll_out,
    }

    return summary, trace_rows, list(source_rows.values()), latest_summary


def _parse_points_table(text: str) -> Tuple[int, List[Dict[str, Any]]]:
    """Parse the wizard-style points_table blob.

    Accepts CSV-ish or TSV-ish rows. Expected (len>=2):
      expected, observed[, passfail]

    Returns: (valid_row_count, rows)
    """
    rows: List[Dict[str, Any]] = []
    if not text:
        return 0, rows

    for raw in str(text).splitlines():
        line = raw.strip()
        if not line:
            continue

        # Try comma first, then tabs/spaces
        parts = [p.strip() for p in line.split(",")]
        if len(parts) < 2:
            parts = [p.strip() for p in line.split("\t")]
        if len(parts) < 2:
            # final fallback: multiple spaces
            parts = [p.strip() for p in line.split()]

        if len(parts) < 2:
            continue

        exp = _safe_float(parts[0])
        obs = _safe_float(parts[1])
        pf = parts[2].strip().upper() if len(parts) >= 3 else ""

        if exp is None or obs is None:
            continue

        rows.append({
            "expected": exp,
            "observed": obs,
            "passfail": pf,
        })

    return len(rows), rows


# -----------------------------
# QA/QC DB queries
# -----------------------------


def _db_table_columns(db_path: Path, table: str) -> List[str]:
    try:
        con = sqlite3.connect(str(db_path))
        cur = con.cursor()
        cur.execute(f"PRAGMA table_info({table})")
        cols = [str(r[1]) for r in cur.fetchall() if r and len(r) >= 2]
        con.close()
        return cols
    except Exception:
        return []


def _query_qaqc_events(db_path: Path, run_id: Optional[str]) -> List[Dict[str, Any]]:
    if not db_path.exists():
        return []
    cols = _db_table_columns(db_path, "qaqc_cal_event")
    if not cols:
        return []

    # Stable column set (schema v1)
    wanted = [
        "event_iso",
        "run_id",
        "phase",
        "pollutant",
        "event_type",
        "meas",
        "target",
        "recovery_pct",
        "drift_abs",
        "drift_recovery_pct",
        "drift_rate_abs_per_hr",
        "drift_rate_recovery_pct_per_hr",
        "combined_drift_ratio",
        "health",
        "manufacturer",
        "model_number",
        "serial_number",
        "site_id",
        "location_id",
        "source_category",
    ]
    sel = [c for c in wanted if c in cols]
    if not sel:
        return []

    where = ""
    params: Tuple[Any, ...] = tuple()
    if run_id:
        where = "WHERE run_id = ?"
        params = (run_id,)

    q = f"SELECT {', '.join(sel)} FROM qaqc_cal_event {where} ORDER BY event_iso ASC"

    out: List[Dict[str, Any]] = []
    try:
        con = sqlite3.connect(str(db_path))
        cur = con.cursor()
        cur.execute(q, params)
        for r in cur.fetchall():
            row = {sel[i]: r[i] for i in range(len(sel))}
            out.append(row)
        con.close()
    except Exception:
        return []

    return out


def _summarize_qaqc(events: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Summarize QA/QC events into per-pollutant pass/fail buckets."""
    per_pol: Dict[str, Dict[str, Any]] = {}
    overall_ok = True

    for ev in events:
        pol = str(ev.get("pollutant") or "").strip().upper() or "(UNKNOWN)"
        phase = str(ev.get("phase") or "").strip().upper() or ""
        etype = str(ev.get("event_type") or "").strip().upper() or ""
        health = str(ev.get("health") or "").strip().upper() or ""
        ok = (health in ("OK", "PASS", "GOOD"))
        if health and not ok:
            overall_ok = False

        bucket = per_pol.setdefault(pol, {
            "pre_zero": None,
            "pre_span": None,
            "post_zero": None,
            "post_span": None,
            "drift": None,
            "events": [],
        })

        rec = {
            "event_iso": ev.get("event_iso"),
            "phase": phase,
            "event_type": etype,
            "meas": ev.get("meas"),
            "target": ev.get("target"),
            "recovery_pct": ev.get("recovery_pct"),
            "drift_abs": ev.get("drift_abs"),
            "drift_recovery_pct": ev.get("drift_recovery_pct"),
            "combined_drift_ratio": ev.get("combined_drift_ratio"),
            "health": health,
            "ok": ok,
        }
        bucket["events"].append(rec)

        # Map into canonical slots (best-effort)
        if phase == "PRE" and etype == "ZERO":
            bucket["pre_zero"] = ok
        elif phase == "PRE" and etype in ("SPAN", "MID"):
            bucket["pre_span"] = ok
        elif phase == "POST" and etype == "ZERO":
            bucket["post_zero"] = ok
        elif phase == "POST" and etype in ("SPAN", "MID"):
            bucket["post_span"] = ok
        elif etype in ("DRIFT",):
            bucket["drift"] = ok

    # Decide pollutant pass/fail
    for pol, b in per_pol.items():
        # If any explicit False in known slots => fail
        slots = [b.get("pre_zero"), b.get("pre_span"), b.get("post_zero"), b.get("post_span"), b.get("drift")]
        if any(x is False for x in slots):
            b["overall_ok"] = False
        else:
            # If we have events but none of the slots are False, consider OK
            b["overall_ok"] = True if b.get("events") else None

    return {
        "overall_ok": bool(overall_ok) if events else None,
        "event_count": len(events),
        "per_pollutant": per_pol,
    }


# -----------------------------
# Test matrix scoring
# -----------------------------


def _score_test_matrix(session: Dict[str, Any]) -> Dict[str, Any]:
    tm = session.get("test_matrix")
    if not isinstance(tm, dict):
        return {
            "present": False,
            "overall_pass": None,
            "score_points": None,
            "score_points_possible": None,
            "tests": [],
        }

    tests_in = tm.get("tests") or []
    tests_out: List[Dict[str, Any]] = []

    score_possible = 0
    score_achieved = 0
    overall_pass = True

    for t in tests_in:
        if not isinstance(t, dict):
            continue
        tid = str(t.get("id") or "").strip() or "(UNKNOWN)"
        name = str(t.get("name") or tid)
        required = bool(t.get("required"))
        performed = bool(t.get("performed"))
        waived = bool(t.get("waived"))
        waive_reason = (t.get("waive_reason") or "")
        pts = int(t.get("points") or 0)
        min_pts = int(t.get("min_points") or 0)

        if pts > 0:
            score_possible += pts

        status = ""
        details = ""

        # Default scoring:
        #  - WAIVED counts as satisfied (but does not add points)
        #  - PASS adds points
        #  - FAIL adds zero
        #  - Missing required => FAIL

        if waived:
            status = "WAIVED"
            details = waive_reason.strip() or "(no reason)"
        elif not performed:
            if required:
                status = "FAIL"
                details = "Required test not performed"
                overall_pass = False
            else:
                status = "SKIP"
                details = "Not performed"
        else:
            # performed
            status = "PASS"
            details = "Performed"

            capture = t.get("capture_fields") if isinstance(t.get("capture_fields"), dict) else {}

            if tid in ("CAL", "LIN"):
                pt_blob = str(capture.get("points_table") or "")
                nrows, rows = _parse_points_table(pt_blob)

                # Enforce minimum point count if specified
                if min_pts and nrows < min_pts:
                    status = "FAIL"
                    details = f"Only {nrows} point(s) captured (min {min_pts})"
                    overall_pass = False
                else:
                    # Optional explicit pass/fail flags per row
                    any_fail = any((str(r.get("passfail") or "").strip().upper() == "FAIL") for r in rows)
                    if any_fail:
                        status = "FAIL"
                        details = "One or more points marked FAIL"
                        overall_pass = False
                    else:
                        details = f"{nrows} point(s) captured"

            else:
                # Generic tests can optionally provide a result
                res = str(capture.get("result") or "").strip().upper()
                if res in ("FAIL", "FAILED"):
                    status = "FAIL"
                    details = "Marked FAIL"
                    overall_pass = False

        if status == "PASS" and pts > 0:
            score_achieved += pts

        tests_out.append({
            "id": tid,
            "name": name,
            "required": required,
            "performed": performed,
            "waived": waived,
            "waive_reason": waive_reason,
            "points": pts,
            "min_points": min_pts,
            "status": status,
            "details": details,
        })

    if not tests_out:
        overall_pass = None

    return {
        "present": True,
        "overall_pass": overall_pass,
        "score_points": score_achieved,
        "score_points_possible": score_possible,
        "tests": tests_out,
    }


def _report_template_contract() -> Dict[str, Any]:
    docs_dir = Path(__file__).resolve().parent.parent / "docs"
    master_template = docs_dir / "MOLE_DAS_MASTER_TEST_REPORT_TEMPLATE_2026_04_09.md"
    crosswalk = docs_dir / "MOLE_DAS_TEST_REPORT_DATA_ELEMENTS_CROSSWALK_2026_04_09.md"
    assessment = docs_dir / "epa_emissions_performance_test_structural_assessment.md"
    return {
        "master_template": {
            "name": "MOLE DAS Master Test Report Template",
            "version": "2026_04_09",
            "path": str(master_template),
            "exists": master_template.exists(),
        },
        "crosswalk": {
            "name": "MOLE DAS Test Report Data Elements Crosswalk",
            "version": "2026_04_09",
            "path": str(crosswalk),
            "exists": crosswalk.exists(),
        },
        "structural_assessment": {
            "name": "EPA Emissions Performance Test Structural Assessment",
            "path": str(assessment),
            "exists": assessment.exists(),
        },
    }


def _report_limit_lookup(session: Dict[str, Any]) -> Dict[str, List[Dict[str, Any]]]:
    reg = session.get("regulatory") if isinstance(session.get("regulatory"), dict) else {}
    limits = reg.get("limits") if isinstance(reg.get("limits"), list) else []
    out: Dict[str, List[Dict[str, Any]]] = {}
    for row in limits:
        if not isinstance(row, dict):
            continue
        code = str(row.get("pollutant") or "").strip().upper()
        if not code:
            continue
        out.setdefault(code, []).append(row)
    return out


def _report_pollutant_method_rows(session: Dict[str, Any]) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    pres = _pollutant_prescriptions(session)
    limits_by_code = _report_limit_lookup(session)
    for code in _selected_pollutants(session):
        cfg = pres.get(code) if isinstance(pres.get(code), dict) else {}
        rows.append({
            "pollutant": code,
            "method": str(cfg.get("method_selected") or ""),
            "instrument": _first_present(cfg.get("instrument_id"), cfg.get("instrument_channel"), cfg.get("channel_id")),
            "units": str(cfg.get("expected_units") or ""),
            "compliance_basis": [
                {
                    "limit_type": str(lim.get("limit_type") or ""),
                    "value": lim.get("value"),
                    "units": str(lim.get("units") or ""),
                    "basis": str(lim.get("basis") or ""),
                    "avg_time": str(lim.get("avg_time") or ""),
                    "o2_ref_pct": lim.get("o2_ref_pct"),
                }
                for lim in limits_by_code.get(code, [])
                if isinstance(lim, dict)
            ],
            "source": "session.pollutants.prescriptions + session.regulatory.limits",
        })
    return rows


def _report_planned_run_rows(session: Dict[str, Any]) -> List[Dict[str, Any]]:
    tm = session.get("test_matrix") if isinstance(session.get("test_matrix"), dict) else {}
    plan = tm.get("plan") if isinstance(tm.get("plan"), dict) else {}
    run_durations = plan.get("run_durations_min") if isinstance(plan.get("run_durations_min"), list) else []
    sample_runs = int(_safe_float(plan.get("sample_runs")) or 0)
    minutes_per_run = _safe_float(plan.get("minutes_per_run"))
    rows: List[Dict[str, Any]] = []
    if run_durations:
        for idx, dur in enumerate(run_durations):
            rows.append({
                "run_no": idx + 1,
                "planned_duration_min": _safe_float(dur),
                "date": None,
                "start": None,
                "stop": None,
                "status": "PLANNED",
                "source": "session.test_matrix.plan.run_durations_min",
            })
        return rows
    if sample_runs > 0 and minutes_per_run is not None:
        for idx in range(sample_runs):
            rows.append({
                "run_no": idx + 1,
                "planned_duration_min": minutes_per_run,
                "date": None,
                "start": None,
                "stop": None,
                "status": "PLANNED",
                "source": "session.test_matrix.plan.sample_runs + minutes_per_run",
            })
    return rows


def _safe_slug(text: Any) -> str:
    s = str(text or "").strip().lower()
    s = re.sub(r"[^a-z0-9]+", "_", s)
    return s.strip("_")


def _report_company_name_from_operator(operator: Any) -> Optional[str]:
    raw = str(operator or "").strip()
    if not raw:
        return None
    return raw.replace("_", " ").strip().title()


def _report_builder_block(session: Dict[str, Any]) -> Dict[str, Any]:
    blk = session.get("report_builder") if isinstance(session.get("report_builder"), dict) else {}
    return blk if isinstance(blk, dict) else {}


def _report_builder_subblock(session: Dict[str, Any], key: str) -> Dict[str, Any]:
    blk = _report_builder_block(session)
    sub = blk.get(key)
    return sub if isinstance(sub, dict) else {}


def _report_builder_list(value: Any) -> List[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item or "").strip()]
    txt = str(value or "").replace("\r", "\n").strip()
    if not txt:
        return []
    return [str(item).strip() for item in re.split(r"[;\n]+", txt) if str(item or "").strip()]


def _report_block_status(required: List[Any], optional: Optional[List[Any]] = None) -> str:
    optional = optional or []
    req_total = len(required)
    req_filled = sum(1 for value in required if _has_value(value))
    opt_filled = sum(1 for value in optional if _has_value(value))
    if req_total > 0 and req_filled == req_total:
        return "Available"
    if req_filled > 0 or opt_filled > 0:
        return "Partial"
    return "Gap"


def _report_parties_block(session: Dict[str, Any]) -> Dict[str, Any]:
    project = session.get("project") if isinstance(session.get("project"), dict) else {}
    intake = project.get("intake") if isinstance(project.get("intake"), dict) else {}
    parties_in = _report_builder_subblock(session, "parties")
    operator_raw = project.get("operator")
    operator_display = _report_company_name_from_operator(operator_raw)
    site_facility = str(project.get("site_facility") or "").strip() or None
    client_name = str(parties_in.get("client_name") or "").strip() or None
    facility_owner_operator_name = str(parties_in.get("facility_owner_operator_name") or "").strip() or site_facility
    test_company_name = str(parties_in.get("test_company_name") or "").strip() or operator_display
    session_operator_name = str(parties_in.get("session_operator_name") or "").strip() or operator_display
    laboratory_name = str(parties_in.get("laboratory_name") or "").strip() or None
    observer_contacts = _report_builder_list(parties_in.get("observer_contacts"))
    responsible_official_name = str(parties_in.get("responsible_official_name") or "").strip() or None
    responsible_official_title = str(parties_in.get("responsible_official_title") or "").strip() or None
    notes: List[str] = []
    status = _report_block_status(
        [
            client_name,
            facility_owner_operator_name,
            test_company_name,
            responsible_official_name,
        ],
        [
            laboratory_name,
            observer_contacts,
            session_operator_name,
            responsible_official_title,
        ],
    )
    if status != "Available":
        notes.append("Client, signatory, laboratory, or observer metadata is still incomplete.")
    return {
        "status": status,
        "client_name": client_name,
        "facility_owner_operator_name": facility_owner_operator_name,
        "test_company_name": test_company_name,
        "session_operator_name": session_operator_name,
        "laboratory_name": laboratory_name,
        "observer_contacts": observer_contacts,
        "responsible_official_name": responsible_official_name,
        "responsible_official_title": responsible_official_title,
        "signatory_required": True,
        "site_contact_provided_flag": bool(((intake.get("items") or {}).get("SITE_CONTACT") or {}).get("provided")),
        "notes": notes,
        "source": "session.report_builder.parties + session.project + session.project.intake",
    }


def _report_process_control_block(session: Dict[str, Any]) -> Dict[str, Any]:
    source = session.get("source") if isinstance(session.get("source"), dict) else {}
    fuel = session.get("fuel") if isinstance(session.get("fuel"), dict) else {}
    process_in = _report_builder_subblock(session, "process_control")
    exhaust_flow = source.get("exhaust_flow") if isinstance(source.get("exhaust_flow"), dict) else {}
    fuel_flow = source.get("fuel_flow") if isinstance(source.get("fuel_flow"), dict) else {}

    narrative_parts = [
        str(source.get("service_class") or "").strip(),
        str(source.get("source_category") or "").strip(),
        f"used for {str(source.get('application') or source.get('source_application') or '').strip()}".strip(),
    ]
    narrative_parts = [part for part in narrative_parts if part and part != "used for"]
    narrative = " ".join(narrative_parts).strip()
    if source.get("manufacturer") or source.get("model_number"):
        unit_desc = " ".join([part for part in [str(source.get("manufacturer") or "").strip(), str(source.get("model_number") or "").strip()] if part]).strip()
        if unit_desc:
            narrative = (narrative + f"; unit {unit_desc}").strip("; ")
    if fuel.get("fuel_category"):
        narrative = (narrative + f"; fuel {fuel.get('fuel_category')}").strip("; ")
    process_narrative = str(process_in.get("process_narrative") or "").strip() or narrative or None
    control_equipment_description = str(process_in.get("control_equipment_description") or "").strip() or None
    status = _report_block_status(
        [process_narrative, control_equipment_description],
        [source.get("notes"), source.get("manufacturer"), source.get("model_number")],
    )

    return {
        "status": status,
        "process_description": {
            "source_category": source.get("source_category"),
            "application": source.get("application"),
            "service_class": source.get("service_class"),
            "output_type": source.get("output_type"),
            "manufacturer": source.get("manufacturer"),
            "model_number": source.get("model_number"),
            "serial_number": source.get("serial_number"),
            "asset_tag": source.get("asset_tag"),
            "engine_cycle": source.get("engine_cycle"),
            "engine_cyl_count": _first_present(source.get("engine_cyl_count"), source.get("cyl_count")),
            "fuel_category": fuel.get("fuel_category"),
            "fuel_flow_basis": fuel_flow.get("basis"),
            "fuel_flow_units": fuel_flow.get("units"),
            "exhaust_flow_method": exhaust_flow.get("method"),
            "unit_notes": source.get("notes"),
        },
        "process_narrative_seed": process_narrative,
        "process_narrative": process_narrative,
        "control_equipment_description": control_equipment_description,
        "control_equipment": {
            "status": "DESCRIBED" if _has_value(control_equipment_description) else "UNKNOWN",
            "description": control_equipment_description,
            "operating_parameters": [],
            "note": None if _has_value(control_equipment_description) else "Control equipment description is not normalized in the current session schema.",
        },
        "source": "session.report_builder.process_control + session.source + session.fuel",
    }


def _report_deviation_approval_block(session: Dict[str, Any], evidence_bundle: Dict[str, Any]) -> Dict[str, Any]:
    source = session.get("source") if isinstance(session.get("source"), dict) else {}
    dev_in = _report_builder_subblock(session, "deviations_approvals")
    stack = source.get("stack") if isinstance(source.get("stack"), dict) else {}
    exhaust_flow = source.get("exhaust_flow") if isinstance(source.get("exhaust_flow"), dict) else {}
    static_files = ((evidence_bundle.get("static_artifacts") or {}).get("files") if isinstance(evidence_bundle.get("static_artifacts"), dict) else []) or []
    candidate_titles = []
    for item in static_files:
        if not isinstance(item, dict):
            continue
        title = str(item.get("title") or item.get("name") or "")
        title_l = title.lower()
        if any(tok in title_l for tok in ["deviation", "approval", "alt method", "alternate method", "agency", "permit", "notice"]):
            candidate_titles.append(title)
    notes = []
    if _has_value(stack.get("hand_notes")):
        notes.append(f"Stack notes: {stack.get('hand_notes')}")
    if _has_value(((exhaust_flow.get("stack_measured") or {}).get("notes") if isinstance(exhaust_flow.get("stack_measured"), dict) else None)):
        notes.append(f"Method notes: {(exhaust_flow.get('stack_measured') or {}).get('notes')}")
    planned_deviations = _report_builder_list(dev_in.get("planned_deviations"))
    field_deviations = _report_builder_list(dev_in.get("field_deviations"))
    alt_approvals = _report_builder_list(dev_in.get("alternative_method_approvals"))
    impact_statement = str(dev_in.get("impact_statement") or "").strip() or None
    status = "Available" if any([planned_deviations, field_deviations, alt_approvals, impact_statement]) else ("Partial" if notes or candidate_titles else "Gap")
    if status != "Available" and not notes:
        notes = ["No dedicated deviation/approval records are normalized in the current session schema."]
    return {
        "status": status,
        "planned_deviations": planned_deviations,
        "field_deviations": field_deviations,
        "alternative_method_approvals": alt_approvals,
        "impact_statement": impact_statement,
        "candidate_evidence_titles": candidate_titles,
        "notes": notes,
        "source": "session.report_builder.deviations_approvals + session.source.stack.hand_notes + session.source.exhaust_flow.stack_measured.notes + evidence_bundle.static_artifacts",
    }


def _report_correspondence_block(session: Dict[str, Any], evidence_bundle: Dict[str, Any]) -> Dict[str, Any]:
    project = session.get("project") if isinstance(session.get("project"), dict) else {}
    intake = project.get("intake") if isinstance(project.get("intake"), dict) else {}
    corr_in = _report_builder_subblock(session, "correspondence")
    evidence_files = intake.get("evidence_files") if isinstance(intake.get("evidence_files"), list) else []
    static_files = ((evidence_bundle.get("static_artifacts") or {}).get("files") if isinstance(evidence_bundle.get("static_artifacts"), dict) else []) or []
    candidates = []
    for item in evidence_files:
        if _has_value(item):
            candidates.append({"title": Path(str(item)).name, "path": str(item), "source": "session.project.intake.evidence_files"})
    for item in static_files:
        if not isinstance(item, dict):
            continue
        title = str(item.get("title") or item.get("name") or "")
        title_l = title.lower()
        if any(tok in title_l for tok in ["notice", "agency", "permit", "cedri", "ert", "email", "correspondence", "approval"]):
            candidates.append({
                "title": title,
                "path": str(item.get("path") or item.get("full_path") or ""),
                "source": "evidence_bundle.static_artifacts.files",
            })
    notice_of_intent_date = str(corr_in.get("notice_of_intent_date") or "").strip() or None
    agency_contact = str(corr_in.get("agency_contact") or "").strip() or None
    approval_dates = _report_builder_list(corr_in.get("approval_dates"))
    submission_status = str(corr_in.get("submission_status") or "").strip() or None
    corr_notes = str(corr_in.get("notes") or "").strip()
    status = "Available" if any([notice_of_intent_date, agency_contact, approval_dates, submission_status, corr_notes]) else ("Partial" if candidates else "Gap")
    notes = [corr_notes] if corr_notes else (
        ["Regulatory correspondence tracking is not normalized yet; candidates are inferred from evidence attachments."] if candidates else ["No normalized correspondence tracking exists in the current session schema."]
    )
    return {
        "status": status,
        "notice_of_intent_date": notice_of_intent_date,
        "agency_contact": agency_contact,
        "approval_dates": approval_dates,
        "submission_status": submission_status,
        "candidate_artifacts": candidates,
        "notes": notes,
        "source": "session.report_builder.correspondence + session.project.intake.evidence_files + evidence_bundle.static_artifacts",
    }


def _report_run_aggregation(
    session: Dict[str, Any],
    evidence_bundle: Dict[str, Any],
    *,
    cfg_path: Path,
    session_dir: Path,
) -> Dict[str, Any]:
    planned_rows = _report_planned_run_rows(session)
    scoped_worksteps = _load_scoped_worksteps(session, cfg_path=cfg_path, session_dir=session_dir)
    worksteps_path = scoped_worksteps.get("path")
    worksteps_summary = ((evidence_bundle.get("sources") or {}).get("worksteps") if isinstance(evidence_bundle.get("sources"), dict) else {}) or {}
    events = scoped_worksteps.get("rows") if isinstance(scoped_worksteps.get("rows"), list) else []

    rows_by_run: Dict[int, Dict[str, Any]] = {}
    current_run_no: Optional[int] = None

    def _ensure_row(run_no: int) -> Dict[str, Any]:
        row = rows_by_run.get(run_no)
        if row is None:
            plan_row = next((r for r in planned_rows if int(r.get("run_no") or 0) == run_no), {})
            row = {
                "run_no": run_no,
                "planned_duration_min": plan_row.get("planned_duration_min"),
                "start_ts_iso": None,
                "end_ts_iso": None,
                "actual_duration_min": None,
                "postcal_required": False,
                "postcal_complete": False,
                "event_counts": {},
                "post_zero_pollutants": [],
                "post_span_pollutants": [],
                "status": "PLANNED",
                "source": "worksteps",
            }
            rows_by_run[run_no] = row
        return row

    def _infer_run_no(evt: Dict[str, Any]) -> Optional[int]:
        payload = evt.get("payload") if isinstance(evt.get("payload"), dict) else {}
        run_no = payload.get("run_no")
        if run_no is not None:
            try:
                return int(run_no)
            except Exception:
                pass
        idx = payload.get("active_run_index")
        if idx is not None:
            try:
                return int(idx) + 1
            except Exception:
                pass
        return current_run_no

    for evt in events:
        if not isinstance(evt, dict):
            continue
        event_name = str(evt.get("event") or evt.get("type") or "").strip().upper()
        ts_iso = str(evt.get("ts_iso") or "")
        payload = evt.get("payload") if isinstance(evt.get("payload"), dict) else {}
        run_no = _infer_run_no(evt)

        if event_name == "TEST_START":
            run_no = run_no or 1
            row = _ensure_row(run_no)
            row["start_ts_iso"] = ts_iso or row.get("start_ts_iso")
            row["status"] = "STARTED"
            current_run_no = run_no
        elif event_name == "TEST_END":
            run_no = run_no or current_run_no or 1
            row = _ensure_row(run_no)
            row["end_ts_iso"] = ts_iso or row.get("end_ts_iso")
            row["status"] = "COMPLETED"
            current_run_no = None
        elif run_no is not None:
            row = _ensure_row(run_no)
            counts = row.get("event_counts") if isinstance(row.get("event_counts"), dict) else {}
            counts[event_name] = int(counts.get(event_name) or 0) + 1
            row["event_counts"] = counts
            if event_name == "POSTCAL_REQUIRED":
                row["postcal_required"] = True
            elif event_name == "POSTCAL_COMPLETE":
                row["postcal_complete"] = True
            elif event_name == "CAPTURE_POST_ZERO":
                pol = str(payload.get("pollutant") or "").strip().upper()
                if pol and pol not in row["post_zero_pollutants"]:
                    row["post_zero_pollutants"].append(pol)
            elif event_name == "CAPTURE_POST_SPAN":
                pol = str(payload.get("pollutant") or "").strip().upper()
                if pol and pol not in row["post_span_pollutants"]:
                    row["post_span_pollutants"].append(pol)

    actual_rows: List[Dict[str, Any]] = []
    for run_no in sorted(rows_by_run):
        row = rows_by_run[run_no]
        start_dt = _parse_iso_dt(row.get("start_ts_iso"))
        end_dt = _parse_iso_dt(row.get("end_ts_iso"))
        if start_dt and end_dt:
            row["actual_duration_min"] = round((end_dt - start_dt).total_seconds() / 60.0, 3)
            row["date"] = start_dt.date().isoformat()
        else:
            row["date"] = None
        actual_rows.append(row)

    if actual_rows:
        coverage_note = "Run summaries are aggregated from session-scoped workstep events."
    else:
        if str(scoped_worksteps.get("path_scope") or "") == "EXTERNAL" and int(worksteps_summary.get("filtered_out_count") or 0) > 0:
            coverage_note = "No actual run events were found within the current session window; external worksteps outside that window were excluded."
        elif str(scoped_worksteps.get("path_scope") or "") == "EXTERNAL":
            coverage_note = "No actual run events were found in the external worksteps source for this session window; only planned Test Matrix rows are available."
        else:
            coverage_note = "No actual run events were found; only planned Test Matrix rows are available."
    return {
        "status": "Partial" if actual_rows else ("Available" if planned_rows else "Gap"),
        "planned_runs": planned_rows,
        "actual_runs": actual_rows,
        "run_count_planned": len(planned_rows),
        "run_count_actual": len(actual_rows),
        "coverage_note": coverage_note,
        "source": "session.test_matrix.plan + session-scoped evidence_bundle.sources.worksteps",
        "worksteps_path": str(worksteps_path) if _has_value(worksteps_path) else "",
        "worksteps_path_scope": scoped_worksteps.get("path_scope"),
        "worksteps_scope_note": scoped_worksteps.get("scope_note"),
        "session_scope": dict(scoped_worksteps.get("session_scope") or {}),
        "worksteps_total_events": int(worksteps_summary.get("total_count") or 0),
        "worksteps_scoped_events": len(events),
        "worksteps_filtered_out": int(worksteps_summary.get("filtered_out_count") or 0),
    }


def _appendix_entry(
    appendix: str,
    title: str,
    artifact_type: str,
    path: Any,
    source: str,
    note: str = "",
) -> Dict[str, Any]:
    p = None
    exists = False
    sha256 = ""
    bytes_size = None
    try:
        if _has_value(path):
            p = Path(str(path)).expanduser()
            exists = p.exists()
            if exists and p.is_file():
                bytes_size = p.stat().st_size
                sha256 = _sha256_file(p)
    except Exception:
        exists = False
        sha256 = ""
        bytes_size = None
    return {
        "appendix": appendix,
        "appendix_title": {
            "A": "Regulatory and administrative support",
            "B": "Source and sampling location support",
            "C": "Method and instrument support",
            "D": "QA/QC support",
            "E": "Field and raw data",
            "F": "Calculations and report outputs",
        }.get(str(appendix or "").upper(), ""),
        "title": title,
        "artifact_type": artifact_type,
        "path": str(p) if p is not None else str(path or ""),
        "exists": exists,
        "sha256": sha256,
        "bytes": bytes_size,
        "include_in_final_report": bool(exists),
        "status": "AVAILABLE" if exists else "MISSING",
        "assignment_basis": source,
        "source": source,
        "note": note,
    }


def _report_appendix_manifest(
    paths: "ReportPackPaths",
    evidence_bundle: Dict[str, Any],
) -> List[Dict[str, Any]]:
    entries: List[Dict[str, Any]] = []
    sources = evidence_bundle.get("sources") if isinstance(evidence_bundle.get("sources"), dict) else {}

    for appendix, title, artifact_type, path, source, note in [
        ("D", "Analyzer validity summary", "csv", paths.analyzer_validity_csv, "report_pack_v1", "QA/QC support export"),
        ("D", "Test matrix scorecard", "csv", paths.test_matrix_csv, "report_pack_v1", "QA/QC support export"),
        ("D", "QA/QC events", "csv", paths.qaqc_events_csv, "report_pack_v1", "QA/QC support export"),
        ("D", "Spike recovery snapshot", "csv", paths.spike_recovery_csv, "report_pack_v1", "QA/QC support export"),
        ("D", "Side-by-side snapshot", "csv", paths.side_by_side_csv, "report_pack_v1", "QA/QC support export"),
        ("D", "Reference audit trace", "csv", paths.reference_trace_csv, "report_pack_v1", "QA/QC support export"),
        ("D", "Reference audit acceptance", "csv", paths.reference_acceptance_csv, "report_pack_v1", "QA/QC support export"),
        ("D", "Reference audit latest snapshot", "json", paths.reference_latest_json, "report_pack_v1", "QA/QC support export"),
        ("F", "FTIR validation summary", "json", paths.ftir_validation_json, "ftir_validation_v1", "Session-scoped FTIR side-by-side validation summary"),
        ("F", "FTIR validation window alignment", "csv", paths.ftir_validation_windows_csv, "ftir_validation_v1", "Aligned MOLE / FTIR comparison windows"),
        ("F", "FTIR validation Method 301 stats", "csv", paths.ftir_validation_method301_csv, "ftir_validation_v1", "Method 301 bias / precision comparison statistics"),
        ("F", "FTIR validation signoff cover", "md", paths.ftir_validation_appendix_cover_md, "ftir_validation_appendix_v1", "Reviewer-facing FTIR validation signoff cover sheet."),
        ("F", "FTIR validation signed comparison ledger", "csv", paths.ftir_validation_appendix_ledger_csv, "ftir_validation_appendix_v1", "Signed aligned-window ledger bound to the frozen FTIR validation snapshot."),
        ("F", "FTIR validation signed Method 301 stats", "csv", paths.ftir_validation_appendix_method301_csv, "ftir_validation_appendix_v1", "Signed Method 301 bias / precision statistics bound to the frozen FTIR validation snapshot."),
        ("F", "FTIR validation signed exclusion register", "csv", paths.ftir_validation_appendix_exclusions_csv, "ftir_validation_appendix_v1", "Signed exclusion register for FTIR validation windows."),
        ("F", "FTIR validation reviewer workbook", "xlsx", paths.ftir_validation_appendix_workbook_xlsx, "ftir_validation_appendix_v1", "Reviewer workbook containing signoff summary, comparison sets, aligned rows, Method 301 stats, and exclusions."),
        ("F", "FTIR validation appendix index", "json", paths.ftir_validation_appendix_index_json, "ftir_validation_appendix_v1", "Index of reviewer-facing FTIR validation appendix artifacts."),
        ("E", "Workstep log", "jsonl", ((sources.get("worksteps") or {}).get("path") if isinstance(sources.get("worksteps"), dict) else ""), "evidence_bundle.sources.worksteps", "Field and activity log"),
        ("E", "Raw samples", "jsonl", ((sources.get("raw_samples") or {}).get("path") if isinstance(sources.get("raw_samples"), dict) else ""), "evidence_bundle.sources.raw_samples", "Raw sample evidence"),
        ("E", "Raw events", "jsonl", ((sources.get("raw_events") or {}).get("path") if isinstance(sources.get("raw_events"), dict) else ""), "evidence_bundle.sources.raw_events", "Runner event evidence"),
        ("E", "Health states", "jsonl", ((sources.get("health_states") or {}).get("path") if isinstance(sources.get("health_states"), dict) else ""), "evidence_bundle.sources.health_states", "Channel health trace"),
        ("E", "Reference audit raw trace", "jsonl", ((sources.get("reference_audit") or {}).get("path") if isinstance(sources.get("reference_audit"), dict) else ""), "evidence_bundle.sources.reference_audit", "Reference audit raw evidence"),
        ("F", "Report pack summary", "json", paths.summary_json, "report_pack_v1", "Compact report-pack summary"),
        ("F", "Report context", "json", paths.report_context_json, "report_context_v1", "Normalized final-report source object"),
        ("F", "Final test report", "md", paths.final_report_md, "final_report_v1", "Markdown render bound from report_context_v1 and the master template."),
        ("F", "Final test report", "docx", paths.final_report_docx, "final_report_v1", "DOCX render bound from report_context_v1 and the master template."),
        ("F", "Final test report", "pdf", paths.final_report_pdf, "final_report_v1", "PDF render bound from report_context_v1 and the master template."),
        ("F", "Evidence bundle", "json", paths.evidence_bundle_json, "report_pack_v1", "Evidence manifest source"),
        ("F", "Evidence step evaluation", "csv", paths.evidence_step_eval_csv, "report_pack_v1", "Calculation / step-evaluation support"),
        ("F", "Fuel analysis snapshot", "csv", paths.fuel_analysis_csv, "report_pack_v1", "Fuel analysis output"),
        ("F", "Pollutant adjustments snapshot", "csv", paths.pollutant_adjustments_csv, "report_pack_v1", "Bias / drift adjustment output"),
        ("F", "Pollutant adjustment history", "csv", paths.pollutant_adjustment_history_csv, "report_pack_v1", "Bias / drift history"),
        ("F", "Regulatory snapshot", "csv", paths.regulatory_snapshot_csv, "report_pack_v1", "Regulatory comparison output"),
        ("F", "Reference audit sources", "csv", paths.reference_sources_csv, "report_pack_v1", "Reference source mapping"),
        ("F", "Reference offset recommendations", "csv", paths.reference_offsets_csv, "report_pack_v1", "Reference offset recommendation output"),
        ("F", "Reference offset recommendations", "json", paths.reference_offsets_json, "report_pack_v1", "Reference offset recommendation output"),
        ("F", "Report pack PDF", "pdf", paths.pdf_path, "report_pack_v1", "Compact report-pack PDF"),
    ]:
        entries.append(_appendix_entry(appendix, title, artifact_type, path, source, note))

    static_block = evidence_bundle.get("static_artifacts") if isinstance(evidence_bundle.get("static_artifacts"), dict) else {}
    for item in static_block.get("files") or []:
        if not isinstance(item, dict):
            continue
        category = str(item.get("category") or "STATIC_ARTIFACT")
        title = str(item.get("title") or item.get("name") or category)
        entries.append(_appendix_entry(
            appendix="B",
            title=title,
            artifact_type=str(item.get("artifact_type") or item.get("type") or "file"),
            path=item.get("path") or item.get("full_path"),
            source="evidence_bundle.static_artifacts.files",
            note=f"Static artifact category: {category}",
        ))
    return entries


def _build_report_context(
    session: Dict[str, Any],
    summary: Dict[str, Any],
    evidence_bundle: Dict[str, Any],
    paths: "ReportPackPaths",
    cfg_path: Path,
    session_dir: Path,
) -> Dict[str, Any]:
    project = session.get("project") if isinstance(session.get("project"), dict) else {}
    source = session.get("source") if isinstance(session.get("source"), dict) else {}
    stack = source.get("stack") if isinstance(source.get("stack"), dict) else {}
    intake = project.get("intake") if isinstance(project.get("intake"), dict) else {}
    regulatory = session.get("regulatory") if isinstance(session.get("regulatory"), dict) else {}
    site = session.get("site_conditions") if isinstance(session.get("site_conditions"), dict) else {}
    location = site.get("location") if isinstance(site.get("location"), dict) else {}
    weather = site.get("weather_compile") if isinstance(site.get("weather_compile"), dict) else {}
    fuel = session.get("fuel") if isinstance(session.get("fuel"), dict) else {}
    run_rows = _report_planned_run_rows(session)
    pollutant_rows = _report_pollutant_method_rows(session)
    rules = regulatory.get("selected_rule_ids") if isinstance(regulatory.get("selected_rule_ids"), list) else []
    regulatory_rows = ((summary.get("regulatory_snapshot") or {}).get("rows") if isinstance(summary.get("regulatory_snapshot"), dict) else []) or []
    appendix_entries = _report_appendix_manifest(paths, evidence_bundle)
    report_parties = _report_parties_block(session)
    process_control = _report_process_control_block(session)
    deviations = _report_deviation_approval_block(session, evidence_bundle)
    correspondence = _report_correspondence_block(session, evidence_bundle)
    run_aggregation = _report_run_aggregation(session, evidence_bundle, cfg_path=cfg_path, session_dir=session_dir)
    ftir_validation = summary.get("ftir_validation") if isinstance(summary.get("ftir_validation"), dict) else {}

    report_title = _first_present(
        project.get("project_name"),
        project.get("job_id"),
        "MOLE DAS Test Report",
    )
    unit_asset = " / ".join([part for part in [
        " ".join([part for part in [str(source.get("manufacturer") or "").strip(), str(source.get("model_number") or "").strip()] if part]),
        str(project.get("asset_unit_id") or "").strip(),
    ] if part])
    if not unit_asset:
        unit_asset = None

    planned_schedule_note = ""
    if run_rows:
        durations = [str(_fmt_num(row.get("planned_duration_min"), 3, "")).rstrip("0").rstrip(".") for row in run_rows]
        durations = [d for d in durations if d]
        planned_schedule_note = (
            f"{len(run_rows)} planned run(s)"
            + (f"; durations (min): {', '.join(durations)}" if durations else "")
        )

    actual_run_rows = run_aggregation.get("actual_runs") if isinstance(run_aggregation.get("actual_runs"), list) else []
    planned_run_rows = run_aggregation.get("planned_runs") if isinstance(run_aggregation.get("planned_runs"), list) else run_rows
    run_schedule_status = str(run_aggregation.get("status") or ("Available" if planned_run_rows else "Gap"))
    results_status = "Partial" if regulatory_rows else "Gap"
    actual_dates = None
    if actual_run_rows:
        start_candidates = [str(row.get("start_ts_iso") or "").strip() for row in actual_run_rows if _has_value(row.get("start_ts_iso"))]
        end_candidates = [str(row.get("end_ts_iso") or "").strip() for row in actual_run_rows if _has_value(row.get("end_ts_iso"))]
        actual_dates = {
            "first_ts": min(start_candidates) if start_candidates else None,
            "last_ts": max(end_candidates) if end_candidates else None,
        }
    if not actual_dates or not (_has_value(actual_dates.get("first_ts")) or _has_value(actual_dates.get("last_ts"))):
        raw_samples = ((summary.get("evidence") or {}).get("raw_samples") if isinstance(summary.get("evidence"), dict) else {}) or {}
        first_ts = raw_samples.get("first_ts")
        last_ts = raw_samples.get("last_ts")
        if _has_value(first_ts) or _has_value(last_ts):
            actual_dates = {"first_ts": first_ts, "last_ts": last_ts}

    responsible_groups = {
        "client_name": report_parties.get("client_name"),
        "facility_owner_operator_name": report_parties.get("facility_owner_operator_name"),
        "test_company_name": report_parties.get("test_company_name"),
        "session_operator_name": report_parties.get("session_operator_name"),
        "laboratory_name": report_parties.get("laboratory_name"),
        "observer_contacts": report_parties.get("observer_contacts"),
    }
    run_matrix_value = {
        "planned_runs": planned_run_rows,
        "actual_runs": actual_run_rows,
        "coverage_note": run_aggregation.get("coverage_note"),
    }
    block_status_counts = {"Available": 0, "Partial": 0, "Gap": 0}
    for block in (report_parties, process_control, deviations, correspondence, run_aggregation):
        status = str(block.get("status") or "")
        if status in block_status_counts:
            block_status_counts[status] += 1

    sections: Dict[str, Any] = {
        "cover_certification": {
            "report_title": _ctx_field(report_title, "session.project.project_name / session.project.job_id", "Available", "Working title for final report generation."),
            "facility_site": _ctx_field(project.get("site_facility"), "session.project.site_facility"),
            "unit_asset": _ctx_field(unit_asset, "session.source.manufacturer + session.source.model_number + session.project.asset_unit_id"),
            "job_id": _ctx_field(project.get("job_id"), "session.project.job_id"),
            "test_dates": _ctx_field(actual_dates, "summary.evidence.raw_samples.first_ts/last_ts", "Gap" if not actual_dates else "Partial", "Actual report-ready run timestamps are not yet normalized; this uses raw evidence timestamps when present."),
            "prepared_for": _ctx_field(report_parties.get("client_name"), report_parties.get("source") or "report.parties.client_name", "Partial" if _has_value(report_parties.get("facility_owner_operator_name")) else "Gap", "Uses normalized parties block. Client legal name is still missing unless separately captured."),
            "prepared_by": _ctx_field(report_parties.get("test_company_name"), report_parties.get("source") or "report.parties.test_company_name", report_parties.get("status"), "Uses normalized parties block."),
            "report_revision": _ctx_field("DRAFT / report_pack_v1", "report_context.defaults", "Partial", "Final revision management is not normalized yet."),
            "issue_date": _ctx_field(str(summary.get("generated_iso") or "")[:10], "summary.generated_iso", "Available"),
            "responsible_official": _ctx_field(report_parties.get("responsible_official_name"), report_parties.get("source") or "report.parties.responsible_official", report_parties.get("status"), "Responsible official / signatory metadata flows from the normalized parties block."),
        },
        "introduction": {
            "purpose_objective": _ctx_field(
                intake.get("regulatory_purpose"),
                "session.project.intake.regulatory_purpose",
                "Partial" if _has_value(intake.get("regulatory_purpose")) else "Gap",
                "Structured regulatory-purpose fields exist, but a report-ready objective narrative is not generated yet.",
            ),
            "responsible_groups": _ctx_field(
                responsible_groups,
                report_parties.get("source") or "report.parties",
                report_parties.get("status"),
                "Uses the normalized parties block. Missing legal/client/laboratory/observer details remain explicit gaps there.",
            ),
            "facility_source_identification": _ctx_field(
                {
                    "facility_site": project.get("site_facility"),
                    "source_category": source.get("source_category"),
                    "application": source.get("application"),
                    "service_class": source.get("service_class"),
                    "asset_unit_id": project.get("asset_unit_id"),
                    "coordinates": {"lat": location.get("lat"), "lon": location.get("lon"), "datum": location.get("datum")},
                    "elevation_ft_msl": location.get("elevation_ft_msl"),
                },
                "session.project + session.source + session.site_conditions.location",
                "Available",
            ),
            "pollutants_and_methods": _ctx_field(
                pollutant_rows,
                "session.pollutants.prescriptions + session.regulatory.limits",
                "Available" if pollutant_rows else "Gap",
            ),
            "test_schedule": _ctx_field(
                run_rows,
                "session.test_matrix.plan",
                run_schedule_status,
                "Planned run schedule is available from the Test Matrix. Actual run timestamps are not normalized yet." if run_rows else "No planned run schedule found in Test Matrix.",
            ),
        },
        "plant_and_sampling_location": {
            "process_source_description": _ctx_field(
                process_control.get("process_description"),
                process_control.get("source") or "session.source + session.fuel",
                process_control.get("status"),
            ),
            "process_narrative": _ctx_field(process_control.get("process_narrative_seed"), process_control.get("source") or "report.process.description", process_control.get("status"), "Uses normalized process/source metadata to build a narrative seed."),
            "control_equipment_description": _ctx_field(process_control.get("control_equipment"), process_control.get("source") or "report.control_equipment", "Partial", "Control equipment is normalized as an explicit placeholder block until detailed control metadata is captured."),
            "stack_sampling_location": _ctx_field(
                {
                    "shape": stack.get("shape"),
                    "diameter_in": stack.get("diameter_in"),
                    "width_in": stack.get("width_in"),
                    "height_in": stack.get("height_in"),
                    "port_count": stack.get("port_count"),
                    "port_angles_deg": stack.get("port_angles_deg"),
                    "port_height_ft_agl": stack.get("port_height_ft_agl"),
                    "upstream_diameters": stack.get("upstream_diameters"),
                    "downstream_diameters": stack.get("downstream_diameters"),
                    "traverse_scheme": stack.get("traverse_scheme"),
                },
                "session.source.stack",
                "Available",
            ),
            "sampling_location_adequacy": _ctx_field(
                {
                    "traverse_scheme": stack.get("traverse_scheme"),
                    "upstream_diameters": stack.get("upstream_diameters"),
                    "downstream_diameters": stack.get("downstream_diameters"),
                    "unstratified_7e": stack.get("unstratified_7e"),
                },
                "session.source.stack",
                "Partial",
                "Geometry inputs exist, but no final adequacy narrative or reviewer statement is generated yet.",
            ),
        },
        "summary_and_results": {
            "run_matrix": _ctx_field(
                run_matrix_value,
                run_aggregation.get("source") or "session.test_matrix.plan",
                run_schedule_status,
                run_aggregation.get("coverage_note") or planned_schedule_note or "Run matrix uses planned Test Matrix rows until actual run aggregation is implemented.",
            ),
            "operating_conditions_summary": _ctx_field(
                {
                    "actual_runs": actual_run_rows,
                    "weather_context": {
                        "mode": site.get("mode"),
                        "provider": weather.get("provider"),
                        "station_id": weather.get("station_id"),
                    },
                } if actual_run_rows else None,
                run_aggregation.get("source") or "aggregated run metrics",
                "Partial" if actual_run_rows else "Gap",
                "Actual run timing is normalized from worksteps. Detailed averaged operating metrics are still pending.",
            ),
            "results_summary": _ctx_field(
                regulatory_rows,
                "summary.regulatory_snapshot.rows",
                results_status,
                "Regulatory comparison rows are available, but a full final-report results table and average aggregation are not normalized yet.",
            ),
            "compliance_discussion_basis": _ctx_field(
                {
                    "selected_rule_ids": rules,
                    "overall_pass": (summary.get("session") or {}).get("overall_pass"),
                    "regulatory_status": (summary.get("regulatory_snapshot") or {}).get("status"),
                },
                "session.regulatory + summary.session + summary.regulatory_snapshot",
                "Partial" if rules else "Gap",
                "Rules and comparison output exist, but no final compliance narrative is generated yet.",
            ),
            "ftir_validation_summary": _ctx_field(
                {
                    "status": ftir_validation.get("status"),
                    "overall_status": ftir_validation.get("overall_status"),
                    "source": ftir_validation.get("source"),
                    "validation_mode": ((ftir_validation.get("config") or {}).get("validation_mode") if isinstance(ftir_validation.get("config"), dict) else None),
                    "comparator_method": ((ftir_validation.get("config") or {}).get("comparator_method") if isinstance(ftir_validation.get("config"), dict) else None),
                    "comparison_set_count": ftir_validation.get("comparison_set_count"),
                    "included_comparison_set_count": ftir_validation.get("included_comparison_set_count"),
                    "excluded_comparison_set_count": ftir_validation.get("excluded_comparison_set_count"),
                    "comparison_sets": ftir_validation.get("comparison_sets"),
                    "paired_window_count": ftir_validation.get("paired_window_count"),
                    "excluded_count": ftir_validation.get("excluded_count"),
                    "coverage_note": ftir_validation.get("coverage_note"),
                    "review_notes": ftir_validation.get("review_notes"),
                    "reviewer": ftir_validation.get("reviewer"),
                    "review_locked": ftir_validation.get("review_locked"),
                    "review_lock_by": ftir_validation.get("review_lock_by"),
                    "review_lock_iso": ftir_validation.get("review_lock_iso"),
                    "review_unlock_by": ftir_validation.get("review_unlock_by"),
                    "review_unlock_iso": ftir_validation.get("review_unlock_iso"),
                    "review_snapshot": ftir_validation.get("review_snapshot"),
                    "signoff": ftir_validation.get("signoff"),
                    "qa": ftir_validation.get("qa"),
                    "excluded_rows": ftir_validation.get("excluded_rows"),
                    "method301": ftir_validation.get("method301"),
                } if ftir_validation else None,
                "summary.ftir_validation",
                str(ftir_validation.get("status") or "Gap") if isinstance(ftir_validation, dict) else "Gap",
                "Session-scoped FTIR side-by-side validation summary and Method 301 comparison statistics.",
            ),
        },
        "sampling_and_analytical_procedures": {
            "methods_used": _ctx_field(
                pollutant_rows,
                "session.pollutants.prescriptions",
                "Available" if pollutant_rows else "Gap",
            ),
            "sample_system_and_analyzers": _ctx_field(
                {
                    "exhaust_flow_method": ((source.get("exhaust_flow") or {}).get("method") if isinstance(source.get("exhaust_flow"), dict) else None),
                    "fuel_flow_basis": ((source.get("fuel_flow") or {}).get("basis") if isinstance(source.get("fuel_flow"), dict) else None),
                    "channels": [
                        {
                            "pollutant": row.get("pollutant"),
                            "instrument": row.get("instrument"),
                            "units": row.get("units"),
                        }
                        for row in pollutant_rows
                    ],
                },
                "session.source + session.pollutants.prescriptions",
                "Partial" if pollutant_rows else "Gap",
                "Analyzer and sample-system fields exist, but no narrative / schematic reference binding is generated yet.",
            ),
            "calibration_gas_summary": _ctx_field(
                [
                    {
                        "pollutant": code,
                        "zero_cylinder_id": str((cfg.get("zero_cylinder_id") or "")),
                        "span_cylinder_id": str((cfg.get("span_cylinder_id") or "")),
                        "mid_cylinder_id": str((cfg.get("mid_cylinder_id") or "")),
                    }
                    for code, cfg in sorted(_pollutant_prescriptions(session).items())
                    if isinstance(cfg, dict)
                ],
                "session.pollutants.prescriptions[*].*_cylinder_id",
                "Partial",
                "Cylinder IDs exist per pollutant, but certificate linkage and unified calibration-gas tables are not normalized yet.",
            ),
            "deviations_alternatives_approvals": _ctx_field(
                deviations,
                deviations.get("source") or "report.deviations",
                deviations.get("status"),
                "Uses the normalized deviations / approvals block.",
            ),
        },
        "qaqc_activities": {
            "pre_test_checks": _ctx_field(
                (summary.get("test_matrix") or {}).get("tests"),
                "summary.test_matrix.tests",
                "Available" if _has_value((summary.get("test_matrix") or {}).get("tests")) else "Gap",
            ),
            "during_test_qaqc": _ctx_field(
                {
                    "analyzer_validity": summary.get("analyzer_validity"),
                    "reference_audit": summary.get("reference_audit"),
                    "side_by_side": ((summary.get("qaqc") or {}).get("side_by_side") if isinstance(summary.get("qaqc"), dict) else {}),
                    "ftir_validation": summary.get("ftir_validation"),
                    "weather_traceability": {
                        "mode": site.get("mode"),
                        "provider": weather.get("provider"),
                        "station_id": weather.get("station_id"),
                        "last_fetch_iso": weather.get("last_fetch_iso"),
                    },
                },
                "summary.analyzer_validity + summary.reference_audit + summary.qaqc.side_by_side + summary.ftir_validation + session.site_conditions.weather_compile",
                "Available",
            ),
            "post_test_checks": _ctx_field(
                {
                    "pollutant_adjustments": summary.get("pollutant_adjustments"),
                    "spike_recovery": ((summary.get("qaqc") or {}).get("spike_recovery") if isinstance(summary.get("qaqc"), dict) else {}),
                },
                "summary.pollutant_adjustments + summary.qaqc.spike_recovery",
                "Available",
            ),
            "qaqc_exceptions": _ctx_field(
                {
                    "failed_test_matrix_items": [
                        row for row in ((summary.get("test_matrix") or {}).get("tests") or [])
                        if isinstance(row, dict) and str(row.get("status") or "").upper() == "FAIL"
                    ],
                    "degraded_pollutants": [
                        {"pollutant": code, "state": row.get("state"), "reason": row.get("reason")}
                        for code, row in (((summary.get("analyzer_validity") or {}).get("per_pollutant") or {}).items())
                        if isinstance(row, dict) and str(row.get("state") or "").upper() not in {"VALID", "OK"}
                    ],
                },
                "summary.test_matrix.tests + summary.analyzer_validity.per_pollutant",
                "Partial",
                "Exception evidence exists, but narrative collation and impact statements are not generated yet.",
            ),
        },
        "appendices": {
            "appendix_manifest": _ctx_field(
                appendix_entries,
                "report_pack outputs + evidence bundle sources/static artifacts",
                "Partial" if appendix_entries else "Gap",
                "Manifest entries are generated from current pack outputs and evidence with explicit appendix titles and assignment basis.",
            ),
            "regulatory_and_correspondence": _ctx_field(
                correspondence,
                correspondence.get("source") or "report.correspondence",
                correspondence.get("status"),
                "Uses the normalized correspondence block.",
            ),
        },
        "implementation_gaps": _ctx_field(
            [
                {
                    "id": "report_parties_and_certification_metadata",
                    "status": report_parties.get("status"),
                    "note": "Parties are now normalized into a structured block, but client/legal signatory metadata is still incomplete.",
                },
                {
                    "id": "process_and_control_narrative",
                    "status": process_control.get("status"),
                    "note": "Process/source description is normalized, but control-equipment detail remains an explicit placeholder.",
                },
                {
                    "id": "deviations_and_approvals",
                    "status": deviations.get("status"),
                    "note": "Deviation and approval candidates are now normalized from session notes and evidence attachments, but formal approval records are still incomplete.",
                },
                {
                    "id": "regulatory_correspondence_tracking",
                    "status": correspondence.get("status"),
                    "note": "Correspondence candidates are now normalized from intake evidence and static artifacts, but formal tracking fields remain incomplete.",
                },
                {
                    "id": "appendix_assignment_model",
                    "status": "Partial",
                    "note": "A manifest is now generated with appendix titles, inclusion flags, and assignment basis, but the assignment rules are still heuristic.",
                },
                {
                    "id": "run_summary_aggregation",
                    "status": run_aggregation.get("status"),
                    "note": "Run summaries are now aggregated from worksteps, but per-run averaged operating metrics are still not normalized.",
                },
                {
                    "id": "ftir_side_by_side_validation",
                    "status": ftir_validation.get("status") if isinstance(ftir_validation, dict) else "Gap",
                    "note": "FTIR validation is session-scoped and exportable, but formal Method 301 acceptance still depends on complete time-matched evidence and six valid comparison sets.",
                },
            ],
            "report-crosswalk-2026-04-09",
            "Partial",
            "These are the tracked implementation gaps remaining between report-pack output and the master final report structure.",
        ),
    }

    normalized_blocks = {
        "report_parties": report_parties,
        "process_control": process_control,
        "deviations_approvals": deviations,
        "regulatory_correspondence": correspondence,
        "run_aggregation": run_aggregation,
        "ftir_validation": ftir_validation,
    }
    coverage = _ctx_counts(sections)
    return {
        "contract_version": "report_context_v1",
        "generated_iso": _now_iso(),
        "template_contract": _report_template_contract(),
        "session_identity": {
            "job_id": project.get("job_id"),
            "run_id": (summary.get("session") or {}).get("run_id"),
            "session_dir": str(session_dir),
            "config_path": str(cfg_path),
            "report_pack_dir": str(paths.out_dir),
        },
        "coverage": {
            "field_status_counts": coverage,
            "available_fields": coverage.get("Available", 0),
            "partial_fields": coverage.get("Partial", 0),
            "gap_fields": coverage.get("Gap", 0),
            "normalized_block_status_counts": block_status_counts,
        },
        "sections": sections,
        "normalized_blocks": normalized_blocks,
    }


def _md_scalar(value: Any) -> str:
    if value is None:
        return "N/A"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        return _fmt_num(value, 3, "")
    if isinstance(value, (int,)):
        return str(value)
    if isinstance(value, str):
        text = value.strip()
        return text if text else "N/A"
    if isinstance(value, list):
        parts = [_md_scalar(item) for item in value if _has_value(item)]
        return ", ".join([p for p in parts if p and p != "N/A"]) or "N/A"
    if isinstance(value, dict):
        parts: List[str] = []
        for key, item in value.items():
            if not _has_value(item):
                continue
            if isinstance(item, (dict, list)):
                continue
            label = str(key).replace("_", " ").strip()
            parts.append(f"{label}: {_md_scalar(item)}")
        return "; ".join(parts) if parts else "N/A"
    return str(value)


def _md_table(headers: List[str], rows: List[List[Any]]) -> str:
    safe_headers = [str(h or "").strip() or " " for h in headers]
    lines = [
        "| " + " | ".join(safe_headers) + " |",
        "| " + " | ".join(["---"] * len(safe_headers)) + " |",
    ]
    for row in rows or []:
        cells = [_md_scalar(cell).replace("\n", " ").replace("|", "\\|") for cell in row]
        lines.append("| " + " | ".join(cells) + " |")
    if len(lines) == 2:
        lines.append("| " + " | ".join(["N/A"] * len(safe_headers)) + " |")
    return "\n".join(lines)


def _write_final_report_markdown(
    paths: "ReportPackPaths",
    report_context: Dict[str, Any],
    summary: Dict[str, Any],
) -> None:
    _ensure_dir(paths.final_report_dir)
    sections = report_context.get("sections") if isinstance(report_context.get("sections"), dict) else {}
    blocks = report_context.get("normalized_blocks") if isinstance(report_context.get("normalized_blocks"), dict) else {}
    cover = sections.get("cover_certification") if isinstance(sections.get("cover_certification"), dict) else {}
    intro = sections.get("introduction") if isinstance(sections.get("introduction"), dict) else {}
    plant = sections.get("plant_and_sampling_location") if isinstance(sections.get("plant_and_sampling_location"), dict) else {}
    results = sections.get("summary_and_results") if isinstance(sections.get("summary_and_results"), dict) else {}
    procedures = sections.get("sampling_and_analytical_procedures") if isinstance(sections.get("sampling_and_analytical_procedures"), dict) else {}
    qaqc = sections.get("qaqc_activities") if isinstance(sections.get("qaqc_activities"), dict) else {}
    appendices = sections.get("appendices") if isinstance(sections.get("appendices"), dict) else {}

    def _ctx_value(container: Dict[str, Any], key: str) -> Any:
        node = container.get(key) if isinstance(container.get(key), dict) else {}
        return node.get("value")

    pollutant_rows = _ctx_value(intro, "pollutants_and_methods") or []
    run_matrix = _ctx_value(results, "run_matrix") or {}
    actual_runs = run_matrix.get("actual_runs") if isinstance(run_matrix, dict) else []
    planned_runs = run_matrix.get("planned_runs") if isinstance(run_matrix, dict) else []
    run_rows = actual_runs or planned_runs or []
    results_rows = _ctx_value(results, "results_summary") or []
    appendix_rows = _ctx_value(appendices, "appendix_manifest") or []
    methods_rows = _ctx_value(procedures, "methods_used") or []
    deviations = _ctx_value(procedures, "deviations_alternatives_approvals") or {}
    correspondence = _ctx_value(appendices, "regulatory_and_correspondence") or {}
    parties = blocks.get("report_parties") if isinstance(blocks.get("report_parties"), dict) else {}
    process_control = blocks.get("process_control") if isinstance(blocks.get("process_control"), dict) else {}
    run_aggregation = blocks.get("run_aggregation") if isinstance(blocks.get("run_aggregation"), dict) else {}
    ftir_validation = blocks.get("ftir_validation") if isinstance(blocks.get("ftir_validation"), dict) else {}
    template_contract = report_context.get("template_contract") if isinstance(report_context.get("template_contract"), dict) else {}
    session_identity = report_context.get("session_identity") if isinstance(report_context.get("session_identity"), dict) else {}
    ftir_validation_summary = _ctx_value(results, "ftir_validation_summary") or {}

    pollutant_table_rows = []
    for row in pollutant_rows:
        if not isinstance(row, dict):
            continue
        basis = row.get("compliance_basis") if isinstance(row.get("compliance_basis"), list) else []
        basis_text = "; ".join(
            [
                " ".join(
                    [
                        _md_scalar(item.get("limit_type")),
                        _md_scalar(item.get("value")),
                        _md_scalar(item.get("units")),
                        f"({ _md_scalar(item.get('basis')) })" if _has_value(item.get("basis")) else "",
                    ]
                ).strip()
                for item in basis
                if isinstance(item, dict)
            ]
        ) or "N/A"
        pollutant_table_rows.append([
            row.get("pollutant"),
            row.get("method"),
            row.get("instrument"),
            row.get("units"),
            basis_text,
        ])

    run_table_rows = []
    for row in run_rows:
        if not isinstance(row, dict):
            continue
        methods_text = ", ".join([str(pol.get("pollutant") or "") for pol in pollutant_rows if isinstance(pol, dict)]) or "N/A"
        run_table_rows.append([
            row.get("run_no"),
            row.get("date"),
            row.get("start_ts_iso") or row.get("start"),
            row.get("end_ts_iso") or row.get("stop"),
            methods_text,
            row.get("operating_level"),
            row.get("status"),
        ])

    results_table_rows = []
    for row in results_rows:
        if not isinstance(row, dict):
            continue
        results_table_rows.append([
            row.get("pollutant") or row.get("code"),
            row.get("run1") or row.get("current_value"),
            row.get("run2"),
            row.get("run3"),
            row.get("avg") or row.get("average"),
            row.get("units"),
            row.get("limit_value") or row.get("limit"),
            row.get("status"),
        ])

    methods_table_rows = []
    for row in methods_rows:
        if not isinstance(row, dict):
            continue
        methods_table_rows.append([
            row.get("pollutant"),
            row.get("method"),
            "Current session basis",
            "Yes" if _has_value((deviations.get("candidate_evidence_titles") if isinstance(deviations, dict) else [])) else "No",
            "Appendix C",
        ])

    appendix_table_rows = []
    for row in appendix_rows:
        if not isinstance(row, dict):
            continue
        appendix_table_rows.append([
            row.get("appendix"),
            row.get("appendix_title"),
            row.get("title"),
            row.get("artifact_type"),
            row.get("status"),
            row.get("path"),
            row.get("sha256"),
        ])

    lines: List[str] = []
    lines.append(f"# {_md_scalar(_ctx_value(cover, 'report_title'))}")
    lines.append("")
    lines.append("Generated from `report_context_v1` and the MOLE DAS master test report template contract.")
    lines.append("")
    lines.append("## Cover / Certification")
    lines.append("")
    lines.append(f"- Facility / site: {_md_scalar(_ctx_value(cover, 'facility_site'))}")
    lines.append(f"- Unit / asset: {_md_scalar(_ctx_value(cover, 'unit_asset'))}")
    lines.append(f"- Job ID: {_md_scalar(_ctx_value(cover, 'job_id'))}")
    lines.append(f"- Test dates: {_md_scalar(_ctx_value(cover, 'test_dates'))}")
    lines.append(f"- Prepared for: {_md_scalar(_ctx_value(cover, 'prepared_for'))}")
    lines.append(f"- Prepared by: {_md_scalar(_ctx_value(cover, 'prepared_by'))}")
    lines.append(f"- Report revision / date: {_md_scalar(_ctx_value(cover, 'report_revision'))} / {_md_scalar(_ctx_value(cover, 'issue_date'))}")
    responsible_official = _md_scalar(_ctx_value(cover, "responsible_official"))
    responsible_title = _md_scalar(parties.get("responsible_official_title"))
    if responsible_title != "N/A":
        responsible_official = f"{responsible_official} ({responsible_title})" if responsible_official != "N/A" else responsible_title
    lines.append(f"- Responsible official: {responsible_official}")
    lines.append("")
    lines.append("## 1. Introduction")
    lines.append("")
    lines.append("### 1.1 Purpose and objective")
    lines.append("")
    lines.append(f"- Regulatory purpose: {_md_scalar(_ctx_value(intro, 'purpose_objective'))}")
    lines.append("")
    lines.append("### 1.2 Responsible groups")
    lines.append("")
    lines.append(f"- Facility owner / operator: {_md_scalar(parties.get('facility_owner_operator_name'))}")
    lines.append(f"- Test company: {_md_scalar(parties.get('test_company_name'))}")
    lines.append(f"- Laboratory: {_md_scalar(parties.get('laboratory_name'))}")
    lines.append(f"- Observer contacts: {_md_scalar(parties.get('observer_contacts'))}")
    lines.append("")
    lines.append("### 1.3 Facility and source identification")
    lines.append("")
    lines.append(f"- Source identification: {_md_scalar(_ctx_value(intro, 'facility_source_identification'))}")
    lines.append("")
    lines.append("### 1.4 Pollutants and methods summary")
    lines.append("")
    lines.append(_md_table(
        ["Pollutant", "Method", "Analyzer / Instrument", "Units", "Compliance basis"],
        pollutant_table_rows,
    ))
    lines.append("")
    lines.append("### 1.5 Test dates and schedule")
    lines.append("")
    lines.append(f"- Schedule note: {_md_scalar((run_aggregation.get('coverage_note') if isinstance(run_aggregation, dict) else None))}")
    lines.append("")
    lines.append("## 2. Plant and Sampling Location Description")
    lines.append("")
    lines.append("### 2.1 Process and source description")
    lines.append("")
    lines.append(f"- Process/source description: {_md_scalar(process_control.get('process_description'))}")
    lines.append(f"- Process narrative: {_md_scalar(process_control.get('process_narrative_seed'))}")
    lines.append("")
    lines.append("### 2.2 Control equipment description")
    lines.append("")
    lines.append(f"- Control equipment: {_md_scalar(process_control.get('control_equipment'))}")
    lines.append("")
    lines.append("### 2.3 Stack / duct and sampling location description")
    lines.append("")
    stack_desc = _ctx_value(plant, "stack_sampling_location") or {}
    lines.append(_md_table(
        ["Item", "Value"],
        [
            ["Stack shape", stack_desc.get("shape") if isinstance(stack_desc, dict) else None],
            ["Diameter / width / height", _md_scalar({"diameter_in": stack_desc.get("diameter_in"), "width_in": stack_desc.get("width_in"), "height_in": stack_desc.get("height_in")}) if isinstance(stack_desc, dict) else None],
            ["Port count / angles", _md_scalar({"port_count": stack_desc.get("port_count"), "port_angles_deg": stack_desc.get("port_angles_deg")}) if isinstance(stack_desc, dict) else None],
            ["Port height AGL", stack_desc.get("port_height_ft_agl") if isinstance(stack_desc, dict) else None],
            ["Upstream disturbance distance", stack_desc.get("upstream_diameters") if isinstance(stack_desc, dict) else None],
            ["Downstream disturbance distance", stack_desc.get("downstream_diameters") if isinstance(stack_desc, dict) else None],
            ["Traverse basis", stack_desc.get("traverse_scheme") if isinstance(stack_desc, dict) else None],
        ],
    ))
    lines.append("")
    lines.append("### 2.4 Sampling-location adequacy narrative")
    lines.append("")
    lines.append(f"- Adequacy basis: {_md_scalar(_ctx_value(plant, 'sampling_location_adequacy'))}")
    lines.append("")
    lines.append("## 3. Summary and Discussion of Results")
    lines.append("")
    lines.append("### 3.1 Test matrix")
    lines.append("")
    lines.append(_md_table(
        ["Run", "Date", "Start", "Stop", "Pollutants / Methods", "Operating level", "Status"],
        run_table_rows,
    ))
    lines.append("")
    lines.append("### 3.2 Operating conditions summary")
    lines.append("")
    lines.append(f"- Operating summary source: {_md_scalar(_ctx_value(results, 'operating_conditions_summary'))}")
    lines.append("")
    lines.append("### 3.3 Results summary")
    lines.append("")
    lines.append(_md_table(
        ["Pollutant", "Run 1", "Run 2", "Run 3", "Average", "Units", "Limit", "Pass / Fail"],
        results_table_rows,
    ))
    lines.append("")
    lines.append("### 3.4 Discussion of results")
    lines.append("")
    lines.append(f"- Compliance discussion basis: {_md_scalar(_ctx_value(results, 'compliance_discussion_basis'))}")
    lines.append(f"- Pollutant adjustments summary: {_md_scalar(((summary.get('pollutant_adjustments') or {}).get('row_count') if isinstance(summary.get('pollutant_adjustments'), dict) else None))} adjustment row(s)")
    lines.append("")
    lines.append("### 3.5 FTIR side-by-side validation")
    lines.append("")
    lines.append(f"- Validation status: {_md_scalar((ftir_validation_summary.get('status') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Overall status: {_md_scalar((ftir_validation_summary.get('overall_status') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Validation source: {_md_scalar((ftir_validation_summary.get('source') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Validation mode: {_md_scalar((((ftir_validation.get('config') or {}) if isinstance(ftir_validation.get('config'), dict) else {}).get('validation_mode')))}")
    lines.append(f"- Comparator method: {_md_scalar((((ftir_validation.get('config') or {}) if isinstance(ftir_validation.get('config'), dict) else {}).get('comparator_method')))}")
    lines.append(f"- Vendor profile requested: {_md_scalar((((ftir_validation.get('config') or {}) if isinstance(ftir_validation.get('config'), dict) else {}).get('ftir_vendor_profile')))}")
    lines.append(f"- Vendor profile used: {_md_scalar((((ftir_validation.get('ftir_source') or {}) if isinstance(ftir_validation.get('ftir_source'), dict) else {}).get('vendor_profile_used')))}")
    lines.append(f"- Comparison set count: {_md_scalar((ftir_validation_summary.get('comparison_set_count') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Included comparison sets: {_md_scalar((ftir_validation_summary.get('included_comparison_set_count') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Excluded comparison sets: {_md_scalar((ftir_validation_summary.get('excluded_comparison_set_count') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Paired window count: {_md_scalar((ftir_validation_summary.get('paired_window_count') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Excluded row count: {_md_scalar((ftir_validation_summary.get('excluded_count') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Coverage note: {_md_scalar((ftir_validation_summary.get('coverage_note') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Acceptance basis: {_md_scalar((ftir_validation_summary.get('acceptance_basis') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Acceptance recommendation: {_md_scalar((ftir_validation_summary.get('acceptance_recommended_decision') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Acceptance basis note: {_md_scalar((ftir_validation_summary.get('acceptance_basis_note') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Reviewer: {_md_scalar((ftir_validation_summary.get('reviewer') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Reviewer notes: {_md_scalar((ftir_validation_summary.get('review_notes') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Review locked: {_md_scalar((ftir_validation_summary.get('review_locked') if isinstance(ftir_validation_summary, dict) else None))}")
    lines.append(f"- Lock by / at: {_md_scalar({'by': (ftir_validation_summary.get('review_lock_by') if isinstance(ftir_validation_summary, dict) else None), 'at': (ftir_validation_summary.get('review_lock_iso') if isinstance(ftir_validation_summary, dict) else None)})}")
    lines.append(f"- Last unlock by / at: {_md_scalar({'by': (ftir_validation_summary.get('review_unlock_by') if isinstance(ftir_validation_summary, dict) else None), 'at': (ftir_validation_summary.get('review_unlock_iso') if isinstance(ftir_validation_summary, dict) else None)})}")
    lines.append(f"- Locked snapshot: {_md_scalar((ftir_validation_summary.get('review_snapshot') if isinstance(ftir_validation_summary, dict) else None))}")
    qa = (ftir_validation_summary.get("qa") if isinstance(ftir_validation_summary, dict) else {}) or {}
    lines.append(f"- QA summary: {_md_scalar((qa.get('summary') if isinstance(qa, dict) else None))}")
    lines.append(f"- QA lock ready: {_md_scalar((qa.get('lock_ready') if isinstance(qa, dict) else None))}")
    lines.append(f"- QA signoff ready: {_md_scalar((qa.get('signoff_ready') if isinstance(qa, dict) else None))}")
    lines.append(f"- QA blocking issue count: {_md_scalar((len(list(qa.get('blocking_issues') or [])) if isinstance(qa, dict) else None))}")
    lines.append(f"- QA warning count: {_md_scalar((len(list(qa.get('warnings') or [])) if isinstance(qa, dict) else None))}")
    signoff = (ftir_validation_summary.get("signoff") if isinstance(ftir_validation_summary, dict) else {}) or {}
    lines.append(f"- Signoff decision: {_md_scalar((signoff.get('decision') if isinstance(signoff, dict) else None))}")
    lines.append(f"- Signoff basis: {_md_scalar((signoff.get('basis') if isinstance(signoff, dict) else None))}")
    lines.append(f"- Signoff by / role / at: {_md_scalar({'by': (signoff.get('by') if isinstance(signoff, dict) else None), 'role': (signoff.get('role') if isinstance(signoff, dict) else None), 'at': (signoff.get('iso') if isinstance(signoff, dict) else None)})}")
    lines.append(f"- Signoff note: {_md_scalar((signoff.get('note') if isinstance(signoff, dict) else None))}")
    lines.append(f"- FTIR appendix package: {_md_scalar((((summary.get('ftir_validation') or {}).get('appendix')) if isinstance(summary.get('ftir_validation'), dict) else None))}")
    method301_rows = []
    for row in list((ftir_validation_summary.get("method301") if isinstance(ftir_validation_summary, dict) else []) or []):
        if not isinstance(row, dict):
            continue
        method301_rows.append([
            row.get("analyte"),
            row.get("paired_window_count"),
            row.get("excluded_window_count"),
            _fmt_num(row.get("relative_bias_pct"), 3, ""),
            _fmt_num(row.get("t_statistic"), 3, ""),
            _fmt_num(row.get("f_statistic"), 3, ""),
            row.get("bias_status"),
            row.get("precision_status"),
            row.get("overall_status"),
        ])
    lines.append(_md_table(
        ["Analyte", "Paired windows", "Excluded", "Relative bias %", "t-stat", "F-stat", "Bias", "Precision", "Overall"],
        method301_rows,
    ))
    comparison_sets = list((ftir_validation_summary.get("comparison_sets") if isinstance(ftir_validation_summary, dict) else []) or [])
    if comparison_sets:
        lines.append("")
        lines.append("#### FTIR Comparison Sets")
        lines.append("")
        comparison_set_rows = []
        for row in comparison_sets:
            if not isinstance(row, dict):
                continue
            comparison_set_rows.append([
                row.get("set_no"),
                row.get("run_no"),
                row.get("window_start_iso"),
                row.get("window_end_iso"),
                row.get("analytes"),
                row.get("inclusion_status"),
                row.get("formal_basis"),
                row.get("review_state"),
                row.get("note"),
            ])
        lines.append(_md_table(
            ["Set", "Run", "Start", "Stop", "Analytes", "Status", "Basis", "Review state", "Note"],
            comparison_set_rows,
        ))
    excluded_rows = list((ftir_validation_summary.get("excluded_rows") if isinstance(ftir_validation_summary, dict) else []) or [])
    if excluded_rows:
        lines.append("")
        lines.append("#### Excluded FTIR Validation Rows")
        lines.append("")
        exclusion_table_rows = []
        for row in excluded_rows:
            if not isinstance(row, dict):
                continue
            exclusion_table_rows.append([
                row.get("run_no"),
                row.get("label"),
                row.get("analyte"),
                row.get("reason"),
                row.get("reviewer"),
                row.get("updated_iso"),
            ])
        lines.append(_md_table(
            ["Run", "Label", "Analyte", "Reason", "Reviewer", "Updated"],
            exclusion_table_rows,
        ))
    lines.append("")
    lines.append("## 4. Sampling and Analytical Procedures")
    lines.append("")
    lines.append("### 4.1 Test methods used")
    lines.append("")
    lines.append(_md_table(
        ["Pollutant / Parameter", "Method", "Basis / Version", "Deviations?", "Reference appendix"],
        methods_table_rows,
    ))
    lines.append("")
    lines.append("### 4.2 Sampling system and analyzer configuration")
    lines.append("")
    lines.append(f"- Sampling system / analyzers: {_md_scalar(_ctx_value(procedures, 'sample_system_and_analyzers'))}")
    lines.append("")
    lines.append("### 4.3 Analytical procedure details")
    lines.append("")
    lines.append(f"- Calibration gas summary: {_md_scalar(_ctx_value(procedures, 'calibration_gas_summary'))}")
    lines.append("")
    lines.append("### 4.4 Deviations, alternatives, and approvals")
    lines.append("")
    lines.append(f"- Planned deviations: {_md_scalar((deviations.get('planned_deviations') if isinstance(deviations, dict) else None))}")
    lines.append(f"- Field deviations: {_md_scalar((deviations.get('field_deviations') if isinstance(deviations, dict) else None))}")
    lines.append(f"- Alternative method / approval references: {_md_scalar((deviations.get('alternative_method_approvals') if isinstance(deviations, dict) else None))}")
    lines.append(f"- Impact statement: {_md_scalar((deviations.get('impact_statement') if isinstance(deviations, dict) else None))}")
    lines.append(f"- Candidate evidence titles: {_md_scalar((deviations.get('candidate_evidence_titles') if isinstance(deviations, dict) else None))}")
    lines.append("")
    lines.append("## 5. QA/QC Activities")
    lines.append("")
    lines.append("### 5.1 Pre-test checks")
    lines.append("")
    lines.append(f"- Pre-test checks: {_md_scalar(_ctx_value(qaqc, 'pre_test_checks'))}")
    lines.append("")
    lines.append("### 5.2 During-test QA/QC")
    lines.append("")
    lines.append(f"- During-test QA/QC: {_md_scalar(_ctx_value(qaqc, 'during_test_qaqc'))}")
    lines.append("")
    lines.append("### 5.3 Post-test checks")
    lines.append("")
    lines.append(f"- Post-test checks: {_md_scalar(_ctx_value(qaqc, 'post_test_checks'))}")
    lines.append("")
    lines.append("### 5.4 QA/QC exceptions")
    lines.append("")
    lines.append(f"- QA/QC exceptions: {_md_scalar(_ctx_value(qaqc, 'qaqc_exceptions'))}")
    lines.append("")
    lines.append("## 6. Appendices")
    lines.append("")
    lines.append(f"- Notice of intent date: {_md_scalar((correspondence.get('notice_of_intent_date') if isinstance(correspondence, dict) else None))}")
    lines.append(f"- Agency contact: {_md_scalar((correspondence.get('agency_contact') if isinstance(correspondence, dict) else None))}")
    lines.append(f"- Approval dates: {_md_scalar((correspondence.get('approval_dates') if isinstance(correspondence, dict) else None))}")
    lines.append(f"- Submission status: {_md_scalar((correspondence.get('submission_status') if isinstance(correspondence, dict) else None))}")
    lines.append(f"- Correspondence notes: {_md_scalar((correspondence.get('notes') if isinstance(correspondence, dict) else None))}")
    lines.append("")
    lines.append(_md_table(
        ["Appendix", "Appendix title", "Document title", "Type", "Status", "Path", "SHA-256"],
        appendix_table_rows,
    ))
    lines.append("")
    lines.append("## Generation Metadata")
    lines.append("")
    lines.append(f"- Report context contract: {_md_scalar(report_context.get('contract_version'))}")
    lines.append(f"- Master template path: {_md_scalar((((template_contract.get('master_template') or {}) if isinstance(template_contract.get('master_template'), dict) else {}).get('path')))}")
    lines.append(f"- Crosswalk path: {_md_scalar((((template_contract.get('crosswalk') or {}) if isinstance(template_contract.get('crosswalk'), dict) else {}).get('path')))}")
    lines.append(f"- Run ID: {_md_scalar(session_identity.get('run_id'))}")
    lines.append(f"- Session dir: {_md_scalar(session_identity.get('session_dir'))}")
    lines.append(f"- Regulatory correspondence block: {_md_scalar(correspondence)}")
    lines.append("")
    paths.final_report_md.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def _strip_markdown_text(text: Any) -> str:
    s = str(text or "")
    s = s.replace("`", "")
    s = s.replace("**", "")
    s = s.replace("__", "")
    return re.sub(r"\s+", " ", s).strip()


def _parse_markdown_table(lines: List[str]) -> Tuple[List[str], List[List[str]]]:
    rows: List[List[str]] = []
    for raw in lines:
        line = str(raw or "").strip()
        if not line.startswith("|"):
            continue
        parts = [p.strip() for p in line.strip("|").split("|")]
        rows.append([_strip_markdown_text(p) for p in parts])
    if not rows:
        return [], []
    headers = rows[0]
    body: List[List[str]] = []
    for row in rows[1:]:
        if row and all(re.fullmatch(r"[:\- ]+", str(cell or "")) for cell in row):
            continue
        body.append(row)
    return headers, body


def _parse_markdown_report_blocks(markdown_text: str) -> List[Dict[str, Any]]:
    lines = markdown_text.splitlines()
    blocks: List[Dict[str, Any]] = []
    i = 0
    while i < len(lines):
        raw = str(lines[i] or "")
        stripped = raw.strip()
        if not stripped:
            i += 1
            continue
        head = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if head:
            blocks.append({
                "type": "heading",
                "level": len(head.group(1)),
                "text": _strip_markdown_text(head.group(2)),
            })
            i += 1
            continue
        if stripped.startswith("|"):
            table_lines: List[str] = []
            while i < len(lines) and str(lines[i] or "").strip().startswith("|"):
                table_lines.append(str(lines[i] or ""))
                i += 1
            headers, rows = _parse_markdown_table(table_lines)
            if headers:
                blocks.append({"type": "table", "headers": headers, "rows": rows})
            continue
        if stripped.startswith("- "):
            items: List[str] = []
            while i < len(lines):
                probe = str(lines[i] or "").strip()
                if not probe.startswith("- "):
                    break
                items.append(_strip_markdown_text(probe[2:]))
                i += 1
            if items:
                blocks.append({"type": "bullet_list", "items": items})
            continue
        para_parts = [stripped]
        i += 1
        while i < len(lines):
            probe = str(lines[i] or "").strip()
            if not probe:
                break
            if probe.startswith("|") or probe.startswith("- ") or re.match(r"^(#{1,6})\s+(.+)$", probe):
                break
            para_parts.append(probe)
            i += 1
        blocks.append({"type": "paragraph", "text": _strip_markdown_text(" ".join(para_parts))})
    return blocks


def _remove_if_exists(path: Path) -> None:
    try:
        if path.exists():
            path.unlink()
    except Exception:
        pass


def _write_final_report_docx(paths: "ReportPackPaths", summary: Dict[str, Any]) -> Dict[str, Any]:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except Exception as exc:
        _remove_if_exists(paths.final_report_docx)
        return {"status": "skipped", "reason": f"python-docx not available: {exc}", "path": str(paths.final_report_docx)}

    try:
        markdown_text = paths.final_report_md.read_text(encoding="utf-8")
        blocks = _parse_markdown_report_blocks(markdown_text)
        doc = Document()
        for section in doc.sections:
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.7)
            section.right_margin = Inches(0.7)

        normal_style = doc.styles["Normal"]
        normal_style.font.name = "Calibri"
        normal_style.font.size = Pt(10)

        title_set = False
        for block in blocks:
            btype = block.get("type")
            if btype == "heading":
                level = int(block.get("level") or 1)
                text = str(block.get("text") or "").strip()
                if not text:
                    continue
                doc.add_heading(text, level=0 if (level == 1 and not title_set) else min(max(level, 1), 4))
                title_set = True
            elif btype == "paragraph":
                text = str(block.get("text") or "").strip()
                if text:
                    doc.add_paragraph(text)
            elif btype == "bullet_list":
                for item in block.get("items") or []:
                    txt = str(item or "").strip()
                    if txt:
                        doc.add_paragraph(txt, style="List Bullet")
            elif btype == "table":
                headers = [str(v or "") for v in (block.get("headers") or [])]
                rows = [list(r) for r in (block.get("rows") or [])]
                if not headers:
                    continue
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                hdr = table.rows[0].cells
                for idx, value in enumerate(headers):
                    hdr[idx].text = value
                for row in rows:
                    cells = table.add_row().cells
                    for idx in range(len(headers)):
                        cells[idx].text = str(row[idx] if idx < len(row) else "")
                doc.add_paragraph("")

        core = doc.core_properties
        core.title = "MOLE DAS Final Test Report"
        core.subject = "Formal deliverable generated from report_context_v1"
        core.author = str((((summary.get("session") or {}).get("operator")) if isinstance(summary.get("session"), dict) else "") or "MOLE DAS")
        doc.save(str(paths.final_report_docx))
        return {
            "status": "generated",
            "path": str(paths.final_report_docx),
            "bytes": paths.final_report_docx.stat().st_size if paths.final_report_docx.exists() else 0,
        }
    except Exception as exc:
        _remove_if_exists(paths.final_report_docx)
        return {"status": "failed", "reason": str(exc), "path": str(paths.final_report_docx)}


def _write_final_report_pdf(paths: "ReportPackPaths") -> Dict[str, Any]:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except Exception as exc:
        _remove_if_exists(paths.final_report_pdf)
        return {"status": "skipped", "reason": f"reportlab not available: {exc}", "path": str(paths.final_report_pdf)}

    try:
        markdown_text = paths.final_report_md.read_text(encoding="utf-8")
        blocks = _parse_markdown_report_blocks(markdown_text)
        styles = getSampleStyleSheet()
        h1 = ParagraphStyle("FinalReportH1", parent=styles["Title"], spaceAfter=10)
        h2 = ParagraphStyle("FinalReportH2", parent=styles["Heading2"], spaceBefore=8, spaceAfter=6)
        h3 = ParagraphStyle("FinalReportH3", parent=styles["Heading3"], spaceBefore=6, spaceAfter=4)
        body = ParagraphStyle("FinalReportBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=12, spaceAfter=4)
        bullet = ParagraphStyle("FinalReportBullet", parent=body, leftIndent=14, firstLineIndent=0)
        doc = SimpleDocTemplate(
            str(paths.final_report_pdf),
            pagesize=letter,
            leftMargin=0.6 * inch,
            rightMargin=0.6 * inch,
            topMargin=0.6 * inch,
            bottomMargin=0.6 * inch,
        )
        story: List[Any] = []
        title_set = False
        for block in blocks:
            btype = block.get("type")
            if btype == "heading":
                level = int(block.get("level") or 1)
                text = _xml_escape(str(block.get("text") or ""))
                if not text:
                    continue
                style = h1 if (level == 1 and not title_set) else (h2 if level <= 2 else h3)
                title_set = True
                story.append(Paragraph(text, style))
                story.append(Spacer(1, 0.08 * inch))
            elif btype == "paragraph":
                text = _xml_escape(str(block.get("text") or ""))
                if text:
                    story.append(Paragraph(text, body))
                    story.append(Spacer(1, 0.04 * inch))
            elif btype == "bullet_list":
                for item in block.get("items") or []:
                    txt = _xml_escape(str(item or ""))
                    if txt:
                        story.append(Paragraph(txt, bullet, bulletText="\u2022"))
                story.append(Spacer(1, 0.04 * inch))
            elif btype == "table":
                headers = [str(v or "") for v in (block.get("headers") or [])]
                rows = [list(r) for r in (block.get("rows") or [])]
                if not headers:
                    continue
                data: List[List[Any]] = [[Paragraph(_xml_escape(v), body) for v in headers]]
                for row in rows:
                    data.append([Paragraph(_xml_escape(str(row[idx] if idx < len(row) else "")), body) for idx in range(len(headers))])
                table = Table(data, repeatRows=1)
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d9d9d9")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#666666")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]))
                story.append(table)
                story.append(Spacer(1, 0.08 * inch))
        doc.build(story)
        return {
            "status": "generated",
            "path": str(paths.final_report_pdf),
            "bytes": paths.final_report_pdf.stat().st_size if paths.final_report_pdf.exists() else 0,
        }
    except Exception as exc:
        _remove_if_exists(paths.final_report_pdf)
        return {"status": "failed", "reason": str(exc), "path": str(paths.final_report_pdf)}


def _write_final_report_index(
    paths: "ReportPackPaths",
    report_context: Dict[str, Any],
    summary: Dict[str, Any],
    render_status: Optional[Dict[str, Any]] = None,
) -> None:
    _ensure_dir(paths.final_report_dir)
    files: List[Dict[str, Any]] = []
    for p in [paths.final_report_md, paths.final_report_docx, paths.final_report_pdf, paths.report_context_json]:
        if not p.exists():
            continue
        try:
            files.append({
                "name": p.name,
                "path": str(p),
                "sha256": _sha256_file(p),
                "bytes": p.stat().st_size,
            })
        except Exception:
            continue
    payload = {
        "generated_iso": summary.get("generated_iso"),
        "contract_version": "final_report_v1",
        "export_dir": str(paths.final_report_dir),
        "markdown_path": str(paths.final_report_md),
        "docx_path": str(paths.final_report_docx),
        "pdf_path": str(paths.final_report_pdf),
        "source_report_context_json": str(paths.report_context_json),
        "source_report_pack_dir": str(paths.out_dir),
        "coverage": report_context.get("coverage"),
        "template_contract": report_context.get("template_contract"),
        "render_status": render_status or {},
        "ftir_validation_appendix": ((summary.get("ftir_validation") or {}).get("appendix") if isinstance(summary.get("ftir_validation"), dict) else None),
        "files": files,
    }
    paths.final_report_index_json.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def _write_ftir_validation_appendix(paths: "ReportPackPaths", ftir_validation_summary: Dict[str, Any]) -> Dict[str, Any]:
    _ensure_dir(paths.ftir_validation_appendix_dir)
    ftir = dict(ftir_validation_summary or {}) if isinstance(ftir_validation_summary, dict) else {}
    signoff = ftir.get("signoff") if isinstance(ftir.get("signoff"), dict) else {}
    snap = ftir.get("review_snapshot") if isinstance(ftir.get("review_snapshot"), dict) else {}
    config = ftir.get("config") if isinstance(ftir.get("config"), dict) else {}
    comparison_sets = list(ftir.get("comparison_sets") or [])
    aligned_rows = list(ftir.get("aligned_rows") or [])
    method301_rows = list(ftir.get("method301") or [])
    excluded_rows = list(ftir.get("excluded_rows") or [])

    cover_lines = [
        "# FTIR Validation Appendix Package",
        "",
        "## Signoff Summary",
        "",
        f"- Validation status: {ftir.get('status') or '(n/a)'}",
        f"- Overall status: {ftir.get('overall_status') or '(n/a)'}",
        f"- Source: {ftir.get('source') or '(n/a)'}",
        f"- Validation mode: {config.get('validation_mode') or '(n/a)'}",
        f"- Comparator method: {config.get('comparator_method') or '(n/a)'}",
        f"- Vendor profile requested: {config.get('ftir_vendor_profile') or '(n/a)'}",
        f"- Vendor profile used: {(ftir.get('ftir_source') or {}).get('vendor_profile_used') if isinstance(ftir.get('ftir_source'), dict) else '(n/a)'}",
        f"- Acceptance basis: {ftir.get('acceptance_basis') or '(n/a)'}",
        f"- Acceptance recommendation: {ftir.get('acceptance_recommended_decision') or '(n/a)'}",
        f"- Acceptance basis note: {ftir.get('acceptance_basis_note') or '(n/a)'}",
        f"- Paired window count: {ftir.get('paired_window_count') or 0}",
        f"- Excluded row count: {ftir.get('excluded_count') or 0}",
        f"- Review locked: {ftir.get('review_locked')}",
        f"- Review lock by / at: {ftir.get('review_lock_by') or '(n/a)'} / {ftir.get('review_lock_iso') or '(n/a)'}",
        f"- Signoff decision: {signoff.get('decision') or 'UNSIGNED'}",
        f"- Signoff basis: {signoff.get('basis') or '(n/a)'}",
        f"- Signoff by / role / at: {signoff.get('by') or '(n/a)'} / {signoff.get('role') or '(n/a)'} / {signoff.get('iso') or '(n/a)'}",
        f"- Signoff note: {signoff.get('note') or '(n/a)'}",
        f"- Frozen snapshot JSON: {snap.get('json_path') or '(n/a)'}",
        f"- Frozen snapshot created by / at: {snap.get('snapshot_by') or '(n/a)'} / {snap.get('snapshot_iso') or '(n/a)'}",
        f"- Coverage note: {ftir.get('coverage_note') or '(n/a)'}",
        "",
        "## Appendix Artifacts",
        "",
        f"- Comparison ledger CSV: {paths.ftir_validation_appendix_ledger_csv}",
        f"- Method 301 stats CSV: {paths.ftir_validation_appendix_method301_csv}",
        f"- Exclusion register CSV: {paths.ftir_validation_appendix_exclusions_csv}",
        f"- Reviewer workbook XLSX: {paths.ftir_validation_appendix_workbook_xlsx}",
    ]
    paths.ftir_validation_appendix_cover_md.write_text("\n".join(cover_lines).strip() + "\n", encoding="utf-8")

    with open(paths.ftir_validation_appendix_ledger_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "run_no", "label", "window_start_iso", "window_end_iso", "analyte", "row_key",
            "mole_count", "ftir_count", "mole_avg", "ftir_avg", "difference",
            "paired", "status", "excluded", "exclusion_reason", "reviewer", "updated_iso",
        ])
        for row in aligned_rows:
            if not isinstance(row, dict):
                continue
            w.writerow([
                row.get("run_no"), row.get("label"), row.get("window_start_iso"), row.get("window_end_iso"),
                row.get("analyte"), row.get("row_key"), row.get("mole_count"), row.get("ftir_count"),
                _fmt_num(row.get("mole_avg"), 6, ""), _fmt_num(row.get("ftir_avg"), 6, ""),
                _fmt_num(row.get("difference"), 6, ""), row.get("paired"), row.get("status"),
                row.get("excluded"), row.get("exclusion_reason"), row.get("reviewer"), row.get("updated_iso"),
            ])

    with open(paths.ftir_validation_appendix_method301_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "analyte", "mode", "paired_window_count", "excluded_window_count", "mole_mean", "ftir_mean",
            "mean_difference", "relative_bias_pct", "correction_factor", "difference_sd",
            "t_statistic", "t_critical_95_two_sided", "candidate_variance", "validated_variance",
            "f_statistic", "f_critical_95", "bias_status", "precision_status", "overall_status", "note",
        ])
        for row in method301_rows:
            if not isinstance(row, dict):
                continue
            w.writerow([
                row.get("analyte"), row.get("mode"), row.get("paired_window_count"), row.get("excluded_window_count"),
                _fmt_num(row.get("mole_mean"), 6, ""), _fmt_num(row.get("ftir_mean"), 6, ""),
                _fmt_num(row.get("mean_difference"), 6, ""), _fmt_num(row.get("relative_bias_pct"), 6, ""),
                _fmt_num(row.get("correction_factor"), 6, ""), _fmt_num(row.get("difference_sd"), 6, ""),
                _fmt_num(row.get("t_statistic"), 6, ""), _fmt_num(row.get("t_critical_95_two_sided"), 6, ""),
                _fmt_num(row.get("candidate_variance"), 6, ""), _fmt_num(row.get("validated_variance"), 6, ""),
                _fmt_num(row.get("f_statistic"), 6, ""), _fmt_num(row.get("f_critical_95"), 6, ""),
                row.get("bias_status"), row.get("precision_status"), row.get("overall_status"), row.get("note"),
            ])

    with open(paths.ftir_validation_appendix_exclusions_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["row_key", "run_no", "label", "analyte", "reason", "reviewer", "updated_iso"])
        for row in excluded_rows:
            if not isinstance(row, dict):
                continue
            w.writerow([
                row.get("row_key"), row.get("run_no"), row.get("label"), row.get("analyte"),
                row.get("reason"), row.get("reviewer"), row.get("updated_iso"),
            ])

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        def _xlsx_value(value: Any) -> Any:
            if value is None:
                return ""
            if isinstance(value, bool):
                return "TRUE" if value else "FALSE"
            if isinstance(value, (list, tuple, set, dict)):
                try:
                    return json.dumps(value, ensure_ascii=True)
                except Exception:
                    return str(value)
            return value

        def _add_sheet(wb: Any, title: str, headers: List[str], rows: Iterable[Iterable[Any]]) -> None:
            ws = wb.create_sheet(title=title[:31])
            ws.append(headers)
            header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            widths = [max(len(str(h or "")), 10) for h in headers]
            for row in rows:
                cleaned = [_xlsx_value(v) for v in row]
                ws.append(cleaned)
                for idx, value in enumerate(cleaned):
                    widths[idx] = min(max(widths[idx], len(str(value or "")) + 2), 60)
            for idx, width in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(idx)].width = width

        wb = Workbook()
        summary_ws = wb.active
        summary_ws.title = "Signoff Summary"
        summary_headers = ["Field", "Value"]
        summary_ws.append(summary_headers)
        header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        for cell in summary_ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        summary_rows = [
            ("Validation status", ftir.get("status")),
            ("Overall status", ftir.get("overall_status")),
            ("Source", ftir.get("source")),
            ("Validation mode", config.get("validation_mode")),
            ("Comparator method", config.get("comparator_method")),
            ("Vendor profile requested", config.get("ftir_vendor_profile")),
            ("Vendor profile used", ((ftir.get("ftir_source") or {}).get("vendor_profile_used") if isinstance(ftir.get("ftir_source"), dict) else "")),
            ("Acceptance basis", ftir.get("acceptance_basis")),
            ("Acceptance recommendation", ftir.get("acceptance_recommended_decision")),
            ("Acceptance basis note", ftir.get("acceptance_basis_note")),
            ("Comparison set count", ftir.get("comparison_set_count")),
            ("Included comparison set count", ftir.get("included_comparison_set_count")),
            ("Excluded comparison set count", ftir.get("excluded_comparison_set_count")),
            ("Paired window count", ftir.get("paired_window_count")),
            ("Excluded row count", ftir.get("excluded_count")),
            ("Review locked", ftir.get("review_locked")),
            ("Review lock by", ftir.get("review_lock_by")),
            ("Review lock at", ftir.get("review_lock_iso")),
            ("Signoff decision", signoff.get("decision") or "UNSIGNED"),
            ("Signoff basis", signoff.get("basis")),
            ("Signoff by", signoff.get("by")),
            ("Signoff role", signoff.get("role")),
            ("Signoff at", signoff.get("iso")),
            ("Signoff note", signoff.get("note")),
            ("Frozen snapshot JSON", snap.get("json_path")),
            ("Frozen snapshot by", snap.get("snapshot_by")),
            ("Frozen snapshot at", snap.get("snapshot_iso")),
            ("Coverage note", ftir.get("coverage_note")),
        ]
        for key, value in summary_rows:
            summary_ws.append([key, _xlsx_value(value)])
        summary_ws.freeze_panes = "A2"
        summary_ws.auto_filter.ref = summary_ws.dimensions
        summary_ws.column_dimensions["A"].width = 30
        summary_ws.column_dimensions["B"].width = 80

        _add_sheet(
            wb,
            "Comparison Sets",
            [
                "set_no", "set_key", "run_no", "label", "window_start_iso", "window_end_iso",
                "validation_mode", "review_state", "inclusion_status", "formal_basis",
                "row_count", "paired_row_count", "included_row_count", "included_paired_row_count",
                "excluded_row_count", "error_row_count", "warning_row_count",
                "analytes", "paired_analytes", "excluded_analytes", "note",
            ],
            [
                [
                    row.get("set_no"), row.get("set_key"), row.get("run_no"), row.get("label"),
                    row.get("window_start_iso"), row.get("window_end_iso"), row.get("validation_mode"),
                    row.get("review_state"), row.get("inclusion_status"), row.get("formal_basis"),
                    row.get("row_count"), row.get("paired_row_count"), row.get("included_row_count"),
                    row.get("included_paired_row_count"), row.get("excluded_row_count"),
                    row.get("error_row_count"), row.get("warning_row_count"),
                    row.get("analytes"), row.get("paired_analytes"), row.get("excluded_analytes"), row.get("note"),
                ]
                for row in comparison_sets
                if isinstance(row, dict)
            ],
        )

        _add_sheet(
            wb,
            "Aligned Rows",
            [
                "comparison_set_no", "comparison_set_key", "comparison_set_status", "comparison_set_basis",
                "run_no", "label", "window_start_iso", "window_end_iso", "analyte", "row_key",
                "mole_count", "ftir_count", "mole_avg", "ftir_avg", "difference", "paired", "status",
                "excluded", "exclusion_reason", "reviewer", "updated_iso", "qa_status", "qa_flags",
                "offset_seconds_adjusted", "drift_seconds_adjusted", "mole_coverage_ratio", "ftir_coverage_ratio",
            ],
            [
                [
                    row.get("comparison_set_no"), row.get("comparison_set_key"), row.get("comparison_set_status"),
                    row.get("comparison_set_basis"), row.get("run_no"), row.get("label"),
                    row.get("window_start_iso"), row.get("window_end_iso"), row.get("analyte"), row.get("row_key"),
                    row.get("mole_count"), row.get("ftir_count"), row.get("mole_avg"), row.get("ftir_avg"),
                    row.get("difference"), row.get("paired"), row.get("status"), row.get("excluded"),
                    row.get("exclusion_reason"), row.get("reviewer"), row.get("updated_iso"),
                    row.get("qa_status"), row.get("qa_flags"), row.get("offset_seconds_adjusted"),
                    row.get("drift_seconds_adjusted"), row.get("mole_coverage_ratio"), row.get("ftir_coverage_ratio"),
                ]
                for row in aligned_rows
                if isinstance(row, dict)
            ],
        )

        _add_sheet(
            wb,
            "Method301 Stats",
            [
                "analyte", "mode", "comparison_set_count", "included_comparison_set_count",
                "excluded_comparison_set_count", "paired_window_count", "excluded_window_count",
                "mole_mean", "ftir_mean", "mean_difference", "relative_bias_pct", "correction_factor",
                "difference_sd", "t_statistic", "t_critical_95_two_sided", "candidate_variance",
                "validated_variance", "f_statistic", "f_critical_95", "bias_status",
                "precision_status", "overall_status", "note",
            ],
            [
                [
                    row.get("analyte"), row.get("mode"), row.get("comparison_set_count"),
                    row.get("included_comparison_set_count"), row.get("excluded_comparison_set_count"),
                    row.get("paired_window_count"), row.get("excluded_window_count"),
                    row.get("mole_mean"), row.get("ftir_mean"), row.get("mean_difference"),
                    row.get("relative_bias_pct"), row.get("correction_factor"), row.get("difference_sd"),
                    row.get("t_statistic"), row.get("t_critical_95_two_sided"), row.get("candidate_variance"),
                    row.get("validated_variance"), row.get("f_statistic"), row.get("f_critical_95"),
                    row.get("bias_status"), row.get("precision_status"), row.get("overall_status"), row.get("note"),
                ]
                for row in method301_rows
                if isinstance(row, dict)
            ],
        )

        _add_sheet(
            wb,
            "Exclusions",
            ["row_key", "run_no", "label", "analyte", "reason", "reviewer", "updated_iso"],
            [
                [
                    row.get("row_key"), row.get("run_no"), row.get("label"),
                    row.get("analyte"), row.get("reason"), row.get("reviewer"), row.get("updated_iso"),
                ]
                for row in excluded_rows
                if isinstance(row, dict)
            ],
        )

        wb.save(str(paths.ftir_validation_appendix_workbook_xlsx))
    except Exception:
        _remove_if_exists(paths.ftir_validation_appendix_workbook_xlsx)

    appendix_files: List[Dict[str, Any]] = []
    for p in [
        paths.ftir_validation_appendix_cover_md,
        paths.ftir_validation_appendix_ledger_csv,
        paths.ftir_validation_appendix_method301_csv,
        paths.ftir_validation_appendix_exclusions_csv,
        paths.ftir_validation_appendix_workbook_xlsx,
    ]:
        try:
            appendix_files.append({
                "name": p.name,
                "path": str(p),
                "sha256": _sha256_file(p),
                "bytes": p.stat().st_size,
            })
        except Exception:
            continue
    appendix_index = {
        "contract_version": "ftir_validation_appendix_v1",
        "export_dir": str(paths.ftir_validation_appendix_dir),
        "source": str(ftir.get("source") or ""),
        "review_locked": bool(ftir.get("review_locked")),
        "signoff": signoff,
        "files": appendix_files,
    }
    paths.ftir_validation_appendix_index_json.write_text(json.dumps(appendix_index, indent=2), encoding="utf-8")
    return {
        "dir": str(paths.ftir_validation_appendix_dir),
        "cover_md": str(paths.ftir_validation_appendix_cover_md),
        "ledger_csv": str(paths.ftir_validation_appendix_ledger_csv),
        "method301_csv": str(paths.ftir_validation_appendix_method301_csv),
        "exclusions_csv": str(paths.ftir_validation_appendix_exclusions_csv),
        "workbook_xlsx": str(paths.ftir_validation_appendix_workbook_xlsx),
        "index_json": str(paths.ftir_validation_appendix_index_json),
    }


# -----------------------------
# Main generator
# -----------------------------


@dataclass
class ReportPackPaths:
    out_dir: Path
    final_report_dir: Path
    ftir_validation_appendix_dir: Path
    summary_json: Path
    report_context_json: Path
    ftir_validation_json: Path
    ftir_validation_windows_csv: Path
    ftir_validation_method301_csv: Path
    ftir_validation_appendix_cover_md: Path
    ftir_validation_appendix_ledger_csv: Path
    ftir_validation_appendix_method301_csv: Path
    ftir_validation_appendix_exclusions_csv: Path
    ftir_validation_appendix_workbook_xlsx: Path
    ftir_validation_appendix_index_json: Path
    final_report_md: Path
    final_report_docx: Path
    final_report_pdf: Path
    final_report_index_json: Path
    evidence_bundle_json: Path
    evidence_step_eval_csv: Path
    analyzer_validity_csv: Path
    fuel_analysis_csv: Path
    pollutant_adjustments_csv: Path
    pollutant_adjustment_history_csv: Path
    spike_recovery_csv: Path
    side_by_side_csv: Path
    test_matrix_csv: Path
    qaqc_events_csv: Path
    regulatory_snapshot_csv: Path
    reference_trace_csv: Path
    reference_acceptance_csv: Path
    reference_sources_csv: Path
    reference_latest_json: Path
    reference_offsets_csv: Path
    reference_offsets_json: Path
    pdf_path: Path
    index_json: Path


def generate_report_pack_v1(
    session: Dict[str, Any],
    cfg_path: Optional[Path] = None,
    session_dir: Optional[Path] = None,
    master_db_path: Optional[Path] = None,
) -> Path:
    """Generate the report pack. Returns the output directory."""

    if cfg_path is None:
        cfg_path = Path.cwd() / "runner_config.json"

    if session_dir is None:
        # Prefer explicit session_dir path in session JSON.
        pths = session.get("paths") if isinstance(session.get("paths"), dict) else {}
        cand = pths.get("session_dir") or pths.get("daq_run_dir")
        if cand:
            try:
                session_dir = Path(str(cand)).expanduser()
                if not session_dir.is_absolute():
                    session_dir = (cfg_path.parent / session_dir).resolve()
                else:
                    session_dir = session_dir.resolve()
            except Exception:
                session_dir = None

    if session_dir is None:
        # If cfg is inside a session folder, use parent; else fallback to cfg folder.
        if cfg_path.name.lower() == "runner_config.json":
            session_dir = cfg_path.parent
        else:
            session_dir = cfg_path.parent

    session_dir = Path(session_dir).expanduser().resolve()

    # Output folder
    out_dir = _ensure_dir(session_dir / "exports" / "report_pack_v1")
    final_report_dir = _ensure_dir(session_dir / "exports" / "final_report_v1")
    ftir_validation_appendix_dir = _ensure_dir(final_report_dir / "ftir_validation_appendix_v1")
    legacy_final_report_md = out_dir / "final_test_report_v1.md"
    try:
        if legacy_final_report_md.exists():
            legacy_final_report_md.unlink()
    except Exception:
        pass

    paths = ReportPackPaths(
        out_dir=out_dir,
        final_report_dir=final_report_dir,
        ftir_validation_appendix_dir=ftir_validation_appendix_dir,
        summary_json=out_dir / "summary.json",
        report_context_json=out_dir / "report_context.json",
        ftir_validation_json=out_dir / "ftir_validation_summary.json",
        ftir_validation_windows_csv=out_dir / "ftir_validation_window_alignment.csv",
        ftir_validation_method301_csv=out_dir / "ftir_validation_method301.csv",
        ftir_validation_appendix_cover_md=ftir_validation_appendix_dir / "ftir_validation_signoff_cover_v1.md",
        ftir_validation_appendix_ledger_csv=ftir_validation_appendix_dir / "ftir_validation_signed_comparison_ledger.csv",
        ftir_validation_appendix_method301_csv=ftir_validation_appendix_dir / "ftir_validation_signed_method301_stats.csv",
        ftir_validation_appendix_exclusions_csv=ftir_validation_appendix_dir / "ftir_validation_signed_exclusion_register.csv",
        ftir_validation_appendix_workbook_xlsx=ftir_validation_appendix_dir / "ftir_validation_reviewer_workbook_v1.xlsx",
        ftir_validation_appendix_index_json=ftir_validation_appendix_dir / "index.json",
        final_report_md=final_report_dir / "final_test_report_v1.md",
        final_report_docx=final_report_dir / "final_test_report_v1.docx",
        final_report_pdf=final_report_dir / "final_test_report_v1.pdf",
        final_report_index_json=final_report_dir / "index.json",
        evidence_bundle_json=out_dir / "evidence_bundle.json",
        evidence_step_eval_csv=out_dir / "evidence_step_eval.csv",
        analyzer_validity_csv=out_dir / "analyzer_validity.csv",
        fuel_analysis_csv=out_dir / "fuel_analysis_snapshot.csv",
        pollutant_adjustments_csv=out_dir / "pollutant_adjustments_snapshot.csv",
        pollutant_adjustment_history_csv=out_dir / "pollutant_adjustment_history.csv",
        spike_recovery_csv=out_dir / "spike_recovery_snapshot.csv",
        side_by_side_csv=out_dir / "side_by_side_snapshot.csv",
        test_matrix_csv=out_dir / "test_matrix_scorecard.csv",
        qaqc_events_csv=out_dir / "qaqc_events.csv",
        regulatory_snapshot_csv=out_dir / "regulatory_snapshot.csv",
        reference_trace_csv=out_dir / "reference_audit_trace.csv",
        reference_acceptance_csv=out_dir / "reference_audit_acceptance.csv",
        reference_sources_csv=out_dir / "reference_audit_sources.csv",
        reference_latest_json=out_dir / "reference_audit_latest.json",
        reference_offsets_csv=out_dir / "reference_audit_offset_recommendations.csv",
        reference_offsets_json=out_dir / "reference_audit_offset_recommendations.json",
        pdf_path=out_dir / "report_pack_v1.pdf",
        index_json=out_dir / "index.json",
    )

    # Locate evidence folders (best effort)
    meta_dir = session_dir / "meta"
    raw_dir = session_dir / "raw"

    build_meta = _read_json(meta_dir / "build.json") or {}
    provenance = _read_json(meta_dir / "provenance.json") or {}

    # Hash of runner_config for tamper-evidence
    cfg_hash = None
    try:
        if cfg_path.exists():
            cfg_hash = _sha256_file(cfg_path)
    except Exception:
        cfg_hash = None

    # DB path
    if master_db_path is None:
        pths = session.get("paths") if isinstance(session.get("paths"), dict) else {}
        cand_db = pths.get("db_path")
        if cand_db:
            try:
                master_db_path = Path(str(cand_db)).expanduser()
                if not master_db_path.is_absolute():
                    master_db_path = (cfg_path.parent / master_db_path).resolve()
                else:
                    master_db_path = master_db_path.resolve()
            except Exception:
                master_db_path = None

    if master_db_path is None:
        # Try local mole_config.json next to cfg_path, then project root
        cand = cfg_path.parent / "mole_config.json"
        if cand.exists():
            cfg = _read_json(cand) or {}
            pths = cfg.get("paths") if isinstance(cfg.get("paths"), dict) else {}
            dbp = pths.get("db_path")
            if dbp:
                try:
                    master_db_path = Path(str(dbp)).expanduser()
                    if not master_db_path.is_absolute():
                        master_db_path = (cand.parent / master_db_path).resolve()
                    else:
                        master_db_path = master_db_path.resolve()
                except Exception:
                    master_db_path = None

    if master_db_path is None:
        # final fallback: ../mole_das_data/db/mole_master.sqlite relative to script
        here = Path(__file__).resolve().parent
        master_db_path = (here.parent / "mole_das_data" / "db" / "mole_master.sqlite").resolve()

    run_id = str(session.get("run_id") or "").strip() or None
    job_id = None
    try:
        proj = session.get("project") if isinstance(session.get("project"), dict) else {}
        job_id = str(proj.get("job_id") or "").strip() or None
    except Exception:
        job_id = None

    # QA/QC events
    qaqc_events = _query_qaqc_events(master_db_path, run_id=run_id)
    qaqc_summary = _summarize_qaqc(qaqc_events)

    # Test matrix scoring
    tm_score = _score_test_matrix(session)

    # Raw samples + health stats
    raw_samples_path = raw_dir / "raw_samples.jsonl"
    health_path = raw_dir / "health_states.jsonl"
    reference_path = raw_dir / "reference_audit.jsonl"

    raw_count = 0
    first_ts = None
    last_ts = None
    if raw_samples_path.exists():
        try:
            with open(raw_samples_path, "r", encoding="utf-8") as f:
                for line in f:
                    if not line.strip():
                        continue
                    raw_count += 1
                    if raw_count == 1:
                        try:
                            obj = json.loads(line)
                            first_ts = obj.get("ts_iso") or obj.get("timestamp_iso")
                        except Exception:
                            first_ts = None
                    last_obj_ts = None
                    try:
                        obj = json.loads(line)
                        last_obj_ts = obj.get("ts_iso") or obj.get("timestamp_iso")
                    except Exception:
                        last_obj_ts = None
                    if last_obj_ts:
                        last_ts = last_obj_ts
        except Exception:
            pass

    health_counts: Dict[str, int] = {}
    if health_path.exists():
        try:
            with open(health_path, "r", encoding="utf-8") as f:
                for line in f:
                    if not line.strip():
                        continue
                    try:
                        obj = json.loads(line)
                    except Exception:
                        continue
                    h = str(obj.get("health") or obj.get("state") or "").strip().upper() or "(UNKNOWN)"
                    health_counts[h] = health_counts.get(h, 0) + 1
        except Exception:
            pass

    reference_summary, reference_trace_rows, reference_source_rows, reference_latest = _summarize_reference_audit(
        session=session,
        reference_path=reference_path,
    )
    evidence_bundle, evidence_step_eval_rows = _build_evidence_bundle(
        session,
        cfg_path=cfg_path,
        session_dir=session_dir,
        raw_samples_path=raw_samples_path,
        raw_events_path=raw_dir / "raw_events.jsonl",
        health_path=health_path,
        reference_path=reference_path,
        raw_count=raw_count,
        first_ts=first_ts,
        last_ts=last_ts,
        health_counts=health_counts,
    )
    analyzer_validity_summary = evidence_bundle.get("analyzer_validity") if isinstance(evidence_bundle.get("analyzer_validity"), dict) else {"status": "NO_DATA", "counts": {}, "per_pollutant": {}}
    regulatory_summary = _build_regulatory_snapshot(session, raw_samples_path, analyzer_validity=analyzer_validity_summary)
    fuel_analysis_summary = _fuel_analysis_snapshot(session)
    pollutant_adjustments_summary = _pollutant_adjustment_summary(session)
    spike_recovery_summary = _spike_recovery_snapshot(session)
    side_by_side_summary = _side_by_side_snapshot(session)
    run_aggregation_preview = _report_run_aggregation(session, evidence_bundle, cfg_path=cfg_path, session_dir=session_dir)
    ftir_validation_cfg = _ftir_validation_cfg_from_session(session)
    ftir_validation_summary = _locked_ftir_validation_snapshot(ftir_validation_cfg)
    ftir_validation_source = "LOCKED_REVIEW_SNAPSHOT" if isinstance(ftir_validation_summary, dict) else "LIVE_COMPUTE"
    if not isinstance(ftir_validation_summary, dict) and mole_ftir_validation is not None:
        ftir_validation_summary = mole_ftir_validation.build_validation_package(
            ftir_validation_cfg,
            run_aggregation=run_aggregation_preview,
            raw_samples_path=raw_samples_path,
        )
    elif not isinstance(ftir_validation_summary, dict):
        ftir_validation_summary = {
            "status": "Gap",
            "config": ftir_validation_cfg,
            "ftir_source": {"status": "Gap", "note": "FTIR validation module not available."},
            "mole_source": {"status": "Gap", "note": "FTIR validation module not available."},
            "windows": {"status": "Gap", "rows": [], "note": "FTIR validation module not available."},
            "aligned_rows": [],
            "method301": [],
            "coverage_note": "FTIR validation module not available.",
            "overall_status": "Gap",
            "paired_window_count": 0,
        }
        ftir_validation_source = "MODULE_UNAVAILABLE"
    if isinstance(ftir_validation_summary, dict):
        ftir_validation_summary["source"] = ftir_validation_source
        if not isinstance(ftir_validation_summary.get("review_snapshot"), dict):
            ftir_validation_summary["review_snapshot"] = (
                dict(ftir_validation_cfg.get("review_snapshot") or {})
                if isinstance(ftir_validation_cfg.get("review_snapshot"), dict)
                else {}
            )

    offset_policy = {}
    offset_store_summary = {}
    offset_recommendation_set: Dict[str, Any] = {}
    offset_promotion_summary: Dict[str, Any] = {
        "attempted": False,
        "applied": False,
        "store_path": "",
        "active_set_id": "",
        "reason": "offset_module_unavailable",
    }
    if mole_ftir_offset_recommendations is not None:
        try:
            ref_cfg = _reference_cfg_from_session(session)
            store = mole_ftir_offset_recommendations.load_store()
            offset_store_summary = mole_ftir_offset_recommendations.summarize_store(store)
            offset_policy = mole_ftir_offset_recommendations.normalize_project_policy(
                ref_cfg.get("offset_recommendations"),
                default_use_global=bool(store.get("auto_apply_future_projects")),
            )
            offset_recommendation_set = mole_ftir_offset_recommendations.derive_recommendation_set(
                session=session,
                reference_summary=reference_summary,
                trace_rows=reference_trace_rows,
                report_pack_dir=out_dir,
                session_dir=session_dir,
                cfg_path=cfg_path,
            )
            if isinstance(offset_recommendation_set, dict):
                report_block = offset_recommendation_set.get("report") if isinstance(offset_recommendation_set.get("report"), dict) else {}
                report_block["report_pack_dir"] = str(out_dir)
                report_block["summary_json"] = str(paths.summary_json)
                report_block["offsets_json"] = str(paths.reference_offsets_json)
                report_block["offsets_csv"] = str(paths.reference_offsets_csv)
                offset_recommendation_set["report"] = report_block

            if bool(offset_policy.get("promote_latest_for_future_projects")):
                offset_promotion_summary["attempted"] = True
                if bool((offset_recommendation_set or {}).get("eligible_for_promotion")):
                    promoted = mole_ftir_offset_recommendations.promote_recommendation_set(
                        offset_recommendation_set,
                        auto_apply_future_projects=True,
                    )
                    active = mole_ftir_offset_recommendations.get_active_set(promoted) or {}
                    offset_store_summary = mole_ftir_offset_recommendations.summarize_store(promoted)
                    offset_promotion_summary = {
                        "attempted": True,
                        "applied": True,
                        "store_path": str(mole_ftir_offset_recommendations.default_store_path()),
                        "active_set_id": str(active.get("set_id") or ""),
                        "reason": "promoted",
                    }
                else:
                    offset_promotion_summary = {
                        "attempted": True,
                        "applied": False,
                        "store_path": str(mole_ftir_offset_recommendations.default_store_path()),
                        "active_set_id": "",
                        "reason": str(((offset_recommendation_set or {}).get("eligibility") or {}).get("reason") or "not_eligible"),
                    }
            else:
                offset_promotion_summary = {
                    "attempted": False,
                    "applied": False,
                    "store_path": str(mole_ftir_offset_recommendations.default_store_path()),
                    "active_set_id": str(offset_store_summary.get("active_set_id") or ""),
                    "reason": "promotion_not_requested",
                }
        except Exception as exc:
            offset_promotion_summary = {
                "attempted": False,
                "applied": False,
                "store_path": str(mole_ftir_offset_recommendations.default_store_path()),
                "active_set_id": "",
                "reason": f"error:{type(exc).__name__}",
            }

    # Overall pass/fail (v1):
    #  - If test matrix required tests fail => FAIL
    #  - If qaqc has explicit failures => FAIL
    overall_pass = None
    if tm_score.get("overall_pass") is None and qaqc_summary.get("overall_ok") is None:
        overall_pass = None
    else:
        overall_pass = True
        if tm_score.get("overall_pass") is False:
            overall_pass = False
        if qaqc_summary.get("overall_ok") is False:
            overall_pass = False
    try:
        validity_blocked = int(((analyzer_validity_summary.get("counts") or {}).get("regulatory_blocked")) or 0)
    except Exception:
        validity_blocked = 0
    if validity_blocked > 0:
        overall_pass = False

    ftir_validation_appendix = _write_ftir_validation_appendix(paths, ftir_validation_summary)

    # Summary JSON
    summary = {
        "generated_iso": _now_iso(),
        "report_pack": {
            "version": "v1",
        },
        "session": {
            "job_id": job_id,
            "run_id": run_id,
            "session_dir": str(session_dir),
            "config_path": str(cfg_path) if cfg_path else None,
            "config_sha256": cfg_hash,
            "overall_pass": overall_pass,
        },
        "evidence": {
            "meta_dir": str(meta_dir),
            "raw_dir": str(raw_dir),
            "worksteps": {
                "path": str((evidence_bundle.get("sources") or {}).get("worksteps", {}).get("path") or ""),
                "count": (evidence_bundle.get("sources") or {}).get("worksteps", {}).get("count"),
                "total_count": (evidence_bundle.get("sources") or {}).get("worksteps", {}).get("total_count"),
                "filtered_out_count": (evidence_bundle.get("sources") or {}).get("worksteps", {}).get("filtered_out_count"),
                "path_scope": (evidence_bundle.get("sources") or {}).get("worksteps", {}).get("path_scope"),
                "session_scope": (evidence_bundle.get("sources") or {}).get("worksteps", {}).get("session_scope"),
            },
            "raw_samples": {
                "path": str(raw_samples_path),
                "count": raw_count,
                "first_ts": first_ts,
                "last_ts": last_ts,
            },
            "health_states": {
                "path": str(health_path),
                "counts": health_counts,
            },
            "reference_audit": {
                "path": str(reference_path),
                "captured_record_count": (reference_summary.get("evidence") or {}).get("captured_record_count"),
                "mode": (reference_summary.get("evidence") or {}).get("mode"),
            },
        },
        "provenance": {
            "build": build_meta,
            "provenance": provenance,
        },
        "test_matrix": tm_score,
        "qaqc": {
            "db_path": str(master_db_path),
            "summary": qaqc_summary,
            "spike_recovery": spike_recovery_summary,
            "side_by_side": side_by_side_summary,
        },
        "ftir_validation": {
            "status": ftir_validation_summary.get("status"),
            "overall_status": ftir_validation_summary.get("overall_status"),
            "acceptance_basis": ftir_validation_summary.get("acceptance_basis"),
            "acceptance_recommended_decision": ftir_validation_summary.get("acceptance_recommended_decision"),
            "acceptance_basis_note": ftir_validation_summary.get("acceptance_basis_note"),
            "source": ftir_validation_source,
            "coverage_note": ftir_validation_summary.get("coverage_note"),
            "paired_window_count": ftir_validation_summary.get("paired_window_count"),
            "excluded_count": ftir_validation_summary.get("excluded_count"),
            "review_notes": ftir_validation_summary.get("review_notes"),
            "reviewer": ftir_validation_summary.get("reviewer"),
            "review_locked": ftir_validation_summary.get("review_locked"),
            "review_lock_by": ftir_validation_summary.get("review_lock_by"),
            "review_lock_iso": ftir_validation_summary.get("review_lock_iso"),
            "review_unlock_by": ftir_validation_summary.get("review_unlock_by"),
            "review_unlock_iso": ftir_validation_summary.get("review_unlock_iso"),
            "review_snapshot": ftir_validation_summary.get("review_snapshot"),
            "signoff": ftir_validation_summary.get("signoff"),
            "config": ftir_validation_summary.get("config"),
            "ftir_source": ftir_validation_summary.get("ftir_source"),
            "mole_source": ftir_validation_summary.get("mole_source"),
            "windows": {
                "status": ((ftir_validation_summary.get("windows") or {}).get("status") if isinstance(ftir_validation_summary.get("windows"), dict) else None),
                "source": ((ftir_validation_summary.get("windows") or {}).get("source") if isinstance(ftir_validation_summary.get("windows"), dict) else None),
                "count": len(((ftir_validation_summary.get("windows") or {}).get("rows") if isinstance(ftir_validation_summary.get("windows"), dict) else []) or []),
                "note": ((ftir_validation_summary.get("windows") or {}).get("note") if isinstance(ftir_validation_summary.get("windows"), dict) else None),
            },
            "excluded_rows": ftir_validation_summary.get("excluded_rows"),
            "summary_json": str(paths.ftir_validation_json),
            "window_alignment_csv": str(paths.ftir_validation_windows_csv),
            "method301_csv": str(paths.ftir_validation_method301_csv),
            "appendix": ftir_validation_appendix,
        },
        "analyzer_validity": analyzer_validity_summary,
        "fuel_analysis": fuel_analysis_summary,
        "pollutant_adjustments": pollutant_adjustments_summary,
        "regulatory_snapshot": regulatory_summary,
        "evidence_bundle": {
            "bundle_json": str(paths.evidence_bundle_json),
            "step_eval_csv": str(paths.evidence_step_eval_csv),
            "analyzer_validity_csv": str(paths.analyzer_validity_csv),
            "fuel_analysis_csv": str(paths.fuel_analysis_csv),
            "pollutant_adjustments_csv": str(paths.pollutant_adjustments_csv),
            "pollutant_adjustment_history_csv": str(paths.pollutant_adjustment_history_csv),
            "spike_recovery_csv": str(paths.spike_recovery_csv),
            "side_by_side_csv": str(paths.side_by_side_csv),
            "ftir_validation_json": str(paths.ftir_validation_json),
            "ftir_validation_windows_csv": str(paths.ftir_validation_windows_csv),
            "ftir_validation_method301_csv": str(paths.ftir_validation_method301_csv),
            "ftir_validation_appendix_cover_md": str(paths.ftir_validation_appendix_cover_md),
            "ftir_validation_appendix_ledger_csv": str(paths.ftir_validation_appendix_ledger_csv),
            "ftir_validation_appendix_method301_csv": str(paths.ftir_validation_appendix_method301_csv),
            "ftir_validation_appendix_exclusions_csv": str(paths.ftir_validation_appendix_exclusions_csv),
            "ftir_validation_appendix_workbook_xlsx": str(paths.ftir_validation_appendix_workbook_xlsx),
            "ftir_validation_appendix_index_json": str(paths.ftir_validation_appendix_index_json),
            "spec_engine_shadow": evidence_bundle.get("spec_engine_shadow"),
            "static_artifacts": {
                "active_count": ((evidence_bundle.get("static_artifacts") or {}).get("active_count")),
                "add_events": ((evidence_bundle.get("static_artifacts") or {}).get("add_events")),
                "remove_events": ((evidence_bundle.get("static_artifacts") or {}).get("remove_events")),
            },
        },
        "reference_audit": reference_summary,
    }

    if mole_ftir_offset_recommendations is not None:
        summary["reference_audit"]["offset_recommendations"] = {
            "project_policy": offset_policy,
            "store": offset_store_summary,
            "latest_evaluation": offset_recommendation_set,
            "promotion": offset_promotion_summary,
        }

    paths.summary_json.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    paths.evidence_bundle_json.write_text(json.dumps(evidence_bundle, indent=2), encoding="utf-8")

    with open(paths.evidence_step_eval_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "ts_iso",
            "event",
            "step",
            "channel",
            "formula_version",
            "status",
            "pass",
            "stable",
            "within_tolerance",
            "avg",
            "std",
            "n",
            "target",
            "tolerance",
            "basis",
            "comparison_value",
            "recovery_pct",
            "legacy_pass",
            "shadow_pass",
            "match",
            "reason",
        ])
        for row in evidence_step_eval_rows:
            w.writerow([
                row.get("ts_iso"),
                row.get("event"),
                row.get("step"),
                row.get("channel"),
                row.get("formula_version"),
                row.get("status"),
                row.get("pass"),
                row.get("stable"),
                row.get("within_tolerance"),
                _fmt_num(row.get("avg"), 6, ""),
                _fmt_num(row.get("std"), 6, ""),
                row.get("n"),
                _fmt_num(row.get("target"), 6, ""),
                _fmt_num(row.get("tolerance"), 6, ""),
                row.get("basis"),
                _fmt_num(row.get("comparison_value"), 6, ""),
                _fmt_num(row.get("recovery_pct"), 4, ""),
                row.get("legacy_pass"),
                row.get("shadow_pass"),
                row.get("match"),
                row.get("reason"),
            ])

    with open(paths.analyzer_validity_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "pollutant",
            "state",
            "reason",
            "allow_display",
            "allow_regulatory",
            "allow_combustion_math",
            "zero_ok",
            "span_ok",
            "drift_ok",
            "linearity_ok",
            "latched_invalid",
            "latest_postcal_run",
            "zero_capture",
            "span_capture",
            "post_zero_capture",
            "post_span_capture",
        ])
        validity_rows = (analyzer_validity_summary.get("per_pollutant") or {}) if isinstance(analyzer_validity_summary.get("per_pollutant"), dict) else {}
        for code, row in sorted(validity_rows.items()):
            gates = row.get("gates") if isinstance(row.get("gates"), dict) else {}
            scores = row.get("scores") if isinstance(row.get("scores"), dict) else {}
            sources = row.get("sources") if isinstance(row.get("sources"), dict) else {}
            w.writerow([
                code,
                row.get("state"),
                row.get("reason"),
                gates.get("allow_display"),
                gates.get("allow_regulatory"),
                gates.get("allow_combustion_math"),
                scores.get("zero"),
                scores.get("span"),
                scores.get("drift"),
                scores.get("linearity"),
                row.get("latched_invalid"),
                row.get("latest_postcal_run"),
                sources.get("zero_capture"),
                sources.get("span_capture"),
                sources.get("post_zero_capture"),
                sources.get("post_span_capture"),
            ])

    with open(paths.fuel_analysis_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "present",
            "capture_enabled",
            "captured",
            "source",
            "captured_by",
            "captured_iso",
            "profile_id",
            "profile_label",
            "family",
            "component_count",
            "methodology_basis_code",
            "methodology_basis_label",
            "default_source_citation",
            "default_source_url",
            "default_source_display",
            "default_source_note",
            "f_factor_basis",
            "molecular_weight",
            "relative_density_air",
            "compressibility_basis",
            "compressibility_factor",
            "hhv_btu_scf",
            "lhv_btu_scf",
            "hhv_btu_lb",
            "lhv_btu_lb",
            "hhv_btu_gal",
            "lhv_btu_gal",
            "f_factor_selected",
            "f_factor_hhv",
            "f_factor_lhv",
            "standard_conditions_json",
            "site_conditions_json",
            "traceability_note",
            "notes",
        ])
        w.writerow([
            fuel_analysis_summary.get("present"),
            fuel_analysis_summary.get("capture_enabled"),
            fuel_analysis_summary.get("captured"),
            fuel_analysis_summary.get("source"),
            fuel_analysis_summary.get("captured_by"),
            fuel_analysis_summary.get("captured_iso"),
            fuel_analysis_summary.get("profile_id"),
            fuel_analysis_summary.get("profile_label"),
            fuel_analysis_summary.get("family"),
            fuel_analysis_summary.get("component_count"),
            fuel_analysis_summary.get("methodology_basis_code"),
            fuel_analysis_summary.get("methodology_basis_label"),
            fuel_analysis_summary.get("default_source_citation"),
            fuel_analysis_summary.get("default_source_url"),
            fuel_analysis_summary.get("default_source_display"),
            fuel_analysis_summary.get("default_source_note"),
            fuel_analysis_summary.get("f_factor_basis"),
            _fmt_num(fuel_analysis_summary.get("molecular_weight"), 6, ""),
            _fmt_num(fuel_analysis_summary.get("relative_density_air"), 6, ""),
            fuel_analysis_summary.get("compressibility_basis"),
            _fmt_num(fuel_analysis_summary.get("compressibility_factor"), 6, ""),
            _fmt_num(fuel_analysis_summary.get("hhv_btu_scf"), 6, ""),
            _fmt_num(fuel_analysis_summary.get("lhv_btu_scf"), 6, ""),
            _fmt_num(fuel_analysis_summary.get("hhv_btu_lb"), 6, ""),
            _fmt_num(fuel_analysis_summary.get("lhv_btu_lb"), 6, ""),
            _fmt_num(fuel_analysis_summary.get("hhv_btu_gal"), 6, ""),
            _fmt_num(fuel_analysis_summary.get("lhv_btu_gal"), 6, ""),
            _fmt_num(fuel_analysis_summary.get("f_factor_selected"), 6, ""),
            _fmt_num(fuel_analysis_summary.get("f_factor_hhv"), 6, ""),
            _fmt_num(fuel_analysis_summary.get("f_factor_lhv"), 6, ""),
            json.dumps(fuel_analysis_summary.get("standard_conditions") or {}, ensure_ascii=False),
            json.dumps(fuel_analysis_summary.get("site_conditions") or {}, ensure_ascii=False),
            fuel_analysis_summary.get("traceability_note"),
            fuel_analysis_summary.get("notes"),
        ])

    with open(paths.pollutant_adjustments_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "adjustments_enabled",
            "drift_basis",
            "formula",
            "elapsed_hours",
            "active_run_no",
            "pollutant",
            "channel_enabled",
            "state",
            "active_now",
            "source",
            "source_label",
            "source_run_no",
            "effective_after_run_no",
            "effective_label",
            "scope",
            "scope_label",
            "expires_after_run_no",
            "expires_after_postcal_run_no",
            "lifecycle_status",
            "status_reason",
            "units",
            "bias",
            "drift_per_hr",
            "total_adjustment",
            "note",
            "updated_by",
            "updated_iso",
        ])
        rows = pollutant_adjustments_summary.get("rows") if isinstance(pollutant_adjustments_summary.get("rows"), list) else []
        if rows:
            for row in rows:
                if not isinstance(row, dict):
                    continue
                w.writerow([
                    pollutant_adjustments_summary.get("enabled"),
                    pollutant_adjustments_summary.get("drift_basis"),
                    pollutant_adjustments_summary.get("formula"),
                    _fmt_num(pollutant_adjustments_summary.get("elapsed_hours"), 6, ""),
                    pollutant_adjustments_summary.get("active_run_no"),
                    row.get("pollutant"),
                    row.get("enabled"),
                    row.get("state"),
                    row.get("active_now"),
                    row.get("source"),
                    row.get("source_label"),
                    row.get("source_run_no"),
                    row.get("effective_after_run_no"),
                    row.get("effective_label"),
                    row.get("scope"),
                    row.get("scope_label"),
                    row.get("expires_after_run_no"),
                    row.get("expires_after_postcal_run_no"),
                    row.get("lifecycle_status"),
                    row.get("status_reason"),
                    row.get("units"),
                    _fmt_num(row.get("bias"), 6, ""),
                    _fmt_num(row.get("drift_per_hr"), 6, ""),
                    _fmt_num(row.get("total_adjustment"), 6, ""),
                    row.get("note"),
                    row.get("updated_by"),
                    row.get("updated_iso"),
                ])
        else:
            w.writerow([
                pollutant_adjustments_summary.get("enabled"),
                pollutant_adjustments_summary.get("drift_basis"),
                pollutant_adjustments_summary.get("formula"),
                _fmt_num(pollutant_adjustments_summary.get("elapsed_hours"), 6, ""),
                pollutant_adjustments_summary.get("active_run_no"),
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
            ])

    with open(paths.pollutant_adjustment_history_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "run_no",
            "pollutant",
            "units",
            "event",
            "pass",
            "policy_mode",
            "method_effective",
            "profile_id",
            "track",
            "bias",
            "drift_per_hr",
            "effective_after_run_no",
            "source_run_no",
            "detail",
            "reason",
            "review_headline",
            "review_overall_pass",
        ])
        history_rows = pollutant_adjustments_summary.get("history") if isinstance(pollutant_adjustments_summary.get("history"), list) else []
        for row in history_rows:
            if not isinstance(row, dict):
                continue
            w.writerow([
                row.get("run_no"),
                row.get("pollutant"),
                row.get("units"),
                row.get("event"),
                row.get("pass"),
                row.get("policy_mode"),
                row.get("method_effective"),
                row.get("profile_id"),
                row.get("track"),
                _fmt_num(row.get("bias"), 6, ""),
                _fmt_num(row.get("drift_per_hr"), 6, ""),
                row.get("effective_after_run_no"),
                row.get("source_run_no"),
                row.get("detail"),
                row.get("reason"),
                row.get("review_headline"),
                row.get("review_overall_pass"),
            ])

    with open(paths.spike_recovery_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "enabled",
            "mandatory",
            "applies_to_mole",
            "applies_to_ftir",
            "requires_flow_restriction",
            "flow_restriction_pct",
            "criterion_summary",
            "flow_restriction_summary",
            "method_standard",
            "method_label",
            "method_basis",
            "active_phase",
            "sample_flow_value",
            "sample_flow_units",
            "sample_flow_source",
            "sample_flow_override",
            "spike_flow_value",
            "target_spike_flow_value",
            "flow_status",
            "flow_reason",
            "channel_count",
            "mole_pass_count",
            "mole_fail_count",
            "ftir_pass_count",
            "ftir_fail_count",
        ])
        w.writerow([
            spike_recovery_summary.get("enabled"),
            spike_recovery_summary.get("mandatory"),
            spike_recovery_summary.get("applies_to_mole"),
            spike_recovery_summary.get("applies_to_ftir"),
            spike_recovery_summary.get("requires_flow_restriction"),
            _fmt_num(spike_recovery_summary.get("flow_restriction_pct"), 4, ""),
            spike_recovery_summary.get("criterion_summary"),
            spike_recovery_summary.get("flow_restriction_summary"),
            spike_recovery_summary.get("method_standard"),
            spike_recovery_summary.get("method_label"),
            spike_recovery_summary.get("method_basis"),
            spike_recovery_summary.get("active_phase"),
            _fmt_num(spike_recovery_summary.get("sample_flow_value"), 6, ""),
            spike_recovery_summary.get("sample_flow_units"),
            spike_recovery_summary.get("sample_flow_source"),
            _fmt_num(spike_recovery_summary.get("sample_flow_override"), 6, ""),
            _fmt_num(spike_recovery_summary.get("spike_flow_value"), 6, ""),
            _fmt_num(spike_recovery_summary.get("target_spike_flow_value"), 6, ""),
            spike_recovery_summary.get("flow_status"),
            spike_recovery_summary.get("flow_reason"),
            spike_recovery_summary.get("channel_count"),
            spike_recovery_summary.get("mole_pass_count"),
            spike_recovery_summary.get("mole_fail_count"),
            spike_recovery_summary.get("ftir_pass_count"),
            spike_recovery_summary.get("ftir_fail_count"),
        ])
        w.writerow([])
        w.writerow([
            "pollutant",
            "units",
            "spike_amount",
            "native_mole",
            "spike_mole",
            "mole_recovery_pct",
            "mole_status",
            "native_ftir",
            "spike_ftir",
            "ftir_recovery_pct",
            "ftir_status",
            "native_timestamp_iso",
            "spike_timestamp_iso",
        ])
        for row in spike_recovery_summary.get("channels") or []:
            w.writerow([
                row.get("pollutant"),
                row.get("units"),
                _fmt_num(row.get("spike_amount"), 6, ""),
                _fmt_num(row.get("native_mole"), 6, ""),
                _fmt_num(row.get("spike_mole"), 6, ""),
                _fmt_num(row.get("mole_recovery_pct"), 4, ""),
                row.get("mole_status"),
                _fmt_num(row.get("native_ftir"), 6, ""),
                _fmt_num(row.get("spike_ftir"), 6, ""),
                _fmt_num(row.get("ftir_recovery_pct"), 4, ""),
                row.get("ftir_status"),
                row.get("native_timestamp_iso"),
                row.get("spike_timestamp_iso"),
            ])

    with open(paths.side_by_side_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        latest_sync = side_by_side_summary.get("latest_sync_marker") if isinstance(side_by_side_summary.get("latest_sync_marker"), dict) else {}
        w.writerow([
            "enabled",
            "bias_valve_interval_min",
            "bias_schedule_minutes",
            "bias_schedule_source",
            "bias_schedule_label",
            "bias_due",
            "bias_due_after_run_no",
            "bias_due_after_iso",
            "ambient_purge_required",
            "project_drift_required",
            "notes",
            "event_count",
            "latest_sync_timestamp_iso",
            "latest_sync_label",
            "latest_sync_note",
        ])
        w.writerow([
            side_by_side_summary.get("enabled"),
            side_by_side_summary.get("bias_valve_interval_min"),
            side_by_side_summary.get("bias_schedule_minutes"),
            side_by_side_summary.get("bias_schedule_source"),
            side_by_side_summary.get("bias_schedule_label"),
            side_by_side_summary.get("bias_due"),
            side_by_side_summary.get("bias_due_after_run_no"),
            side_by_side_summary.get("bias_due_after_iso"),
            side_by_side_summary.get("ambient_purge_required"),
            side_by_side_summary.get("project_drift_required"),
            side_by_side_summary.get("notes"),
            side_by_side_summary.get("event_count"),
            latest_sync.get("timestamp_iso"),
            latest_sync.get("label"),
            latest_sync.get("note"),
        ])
        w.writerow([])
        w.writerow([
            "timestamp_iso",
            "event_type",
            "label",
            "note",
            "run_id",
            "run_no",
            "active_run_index",
            "reference_enabled",
            "reference_role",
            "method_standard",
        ])
        for row in side_by_side_summary.get("events") or []:
            if not isinstance(row, dict):
                continue
            w.writerow([
                row.get("timestamp_iso"),
                row.get("event_type"),
                row.get("label"),
                row.get("note"),
                row.get("run_id"),
                row.get("run_no"),
                row.get("active_run_index"),
                row.get("reference_enabled"),
                row.get("reference_role"),
                row.get("method_standard"),
            ])

    if mole_ftir_validation is not None:
        mole_ftir_validation.write_validation_exports(
            ftir_validation_summary,
            json_path=paths.ftir_validation_json,
            windows_csv_path=paths.ftir_validation_windows_csv,
            method301_csv_path=paths.ftir_validation_method301_csv,
        )
    else:
        paths.ftir_validation_json.write_text(json.dumps(ftir_validation_summary, indent=2), encoding="utf-8")

    # Test matrix CSV
    with open(paths.test_matrix_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "name", "required", "performed", "waived", "waive_reason", "points", "min_points", "status", "details"])
        for t in tm_score.get("tests") or []:
            w.writerow([
                t.get("id"),
                t.get("name"),
                "Y" if t.get("required") else "N",
                "Y" if t.get("performed") else "N",
                "Y" if t.get("waived") else "N",
                t.get("waive_reason") or "",
                t.get("points"),
                t.get("min_points"),
                t.get("status"),
                t.get("details"),
            ])

    # QA/QC events CSV
    with open(paths.qaqc_events_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "event_iso",
            "run_id",
            "phase",
            "pollutant",
            "event_type",
            "meas",
            "target",
            "recovery_pct",
            "drift_abs",
            "drift_recovery_pct",
            "combined_drift_ratio",
            "health",
        ])
        for ev in qaqc_events:
            w.writerow([
                ev.get("event_iso"),
                ev.get("run_id"),
                ev.get("phase"),
                ev.get("pollutant"),
                ev.get("event_type"),
                _fmt_num(ev.get("meas"), 3, ""),
                _fmt_num(ev.get("target"), 3, ""),
                _fmt_num(ev.get("recovery_pct"), 2, ""),
                _fmt_num(ev.get("drift_abs"), 3, ""),
                _fmt_num(ev.get("drift_recovery_pct"), 2, ""),
                _fmt_num(ev.get("combined_drift_ratio"), 3, ""),
                ev.get("health"),
            ])

    with open(paths.regulatory_snapshot_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "pollutant",
            "units",
            "raw",
            "dry",
            "corr",
            "o2_ref",
            "validity_state",
            "regulatory_allowed",
            "limit_value",
            "limit_units",
            "compare_value",
            "compare_source",
            "pct_limit",
            "status",
            "mass_value",
            "mass_units",
            "mass_limit_value",
            "mass_limit_units",
            "mass_pct_limit",
            "mass_status",
        ])
        for row in regulatory_summary.get("rows") or []:
            w.writerow([
                row.get("pollutant"),
                row.get("units"),
                _fmt_num(row.get("raw"), 6, ""),
                _fmt_num(row.get("dry"), 6, ""),
                _fmt_num(row.get("corr"), 6, ""),
                _fmt_num(row.get("o2_ref"), 3, ""),
                row.get("validity_state"),
                row.get("regulatory_allowed"),
                _fmt_num(row.get("limit_value"), 6, ""),
                row.get("limit_units"),
                _fmt_num(row.get("compare_value"), 6, ""),
                row.get("compare_source"),
                _fmt_num(row.get("pct_limit"), 3, ""),
                row.get("status"),
                _fmt_num(row.get("mass_value"), 6, ""),
                row.get("mass_units"),
                _fmt_num(row.get("mass_limit_value"), 6, ""),
                row.get("mass_limit_units"),
                _fmt_num(row.get("mass_pct_limit"), 3, ""),
                row.get("mass_status"),
            ])

    with open(paths.reference_trace_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "record_origin",
            "ts_utc",
            "frame_ts_iso",
            "status",
            "ok",
            "provider",
            "role",
            "pollutant",
            "primary_value",
            "reference_value",
            "delta",
            "units",
            "source_column",
            "mapping_source",
            "source_path",
            "source_sha256",
            "source_mtime_iso",
            "age_s",
            "header_index",
            "row_index",
            "delimiter",
            "acceptance_status",
            "acceptance_ok",
            "acceptance_tolerance_expr",
            "acceptance_tolerance_source",
            "acceptance_tolerance_type",
            "acceptance_tolerance_abs",
            "acceptance_tolerance_units",
            "acceptance_expected_units",
            "acceptance_reference_units",
            "acceptance_units_match",
            "acceptance_span_basis_value",
            "acceptance_span_basis_source",
            "acceptance_delta_abs",
            "acceptance_delta_pct_span",
            "acceptance_selection_strategy",
            "acceptance_tolerance_step_type",
            "acceptance_tolerance_step_label",
            "error",
        ])
        for row in reference_trace_rows:
            w.writerow([
                row.get("record_origin"),
                row.get("ts_utc"),
                row.get("frame_ts_iso"),
                row.get("status"),
                row.get("ok"),
                row.get("provider"),
                row.get("role"),
                row.get("pollutant"),
                _fmt_num(row.get("primary_value"), 3, ""),
                _fmt_num(row.get("reference_value"), 3, ""),
                _fmt_num(row.get("delta"), 3, ""),
                row.get("units"),
                row.get("source_column"),
                row.get("mapping_source"),
                row.get("source_path"),
                row.get("source_sha256"),
                row.get("source_mtime_iso"),
                _fmt_num(row.get("age_s"), 1, ""),
                row.get("header_index"),
                row.get("row_index"),
                row.get("delimiter"),
                row.get("acceptance_status"),
                row.get("acceptance_ok"),
                row.get("acceptance_tolerance_expr"),
                row.get("acceptance_tolerance_source"),
                row.get("acceptance_tolerance_type"),
                _fmt_num(row.get("acceptance_tolerance_abs"), 3, ""),
                row.get("acceptance_tolerance_units"),
                row.get("acceptance_expected_units"),
                row.get("acceptance_reference_units"),
                row.get("acceptance_units_match"),
                _fmt_num(row.get("acceptance_span_basis_value"), 3, ""),
                row.get("acceptance_span_basis_source"),
                _fmt_num(row.get("acceptance_delta_abs"), 3, ""),
                _fmt_num(row.get("acceptance_delta_pct_span"), 3, ""),
                row.get("acceptance_selection_strategy"),
                row.get("acceptance_tolerance_step_type"),
                row.get("acceptance_tolerance_step_label"),
                row.get("error"),
            ])

    with open(paths.reference_acceptance_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "pollutant",
            "units",
            "reference_count",
            "delta_count",
            "pass_count",
            "fail_count",
            "unknown_count",
            "latest_status",
            "latest_ok",
            "latest_primary_value",
            "latest_reference_value",
            "latest_delta",
            "latest_delta_abs",
            "latest_delta_pct_span",
            "tolerance_expr",
            "tolerance_source",
            "tolerance_type",
            "tolerance_abs",
            "tolerance_units",
            "span_basis_value",
            "span_basis_source",
            "selection_strategy",
            "status_counts_json",
        ])
        for pollutant, rec in sorted((reference_summary.get("per_pollutant") or {}).items()):
            acc = rec.get("acceptance") if isinstance(rec.get("acceptance"), dict) else {}
            basis = acc.get("basis") if isinstance(acc.get("basis"), dict) else {}
            w.writerow([
                pollutant,
                rec.get("units"),
                rec.get("reference_count"),
                rec.get("delta_count"),
                acc.get("pass_count"),
                acc.get("fail_count"),
                acc.get("unknown_count"),
                acc.get("latest_status"),
                acc.get("latest_ok"),
                _fmt_num(acc.get("latest_primary_value"), 3, ""),
                _fmt_num(acc.get("latest_reference_value"), 3, ""),
                _fmt_num(acc.get("latest_delta"), 3, ""),
                _fmt_num(acc.get("latest_delta_abs"), 3, ""),
                _fmt_num(acc.get("latest_delta_pct_span"), 3, ""),
                basis.get("tolerance_expr"),
                basis.get("tolerance_source"),
                basis.get("tolerance_type"),
                _fmt_num(basis.get("tolerance_abs"), 3, ""),
                basis.get("tolerance_units"),
                _fmt_num(basis.get("span_basis_value"), 3, ""),
                basis.get("span_basis_source"),
                basis.get("selection_strategy"),
                json.dumps(acc.get("status_counts") or {}, ensure_ascii=False),
            ])

    with open(paths.reference_sources_csv, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "record_origin",
            "source_path",
            "source_sha256",
            "source_bytes",
            "provider",
            "role",
            "first_ts",
            "last_ts",
            "record_count",
            "latest_status",
            "source_mtime_iso",
        ])
        for row in reference_source_rows:
            w.writerow([
                row.get("record_origin"),
                row.get("source_path"),
                row.get("source_sha256"),
                row.get("source_bytes"),
                row.get("provider"),
                row.get("role"),
                row.get("first_ts"),
                row.get("last_ts"),
                row.get("record_count"),
                row.get("latest_status"),
                row.get("source_mtime_iso"),
            ])

    paths.reference_latest_json.write_text(json.dumps(reference_latest or {}, indent=2), encoding="utf-8")
    if mole_ftir_offset_recommendations is not None:
        with open(paths.reference_offsets_csv, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow([
                "set_id",
                "generated_iso",
                "eligible_for_promotion",
                "pollutant",
                "units",
                "recommended_offset",
                "recommended_scale",
                "delta_avg",
                "abs_delta_avg",
                "delta_count",
                "reference_count",
                "latest_delta",
                "latest_status",
                "pass_count",
                "fail_count",
                "unknown_count",
                "basis_json",
            ])
            for pollutant, rec in sorted(((offset_recommendation_set or {}).get("channels") or {}).items()):
                w.writerow([
                    (offset_recommendation_set or {}).get("set_id"),
                    (offset_recommendation_set or {}).get("generated_iso"),
                    (offset_recommendation_set or {}).get("eligible_for_promotion"),
                    pollutant,
                    rec.get("units"),
                    _fmt_num(rec.get("recommended_offset"), 6, ""),
                    _fmt_num(rec.get("recommended_scale"), 6, ""),
                    _fmt_num(rec.get("delta_avg"), 6, ""),
                    _fmt_num(rec.get("abs_delta_avg"), 6, ""),
                    rec.get("delta_count"),
                    rec.get("reference_count"),
                    _fmt_num(rec.get("latest_delta"), 6, ""),
                    rec.get("latest_status"),
                    rec.get("pass_count"),
                    rec.get("fail_count"),
                    rec.get("unknown_count"),
                    json.dumps(rec.get("basis") or {}, ensure_ascii=False),
                ])
        paths.reference_offsets_json.write_text(json.dumps(offset_recommendation_set or {}, indent=2), encoding="utf-8")

    # PDF (optional)
    _write_pdf(paths, summary)

    # Report context (built after pack outputs so appendix manifest sees the final files)
    report_context = _build_report_context(
        session=session,
        summary=summary,
        evidence_bundle=evidence_bundle,
        paths=paths,
        cfg_path=cfg_path,
        session_dir=session_dir,
    )
    summary["report_context"] = {
        "contract_version": report_context.get("contract_version"),
        "json_path": str(paths.report_context_json),
        "coverage": report_context.get("coverage"),
    }
    paths.report_context_json.write_text(json.dumps(report_context, indent=2), encoding="utf-8")
    _write_final_report_markdown(paths, report_context, summary)
    _write_final_report_docx(paths, summary)
    _write_final_report_pdf(paths)

    # Rebuild once so the manifest sees report_context.json and final-report artifacts.
    report_context = _build_report_context(
        session=session,
        summary=summary,
        evidence_bundle=evidence_bundle,
        paths=paths,
        cfg_path=cfg_path,
        session_dir=session_dir,
    )
    summary["report_context"] = {
        "contract_version": report_context.get("contract_version"),
        "json_path": str(paths.report_context_json),
        "coverage": report_context.get("coverage"),
    }
    docx_status = _write_final_report_docx(paths, summary)
    pdf_status = _write_final_report_pdf(paths)
    summary["final_report"] = {
        "contract_version": "final_report_v1",
        "export_dir": str(paths.final_report_dir),
        "markdown_path": str(paths.final_report_md),
        "docx_path": str(paths.final_report_docx),
        "pdf_path": str(paths.final_report_pdf),
        "index_path": str(paths.final_report_index_json),
        "ftir_validation_appendix_dir": str(paths.ftir_validation_appendix_dir),
        "source_report_context_json": str(paths.report_context_json),
        "source_report_pack_dir": str(paths.out_dir),
        "render_status": {
            "docx": docx_status,
            "pdf": pdf_status,
        },
    }
    paths.report_context_json.write_text(json.dumps(report_context, indent=2), encoding="utf-8")
    _write_final_report_markdown(paths, report_context, summary)
    docx_status = _write_final_report_docx(paths, summary)
    pdf_status = _write_final_report_pdf(paths)
    summary["final_report"]["render_status"] = {
        "docx": docx_status,
        "pdf": pdf_status,
    }
    _write_final_report_index(paths, report_context, summary, render_status=summary["final_report"]["render_status"])
    paths.summary_json.write_text(json.dumps(summary, indent=2), encoding="utf-8")

    # Index JSON (hashes)
    index: Dict[str, Any] = {
        "generated_iso": summary.get("generated_iso"),
        "files": [],
    }
    for p in [
        paths.summary_json,
        paths.report_context_json,
        paths.final_report_md,
        paths.final_report_docx,
        paths.final_report_pdf,
        paths.final_report_index_json,
        paths.ftir_validation_appendix_cover_md,
        paths.ftir_validation_appendix_ledger_csv,
        paths.ftir_validation_appendix_method301_csv,
        paths.ftir_validation_appendix_exclusions_csv,
        paths.ftir_validation_appendix_index_json,
        paths.evidence_bundle_json,
        paths.evidence_step_eval_csv,
        paths.analyzer_validity_csv,
        paths.fuel_analysis_csv,
        paths.pollutant_adjustments_csv,
        paths.pollutant_adjustment_history_csv,
        paths.spike_recovery_csv,
        paths.side_by_side_csv,
        paths.ftir_validation_json,
        paths.ftir_validation_windows_csv,
        paths.ftir_validation_method301_csv,
        paths.test_matrix_csv,
        paths.qaqc_events_csv,
        paths.regulatory_snapshot_csv,
        paths.reference_trace_csv,
        paths.reference_acceptance_csv,
        paths.reference_sources_csv,
        paths.reference_latest_json,
        paths.reference_offsets_csv,
        paths.reference_offsets_json,
        paths.pdf_path,
    ]:
        if not p.exists():
            continue
        try:
            index["files"].append({
                "name": p.name,
                "sha256": _sha256_file(p),
                "bytes": p.stat().st_size,
            })
        except Exception:
            pass

    paths.index_json.write_text(json.dumps(index, indent=2), encoding="utf-8")

    return out_dir


def _write_pdf(paths: ReportPackPaths, summary: Dict[str, Any]) -> None:
    """Best-effort PDF generation (reportlab optional)."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
    except Exception:
        return

    try:
        styles = getSampleStyleSheet()
        doc = SimpleDocTemplate(str(paths.pdf_path), pagesize=letter)
        story: List[Any] = []

        # Logo (best-effort)
        try:
            here = Path(__file__).resolve().parent
            logo = here.parent / "mole_assets" / "branding" / "mole_logo_130.png"
            if logo.exists():
                story.append(Image(str(logo), width=130, height=130))
                story.append(Spacer(1, 8))
        except Exception:
            pass

        story.append(Paragraph("MOLE DAS - Report Pack v1", styles["Title"]))
        story.append(Spacer(1, 12))

        sess = summary.get("session") or {}
        story.append(Paragraph(f"Job ID: {sess.get('job_id') or ''}", styles["Normal"]))
        story.append(Paragraph(f"Run ID: {sess.get('run_id') or ''}", styles["Normal"]))
        story.append(Paragraph(f"Generated: {summary.get('generated_iso')}", styles["Normal"]))
        story.append(Paragraph(f"Overall: {sess.get('overall_pass')}", styles["Normal"]))
        story.append(Spacer(1, 12))

        # Test matrix table
        tm = (summary.get("test_matrix") or {}).get("tests") or []
        story.append(Paragraph("Test Matrix Scorecard", styles["Heading2"]))
        if not tm:
            story.append(Paragraph("(no test matrix present)", styles["Normal"]))
        else:
            data = [["ID", "Req", "Perf", "Waived", "Pts", "Min", "Status", "Details"]]
            for t in tm:
                data.append([
                    t.get("id"),
                    "Y" if t.get("required") else "N",
                    "Y" if t.get("performed") else "N",
                    "Y" if t.get("waived") else "N",
                    str(t.get("points")),
                    str(t.get("min_points")),
                    t.get("status"),
                    t.get("details"),
                ])
            tbl = Table(data, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("FONT", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(tbl)

        story.append(Spacer(1, 12))

        # QA/QC summary
        qaqc = (summary.get("qaqc") or {}).get("summary") or {}
        story.append(Paragraph("QA/QC Trending Summary", styles["Heading2"]))
        story.append(Paragraph(f"Events captured: {qaqc.get('event_count')}", styles["Normal"]))
        story.append(Paragraph(f"Overall OK: {qaqc.get('overall_ok')}", styles["Normal"]))
        story.append(Spacer(1, 6))

        per_pol = qaqc.get("per_pollutant") or {}
        if per_pol:
            data = [["Pollutant", "PRE ZERO", "PRE SPAN", "POST ZERO", "POST SPAN", "DRIFT", "OK"]]
            for pol, b in per_pol.items():
                data.append([
                    pol,
                    str(b.get("pre_zero")),
                    str(b.get("pre_span")),
                    str(b.get("post_zero")),
                    str(b.get("post_span")),
                    str(b.get("drift")),
                    str(b.get("overall_ok")),
                ])
            tbl = Table(data, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("FONT", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]))
            story.append(tbl)
        else:
            story.append(Paragraph("(no QA/QC events found for this run_id)", styles["Normal"]))

        story.append(Spacer(1, 12))
        spike = (summary.get("qaqc") or {}).get("spike_recovery") or {}
        story.append(Paragraph("Spike Recovery / Flow Restriction", styles["Heading2"]))
        if not spike:
            story.append(Paragraph("(no spike recovery summary found)", styles["Normal"]))
        else:
            story.append(Paragraph(
                f"Enabled: {spike.get('enabled')}  Mandatory: {spike.get('mandatory')}  Applies to FTIR: {spike.get('applies_to_ftir')}",
                styles["Normal"],
            ))
            story.append(Paragraph(
                f"Flow restriction: {spike.get('requires_flow_restriction')}  Limit: {_fmt_num(spike.get('flow_restriction_pct'), 4, '')}%",
                styles["Normal"],
            ))
            story.append(Paragraph(
                f"Sample flow: {_fmt_num(spike.get('sample_flow_value'), 6, '(n/a)')} {spike.get('sample_flow_units') or ''}".strip(),
                styles["Normal"],
            ))
            story.append(Paragraph(
                f"Spike cal flow: {_fmt_num(spike.get('spike_flow_value'), 6, '(n/a)')}  Target spike flow: {_fmt_num(spike.get('target_spike_flow_value'), 6, '(n/a)')}",
                styles["Normal"],
            ))
            story.append(Paragraph(
                f"Flow status: {spike.get('flow_status') or '(n/a)'}  Reason: {spike.get('flow_reason') or '(n/a)'}",
                styles["Normal"],
            ))
            story.append(Paragraph(
                f"Method basis: {spike.get('method_label') or spike.get('method_standard') or '(n/a)'}",
                styles["Normal"],
            ))
            spike_rows = spike.get("channels") or []
            if spike_rows:
                data = [["Pollutant", "Spike Amt", "MOLE %", "MOLE", "FTIR %", "FTIR"]]
                for row in spike_rows:
                    data.append([
                        row.get("pollutant"),
                        _fmt_num(row.get("spike_amount"), 6, ""),
                        _fmt_num(row.get("mole_recovery_pct"), 4, ""),
                        row.get("mole_status"),
                        _fmt_num(row.get("ftir_recovery_pct"), 4, ""),
                        row.get("ftir_status"),
                    ])
                tbl = Table(data, repeatRows=1)
                tbl.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("FONT", (0, 0), (-1, 0), "Helvetica-Bold"),
                ]))
                story.append(Spacer(1, 6))
                story.append(tbl)

        story.append(Spacer(1, 12))
        ref = summary.get("reference_audit") or {}
        story.append(Paragraph("Reference / Audit FTIR Traceability", styles["Heading2"]))
        story.append(Paragraph(f"Configured: {ref.get('configured')}", styles["Normal"]))
        ref_ev = ref.get("evidence") or {}
        story.append(Paragraph(f"Evidence mode: {ref_ev.get('mode')}", styles["Normal"]))
        story.append(Paragraph(f"Captured FTIR records: {ref_ev.get('captured_record_count')}", styles["Normal"]))
        ref_acc = ref.get("acceptance") or {}
        story.append(Paragraph(
            f"Acceptance summary: PASS={ref_acc.get('pass_count')}  FAIL={ref_acc.get('fail_count')}  UNKNOWN={ref_acc.get('unknown_count')}",
            styles["Normal"],
        ))
        ref_latest = ref.get("latest") or {}
        if ref_latest:
            story.append(Paragraph(f"Latest status: {ref_latest.get('status')}", styles["Normal"]))
            story.append(Paragraph(f"Source: {ref_latest.get('source_path') or '(none)'}", styles["Normal"]))
            story.append(Paragraph(f"Source SHA-256: {ref_latest.get('source_sha256') or '(n/a)'}", styles["Normal"]))
        else:
            story.append(Paragraph("(no FTIR trace evidence found)", styles["Normal"]))

        ref_poll = ref.get("per_pollutant") or {}
        if ref_poll:
            data = [["Pollutant", "Units", "Ref Cnt", "Delta Cnt", "Latest", "Tol", "|Delta|"]]
            for pol, rec in ref_poll.items():
                acc = rec.get("acceptance") or {}
                basis = acc.get("basis") or {}
                data.append([
                    pol,
                    rec.get("units"),
                    str(rec.get("reference_count")),
                    str(rec.get("delta_count")),
                    str(acc.get("latest_status")),
                    _fmt_num(basis.get("tolerance_abs"), 3, ""),
                    _fmt_num(acc.get("latest_delta_abs"), 3, ""),
                ])
            tbl = Table(data, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("FONT", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]))
            story.append(Spacer(1, 6))
            story.append(tbl)

        story.append(Spacer(1, 12))
        fuel = summary.get("fuel_analysis") or {}
        story.append(Paragraph("Fuel Analysis Methodology and Source Traceability", styles["Heading2"]))
        if not fuel:
            story.append(Paragraph("(no fuel analysis summary found)", styles["Normal"]))
        else:
            story.append(Paragraph(
                f"Methodology basis: {fuel.get('methodology_basis_label') or fuel.get('methodology_basis_code') or '(n/a)'}",
                styles["Normal"],
            ))
            story.append(Paragraph(
                f"Fuel family/profile: {fuel.get('family') or '(n/a)'} / {fuel.get('profile_label') or fuel.get('profile_id') or '(n/a)'}",
                styles["Normal"],
            ))
            story.append(Paragraph(
                f"Fuel basis source: {fuel.get('default_source_display') or '(n/a)'}",
                styles["Normal"],
            ))
            default_source_note = str(fuel.get("default_source_note") or "").strip()
            if default_source_note:
                story.append(Paragraph(
                    f"Fuel source note: {default_source_note}",
                    styles["Normal"],
                ))
            story.append(Paragraph(
                f"Capture source: {fuel.get('source') or '(n/a)'}  Captured: {fuel.get('captured')}",
                styles["Normal"],
            ))
            fuel_rows = [
                ["Property", "Value"],
                ["Relative density (air=1)", _fmt_num(fuel.get("relative_density_air"), 6, "")],
                ["Compressibility basis", str(fuel.get("compressibility_basis") or "")],
                ["Compressibility factor", _fmt_num(fuel.get("compressibility_factor"), 6, "")],
                ["HHV (Btu/scf)", _fmt_num(fuel.get("hhv_btu_scf"), 6, "")],
                ["LHV (Btu/scf)", _fmt_num(fuel.get("lhv_btu_scf"), 6, "")],
                ["Selected F-factor", _fmt_num(fuel.get("f_factor_selected"), 6, "")],
            ]
            tbl = Table(fuel_rows, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("FONT", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]))
            story.append(Spacer(1, 6))
            story.append(tbl)
            traceability_note = str(fuel.get("traceability_note") or "").strip()
            if traceability_note:
                story.append(Spacer(1, 6))
                story.append(Paragraph(f"Traceability note: {traceability_note}", styles["Normal"]))

        story.append(Spacer(1, 12))
        padj = summary.get("pollutant_adjustments") or {}
        story.append(Paragraph("Pollutant Bias / Drift Adjustments", styles["Heading2"]))
        story.append(Paragraph(
            f"Adjustments enabled: {'YES' if padj.get('enabled') else 'NO'}",
            styles["Normal"],
        ))
        story.append(Paragraph(
            f"Drift basis: {padj.get('drift_basis') or '(n/a)'}",
            styles["Normal"],
        ))
        story.append(Paragraph(
            f"Elapsed hours reference: {_fmt_num(padj.get('elapsed_hours'), 4, '(n/a)')}",
            styles["Normal"],
        ))
        story.append(Paragraph(
            f"Formula: {padj.get('formula') or '(n/a)'}",
            styles["Normal"],
        ))
        note = str(padj.get("note") or "").strip()
        if note:
            story.append(Paragraph(note, styles["Normal"]))
        padj_rows = padj.get("rows") if isinstance(padj.get("rows"), list) else []
        if not padj_rows:
            story.append(Paragraph("(no pollutant-specific adjustments configured)", styles["Normal"]))
        else:
            data = [["Pollutant", "Enabled", "State", "Source", "Effective", "Scope", "Units", "Bias", "Drift / hr", "Elapsed hr", "Total adj"]]
            for row in padj_rows:
                if not isinstance(row, dict):
                    continue
                data.append([
                    str(row.get("pollutant") or ""),
                    "YES" if row.get("enabled") else "NO",
                    str(row.get("state") or ""),
                    str(row.get("source_label") or row.get("source") or ""),
                    str(row.get("effective_label") or ""),
                    str(row.get("scope_summary") or row.get("scope_label") or row.get("scope") or ""),
                    str(row.get("units") or ""),
                    _fmt_num(row.get("bias"), 6, ""),
                    _fmt_num(row.get("drift_per_hr"), 6, ""),
                    _fmt_num(row.get("elapsed_hours"), 4, ""),
                    _fmt_num(row.get("total_adjustment"), 6, ""),
                ])
            tbl = Table(data, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("FONT", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]))
            story.append(Spacer(1, 6))
            story.append(tbl)
            note_rows = [row for row in padj_rows if isinstance(row, dict) and str(row.get("note") or "").strip()]
            if note_rows:
                story.append(Spacer(1, 6))
                story.append(Paragraph("Operator adjustment notes", styles["Normal"]))
                for row in note_rows:
                    stamp_parts = []
                    if str(row.get("updated_by") or "").strip():
                        stamp_parts.append(f"by {str(row.get('updated_by') or '').strip()}")
                    if str(row.get("updated_iso") or "").strip():
                        stamp_parts.append(f"at {str(row.get('updated_iso') or '').strip()}")
                    stamp_txt = f" ({' '.join(stamp_parts)})" if stamp_parts else ""
                    story.append(Paragraph(
                        f"{str(row.get('pollutant') or '(n/a)')}: {str(row.get('note') or '').strip()}{stamp_txt}",
                        styles["Normal"],
                    ))
        latest_review = padj.get("latest_review") if isinstance(padj.get("latest_review"), dict) else {}
        if latest_review:
            story.append(Spacer(1, 6))
            story.append(Paragraph("Latest post-cal decision review", styles["Normal"]))
            story.append(Paragraph(str(latest_review.get("headline") or "(no headline)"), styles["Normal"]))
            review_message = str(latest_review.get("message") or "").strip()
            if review_message:
                story.append(Paragraph(review_message, styles["Normal"]))
            story.append(Paragraph(
                f"Overall pass: {latest_review.get('overall_pass')}  Health: {latest_review.get('postcal_health') or '(n/a)'}",
                styles["Normal"],
            ))
            for label, key in [
                ("Promoted", "promoted_codes"),
                ("Manual only", "manual_only_codes"),
                ("Invalidated", "invalidated_codes"),
                ("Suspended", "suspended_codes"),
            ]:
                codes = latest_review.get(key)
                code_text = ", ".join(str(code).strip().upper() for code in (codes or []) if str(code).strip()) or "(none)"
                story.append(Paragraph(f"{label}: {code_text}", styles["Normal"]))
        history_rows = padj.get("history") if isinstance(padj.get("history"), list) else []
        if history_rows:
            story.append(Spacer(1, 6))
            story.append(Paragraph("Adjustment provenance history", styles["Normal"]))
            data = [["Run", "Pollutant", "Event", "Policy", "Detail"]]
            for row in history_rows[:15]:
                if not isinstance(row, dict):
                    continue
                detail_parts = []
                detail = str(row.get("detail") or "").strip()
                reason = str(row.get("reason") or "").strip()
                if detail:
                    detail_parts.append(detail)
                if reason:
                    detail_parts.append(reason)
                data.append([
                    str(row.get("run_no") or ""),
                    str(row.get("pollutant") or ""),
                    str(row.get("event") or ""),
                    str(row.get("policy_mode") or ""),
                    " | ".join(detail_parts),
                ])
            tbl = Table(data, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("FONT", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(Spacer(1, 6))
            story.append(tbl)
            if len(history_rows) > 15:
                story.append(Spacer(1, 4))
                story.append(Paragraph(
                    f"Showing 15 of {len(history_rows)} adjustment-history rows. See pollutant_adjustment_history.csv for the full ledger.",
                    styles["Normal"],
                ))

        story.append(Spacer(1, 12))
        story.append(Paragraph("This is a v1 report pack. Future versions can add charts and full calculation outputs.", styles["Normal"]))

        doc.build(story)
    except Exception:
        try:
            if paths.pdf_path.exists():
                paths.pdf_path.unlink()
        except Exception:
            pass


# -----------------------------
# CLI wrapper
# -----------------------------


def main() -> int:
    ap = argparse.ArgumentParser(description="Generate MOLE DAS Report Pack v1")
    ap.add_argument("--config", default=None, help="Path to session config JSON (runner_config.json or legacy session JSON)")
    ap.add_argument("--session-dir", default=None, help="Session/run folder (contains meta/, raw/, exports/)")
    ap.add_argument("--db", default=None, help="Override master DB path (mole_master.sqlite)")
    args = ap.parse_args()

    cfg_path: Optional[Path] = None
    if args.config:
        cfg_path = Path(str(args.config)).expanduser().resolve()
        if not cfg_path.exists():
            raise SystemExit(f"Config not found: {cfg_path}")

    sess_dir: Optional[Path] = None
    if args.session_dir:
        sess_dir = Path(str(args.session_dir)).expanduser().resolve()
        if not sess_dir.exists():
            raise SystemExit(f"Session dir not found: {sess_dir}")

    db_path: Optional[Path] = None
    if args.db:
        db_path = Path(str(args.db)).expanduser().resolve()

    if cfg_path is None and sess_dir is None:
        raise SystemExit("Provide --config or --session-dir")

    if cfg_path is not None:
        session = json.loads(cfg_path.read_text(encoding="utf-8-sig"))
    else:
        # Try to find runner_config.json inside session dir
        rc = Path(sess_dir) / "runner_config.json"  # type: ignore[arg-type]
        if not rc.exists():
            raise SystemExit(f"runner_config.json not found in: {sess_dir}")
        cfg_path = rc
        session = json.loads(rc.read_text(encoding="utf-8-sig"))

    out = generate_report_pack_v1(session=session, cfg_path=cfg_path, session_dir=sess_dir, master_db_path=db_path)
    print(f"Report pack written to: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
